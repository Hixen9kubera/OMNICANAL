"""
Radar de precios · F1 — ESTADO POR CONTENEDOR (informe fuera del panel).

¿Cuánto lleva recuperado cada contenedor de su cifra X, a qué ritmo, cuándo
termina, qué SKUs lo frenan y qué palancas (que no son el precio) quedan?
X es UNA cifra total por contenedor (decisión de Eduardo, 28-sep-2026).

SOLO LECTURA DEL SANDBOX
  · La DSN sale de `env.staging` (SUPABASE_DB_URL) en la raíz del worktree y
    SOLO se acepta el sandbox (`yvootpbz`): si trae `tukwcvsi` (producción) en
    cualquier parte, aborta antes de conectar. Nunca se imprime la DSN.
  · Una sola transacción `repeatable read, read only` (por TRANSACCIÓN, jamás
    `SET SESSION` ni `set_session(readonly=True)`: regla 13) sobre el puerto de
    SESIÓN 5432 del mismo host (el 6543 no sostiene transacciones largas).
  · Cero escrituras en cualquier base. Lo único que escribe es el Excel y, la
    primera vez, la plantilla de X; ambos FUERA del repo (el repo es público:
    ninguna cifra de X entra aquí). Si la ruta cae dentro del repo, aborta.

LA CIFRA X
  `x_contenedores.csv` (contenedor,x_total_mxn,fecha_liberacion,fuente,
  capturado_por) en la carpeta de salida. Si no existe se CREA como plantilla
  con todos los contenedores de `costing.sku_contenedor` y X vacío; si existe,
  jamás se sobrescribe. X vacío = «sin dato», nunca 0.

DEFINICIONES (las mismas del §3 de la especificación del radar)
  · Contribución de una venta = precio pagado / (1 + IVA) − comisión − envío.
      comisión: la real de `channel.order_items` (`comision` = sale_fee de ML:
        porcentaje sobre el precio CON IVA; el IVA del propio cargo ML lo
        factura aparte y es acreditable, así que se resta tal cual, sin
        volverle a quitar IVA). `comision = 0` NO es «sin comisión»: en ML es
        la venta cuyo fee aún no llega (paso 0 → real) y en Amazon/TikTok/Temu
        el canal no la reporta. Ahí se ESTIMA con la tasa real media de esa
        publicación, o `comision_estimada` si no hay ninguna.
      envío (solo ML): el real de `enrich.order_shipping_cost` (por pedido,
        repartido entre sus líneas por importe); si ese pedido no lo tiene, el
        promedio real por pieza del mismo SKU; si tampoco, la tabla oficial de
        `services/costos.py` (`calc_fee_envio_ml` con `_peso_efectivo` de la
        pieza, una vez por línea con el peso de todas sus piezas). Un peso más
        denso que `densidad_max_kg_l` es de caja master (871 kg en 101×31×18
        cm): ahí cuenta solo el volumétrico, o la tabla inventa $3,000 por
        envío. Otros canales: sin dato (no se resta).
      Full, publicidad y devoluciones: sin dato → TODA contribución es COTA
      SUPERIOR.
  · Ventas: filtro idéntico a `channel.sales_daily` (fuera cancelled/invalid/
    canceled sin importar la caja; `partially_refunded` se queda, como allá).
    Vivo = `channel.order_items` desde el 16-jul; archivo = `analytics.
    sales_daily_hist` hasta el 15-jul (el mismo empalme que `channel.
    sales_daily_completa`), con su `sale_fee` real y el envío estimado.
  · Atribución: SKU en un solo contenedor → a ese. SKU en varios → cubeta
    «multi: sin asignar»: no hay piezas por contenedor ni fechas confiables
    para un FIFO, así que no se reparte.
  · Recuperado = suma de la contribución desde la liberación (si está
    capturada) o desde el inicio de los datos. Falta = max(0, X − recuperado).
    Ritmo = contribución de las últimas 4 semanas hasta el corte (último día
    con ventas) / semanas. ETA = corte + falta / ritmo.

Uso (desde la raíz del worktree):
  backend/.venv/Scripts/python.exe backend/scripts/radar_estado_contenedores.py
      [--env env.staging] [--salida ../_radar_f0] [--x-csv …/x_contenedores.csv]
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import math
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT / "backend"))

from services import costos  # noqa: E402  (solo funciones puras: tabla de envío ML)
from services.temu import VENDIBLES as TEMU_VENDIBLES  # noqa: E402
from services.ubicar_contenedores import REF_PRODUCCION, REF_SANDBOX, ref_de  # noqa: E402
from services.walmart_panel import ESTADO_VIVO as WALMART_VIVO  # noqa: E402

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass

# ── Parámetros ────────────────────────────────────────────────────────────────
# Cada uno es una DECISIÓN, NO UNA MEDICIÓN: cambiarlo cambia el informe, no la
# realidad. El Excel los imprime en «Notas» para que se lean como tales.
PARAMS: dict[str, Any] = {
    "iva": 0.16,                     # decisión, no medición: IVA MX
    "comision_estimada": 0.175,      # decisión, no medición: tasa si no hay comisión real
    "semanas_ritmo": 4,              # decisión, no medición: ventana del ritmo
    "horizonte_exceso_dias": 180,    # decisión, no medición: cobertura que ya es exceso
    "cobertura_frena_dias": 90,      # decisión, no medición: cobertura que ya frena
    "dias_velocidad": 90,            # decisión, no medición: ventana de ventas/día
    "horizonte_lento_semanas": 26,   # decisión, no medición: ETA más lejana = «lento»
    "rotacion_alta_dia": 0.5,        # decisión, no medición: piezas/día que ya pide Full
    "visitas_min": 300,              # decisión, no medición: «muchas visitas» (30 d)
    "conversion_baja": 0.01,         # decisión, no medición: «poca conversión» (u/visita)
    "densidad_max_kg_l": 1.5,        # decisión, no medición: más denso = peso de caja master
    "min_piezas_promedio_envio": 3,  # decisión, no medición: piezas con envío real para promediar
}
# SKUs de recompra: lo decide el negocio. Vacío a propósito.
RECOMPRA: set[str] = set()

INICIO_VIVO = dt.date(2026, 7, 16)   # empalme de channel.sales_daily_completa (acta 59)
MULTI = "multi: sin asignar"
SIN_CONTENEDOR = "sin contenedor (SKU no ubicado)"
COLUMNAS_X = ("contenedor", "x_total_mxn", "fecha_liberacion", "fuente", "capturado_por")
SIN_DATO = "sin dato"


# ══════════════════════════════════════════════════════════════════════════════
# Funciones PURAS (probadas en tests/test_radar_estado_contenedores.py)
# ══════════════════════════════════════════════════════════════════════════════

def validar_dsn(dsn: str) -> str:
    """La DSN del sandbox o SystemExit. Nunca incluye la DSN en el mensaje."""
    dsn = (dsn or "").strip()
    if not dsn:
        raise SystemExit("ABORT: env.staging no trae SUPABASE_DB_URL.")
    if REF_PRODUCCION in dsn.lower():
        raise SystemExit("ABORT: la DSN apunta a PRODUCCIÓN (tukwcvsi). Este informe solo lee el sandbox.")
    if not ref_de(dsn).startswith(REF_SANDBOX):
        raise SystemExit("ABORT: la DSN no es la del sandbox (yvootpbz).")
    return dsn


def a_puerto_sesion(dsn: str) -> str:
    """El pooler en modo transacción (6543) → el de sesión (5432) del mismo host."""
    return re.sub(r":6543(?=[/?]|$)", ":5432", dsn)


def atribuir(filas: Iterable[dict]) -> dict[str, int | str]:
    """{sku: número de contenedor | MULTI}. Multi si el SKU trae más de un
    número o cualquiera de sus filas dice `multi` (otra N que no se cargó)."""
    nums: dict[str, set[int]] = defaultdict(set)
    marcado: set[str] = set()
    for f in filas:
        sku = str(f["sku"]).upper()
        nums[sku].add(int(f["numero"]))
        if f.get("multi"):
            marcado.add(sku)
    return {s: (MULTI if len(n) > 1 or s in marcado else next(iter(n)))
            for s, n in nums.items()}


def cubeta_de(sku: str | None, atribucion: dict[str, int | str]) -> int | str:
    return atribucion.get(str(sku or "").upper(), SIN_CONTENEDOR)


def comision_de(ingreso: float, comision: float | None, tasa_publicacion: float | None,
                estimada: float) -> tuple[float, str]:
    """(comisión, 'real'|'estimado'). 0 o None = aún no llega o el canal no la da."""
    if comision is not None and comision > 0:
        return float(comision), "real"
    tasa = tasa_publicacion if tasa_publicacion is not None else estimada
    return ingreso * tasa, "estimado"


def peso_dudoso(medidas: dict | None, densidad_max: float) -> bool:
    """Peso de CAJA MASTER registrado como de pieza: densidad > `densidad_max`
    kg/L (las medidas sí son de la pieza; ver memoria «peso de caja en
    costos_validados»: 871 kg en 101×31×18 cm)."""
    if not medidas:
        return False
    vol_l = (float(medidas.get("largo") or 0) * float(medidas.get("ancho") or 0)
             * float(medidas.get("alto") or 0)) / 1000.0
    return vol_l > 0 and float(medidas.get("peso") or 0) / vol_l > densidad_max


def envio_estimado_ml(medidas: dict | None, cantidad: int, ingreso: float,
                      densidad_max: float = 1.5) -> tuple[float, str]:
    """Tabla oficial de `services/costos.py`: un envío por línea con el peso
    efectivo de todas sus piezas y el importe de la línea como tramo. Si el
    peso es de caja master, se usa solo el volumétrico de la pieza."""
    if medidas and any(float(medidas.get(k) or 0) > 0 for k in ("peso", "largo", "ancho", "alto")):
        dudoso = peso_dudoso(medidas, densidad_max)
        peso, _ = costos._peso_efectivo(0.0 if dudoso else float(medidas.get("peso") or 0),
                                        float(medidas.get("largo") or 0),
                                        float(medidas.get("ancho") or 0),
                                        float(medidas.get("alto") or 0))
        estado = "tabla_peso_dudoso" if dudoso else "tabla"
    else:
        peso, _ = costos._peso_efectivo(0.0, 0.0, 0.0, 0.0)   # 0.5 kg por defecto del motor
        estado = "tabla_sin_medidas"
    return costos.calc_fee_envio_ml(peso * max(int(cantidad or 1), 1), ingreso), estado


def promedio_envio_real(lineas: Iterable[tuple[str, int, float]], min_piezas: int) -> dict[str, float]:
    """{sku: envío real por pieza} de las líneas con envío real (sku, piezas,
    envío de la línea). Solo con al menos `min_piezas` piezas medidas."""
    tot: dict[str, list[float]] = defaultdict(lambda: [0.0, 0])
    for sku, piezas, envio in lineas:
        tot[sku][0] += envio
        tot[sku][1] += piezas
    return {s: e / n for s, (e, n) in tot.items() if n >= min_piezas and n > 0}


def contribucion(ingreso: float, comision: float, envio: float | None, iva: float) -> float:
    """precio pagado / (1 + IVA) − comisión − envío (envío None = sin dato, no se resta)."""
    return ingreso / (1 + iva) - comision - (envio or 0.0)


def estado_contribucion(comision_estado: str, envio_estado: str) -> str:
    reales = (comision_estado == "real") + (envio_estado == "real")
    return ("estimado", "parcial", "real")[reales]


def desde_efectivo(fecha_liberacion: dt.date | None, inicio_datos: dt.date) -> tuple[dt.date, bool]:
    """(desde cuándo se cuenta, ¿historia truncada?). Truncada = no hay fecha de
    liberación o es anterior al inicio de los datos: lo de antes no se ve."""
    if fecha_liberacion is None:
        return inicio_datos, True
    if fecha_liberacion < inicio_datos:
        return inicio_datos, True
    return fecha_liberacion, False


def recuperado(serie: Iterable[tuple[dt.date, float]], desde: dt.date,
               hasta: dt.date | None = None) -> float:
    return sum(v for d, v in serie if d >= desde and (hasta is None or d <= hasta))


def ritmo_semanal(serie: Iterable[tuple[dt.date, float]], corte: dt.date, semanas: int,
                  desde: dt.date | None = None) -> float:
    """Contribución por semana en las últimas `semanas` hasta el corte. Si la
    liberación cae dentro de la ventana, se divide solo entre lo transcurrido."""
    inicio = corte - dt.timedelta(days=7 * semanas - 1)
    if desde is not None and desde > inicio:
        inicio = desde
    dias = (corte - inicio).days + 1
    if dias <= 0:
        return 0.0
    return recuperado(serie, inicio, corte) / (dias / 7.0)


def falta_de(x: float | None, rec: float) -> float | None:
    """None si no hay X (sin dato ≠ 0)."""
    if x is None:
        return None
    return max(0.0, x - rec)


def eta_de(falta: float | None, ritmo: float, corte: dt.date) -> tuple[dt.date | None, float | None]:
    """(fecha estimada, semanas). None si no hay X, ya se recuperó o no hay ritmo."""
    if falta is None or falta <= 0 or ritmo <= 0:
        return None, None
    semanas = falta / ritmo
    return corte + dt.timedelta(days=math.ceil(semanas * 7)), semanas


def estado_contenedor(x: float | None, falta: float | None, ritmo: float,
                      semanas_eta: float | None, horizonte_lento: float) -> str:
    if x is None:
        return "sin X"
    if falta is not None and falta <= 0:
        return "recuperado"
    if ritmo <= 0:
        return "sin ritmo"
    if semanas_eta is not None and semanas_eta > horizonte_lento:
        return "lento"
    return "en curso"


def cobertura_dias(stock: float, ventas_dia: float) -> float | None:
    """Días de stock al ritmo actual; None = sin ventas (infinita)."""
    if ventas_dia <= 0:
        return None
    return stock / ventas_dia


def clasificar(sku: str, stock: float, unidades_90d: float, cobertura: float | None,
               horizonte_exceso: float) -> str:
    """exceso (0 ventas en 90 d con stock, o cobertura > horizonte) · recompra
    (lista del negocio) · normal. El mismo orden del §3 de la especificación."""
    if stock > 0 and (unidades_90d <= 0 or (cobertura is not None and cobertura > horizonte_exceso)):
        return "exceso"
    if str(sku).upper() in RECOMPRA:
        return "recompra"
    return "normal"


def parsear_monto(texto: str | None) -> float | None:
    """'1,234,567.89' · '$1234567' · '' → float | None. Negativo o 0 → ValueError."""
    s = (texto or "").strip().replace("$", "").replace(",", "").replace(" ", "")
    if not s:
        return None
    v = float(s)
    if v <= 0:
        raise ValueError("X debe ser positivo")
    return v


def parsear_fecha(texto: str | None) -> dt.date | None:
    s = (texto or "").strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"fecha no reconocida: {s!r} (usa AAAA-MM-DD)")


def numero_contenedor(texto: str) -> int | None:
    m = re.search(r"\d+", str(texto or ""))
    return int(m.group()) if m else None


def crear_plantilla_x(ruta: Path, contenedores: Iterable[int]) -> bool:
    """Crea la plantilla SOLO si no existe (modo 'x': jamás sobrescribe).
    True si la creó."""
    try:
        with open(ruta, "x", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(COLUMNAS_X)
            for n in sorted(set(contenedores)):
                w.writerow([n, "", "", "", ""])
        return True
    except FileExistsError:
        return False


def leer_x(ruta: Path) -> dict[int, dict]:
    """{número: {x, fecha_liberacion, fuente, capturado_por}}. Errores con el
    contenedor y la columna, nunca con la cifra."""
    salida: dict[int, dict] = {}
    with open(ruta, encoding="utf-8-sig", newline="") as fh:
        for i, fila in enumerate(csv.DictReader(fh), start=2):
            n = numero_contenedor(fila.get("contenedor", ""))
            if n is None:
                continue
            try:
                x = parsear_monto(fila.get("x_total_mxn"))
            except ValueError:
                raise SystemExit(f"x_contenedores.csv renglón {i} (contenedor {n}): x_total_mxn no es un monto positivo.")
            try:
                fecha = parsear_fecha(fila.get("fecha_liberacion"))
            except ValueError as e:
                raise SystemExit(f"x_contenedores.csv renglón {i} (contenedor {n}): {e}")
            salida[n] = {"x": x, "fecha_liberacion": fecha,
                         "fuente": (fila.get("fuente") or "").strip(),
                         "capturado_por": (fila.get("capturado_por") or "").strip()}
    return salida


def dentro_de(ruta: Path, raiz: Path) -> bool:
    try:
        ruta.resolve().relative_to(raiz.resolve())
        return True
    except ValueError:
        return False


# ══════════════════════════════════════════════════════════════════════════════
# Lectura del sandbox (una sola foto)
# ══════════════════════════════════════════════════════════════════════════════

_FILTRO_VENTA = "lower(coalesce(o.estado_canal, '')) not in ('cancelled', 'invalid', 'canceled')"

CONSULTAS: dict[str, str] = {
    "sku_contenedor": """select sku::text as sku, numero, codigo, nivel, multi
                           from costing.sku_contenedor""",
    "productos": """select sku::text as sku, name, wc_id, wc_parent_id from core.products""",
    "medidas": """select sku::text as sku, peso, largo, ancho, alto from costing.costos_validados""",
    "ventas_vivo": f"""
        select i.canal, i.cuenta, i.external_order_id, i.linea, i.item_id, i.sku::text as sku,
               i.cantidad, i.precio_unitario, i.comision,
               (o.creado_at at time zone 'America/Mexico_City')::date as fecha
          from channel.order_items i
          join channel.orders o using (canal, cuenta, external_order_id)
         where {_FILTRO_VENTA}
           and (o.creado_at at time zone 'America/Mexico_City')::date >= %(inicio_vivo)s""",
    "vivo_sin_sku": f"""
        select count(*) as lineas, coalesce(sum(i.cantidad), 0) as piezas
          from channel.order_items i
          join channel.orders o using (canal, cuenta, external_order_id)
         where {_FILTRO_VENTA} and i.sku is null
           and (o.creado_at at time zone 'America/Mexico_City')::date >= %(inicio_vivo)s""",
    # costo_vendedor NULL con fila = «ML dijo que no hay costo» (0020): es 0 real.
    "envios": """select cuenta, external_order_id, coalesce(costo_vendedor, 0) as costo,
                        consultado_at::date as consultado
                   from enrich.order_shipping_cost""",
    "archivo": """select date as fecha, cuenta, item_id, sku::text as sku, units_sold, revenue, sale_fee
                    from analytics.sales_daily_hist
                   where date < %(inicio_vivo)s and sku is not null""",
    "listings": """select l.sku::text as sku, l.canal, a.legacy_code as cuenta, l.listing_id,
                          l.status, l.situacion, l.stock_own, l.stock_full, l.stock_fba,
                          l.is_fulfillment, l.logistic_type, l.price, l.price_sale
                     from channel.listings l
                     left join core.accounts a on a.id = l.account_id""",
    "publicaciones": """select sku::text as sku, cuenta, ml_item_id, estado, precio,
                               visitas_30d, unidades_30d, periodo
                          from enrich.market_publicaciones_v""",
}


def leer_sandbox(dsn: str) -> dict[str, list[dict]]:
    import psycopg2
    import psycopg2.extras

    params = {"inicio_vivo": INICIO_VIVO}
    try:
        conn = psycopg2.connect(a_puerto_sesion(dsn), connect_timeout=30)
    except psycopg2.OperationalError as e:
        raise SystemExit(f"ABORT: no conecta al sandbox ({type(e).__name__}).")
    salida: dict[str, list[dict]] = {}
    try:
        cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cur.execute("set transaction isolation level repeatable read, read only")
        cur.execute("set local statement_timeout = '180s'")
        cur.execute("select to_regclass('costing.sku_contenedor') is not null as hay")
        if not cur.fetchone()["hay"]:
            raise SystemExit("ABORT: el sandbox no tiene costing.sku_contenedor (0060). "
                             "¿Se re-clonó core.products? Ver el encabezado de la 0060.")
        for nombre, sql in CONSULTAS.items():
            cur.execute(sql, params)
            salida[nombre] = [dict(r) for r in cur.fetchall()]
    finally:
        conn.rollback()
        conn.close()
    return salida


# ══════════════════════════════════════════════════════════════════════════════
# Armado
# ══════════════════════════════════════════════════════════════════════════════

def _f(v: Any) -> float:
    return float(v) if v is not None else 0.0


def ventas_con_contribucion(datos: dict[str, list[dict]]) -> list[dict]:
    """Cada venta (línea viva o día×publicación del archivo) con su contribución."""
    iva, est = PARAMS["iva"], PARAMS["comision_estimada"]
    medidas = {str(m["sku"]).upper(): m for m in datos["medidas"]}
    envios = {(e["cuenta"], e["external_order_id"]): _f(e["costo"]) for e in datos["envios"]}

    # Tasa real media por publicación (ponderada por importe), vivo + archivo.
    num: dict[tuple, float] = defaultdict(float)
    den: dict[tuple, float] = defaultdict(float)
    for v in datos["ventas_vivo"]:
        ing = _f(v["precio_unitario"]) * int(v["cantidad"] or 0)
        if _f(v["comision"]) > 0 and ing > 0:
            k = (v["canal"], v["cuenta"], v["item_id"])
            num[k] += _f(v["comision"])
            den[k] += ing
    for h in datos["archivo"]:
        if _f(h["sale_fee"]) > 0 and _f(h["revenue"]) > 0:
            k = ("mercado_libre", h["cuenta"], h["item_id"])
            num[k] += _f(h["sale_fee"])
            den[k] += _f(h["revenue"])
    tasa = {k: num[k] / den[k] for k in num if den[k] > 0}

    # Importe por pedido (solo las líneas con SKU) para repartir el envío real.
    importe_pedido: dict[tuple, float] = defaultdict(float)
    lineas_pedido: dict[tuple, int] = defaultdict(int)
    for v in datos["ventas_vivo"]:
        if v["sku"] is None:
            continue
        k = (v["cuenta"], v["external_order_id"])
        importe_pedido[k] += _f(v["precio_unitario"]) * int(v["cantidad"] or 0)
        lineas_pedido[k] += 1

    def envio_real(v: dict) -> float | None:
        """El envío real de la línea (su parte del pedido), o None."""
        k = (v["cuenta"], v["external_order_id"])
        if v["canal"] != "mercado_libre" or k not in envios:
            return None
        tot = importe_pedido[k]
        ing = _f(v["precio_unitario"]) * int(v["cantidad"] or 0)
        return envios[k] * ((ing / tot) if tot > 0 else 1.0 / max(lineas_pedido[k], 1))

    # Envío real por pieza de cada SKU: la mejor estimación para SUS ventas sin
    # envío propio (la misma idea que el radar: el promedio real del item).
    promedio = promedio_envio_real(
        ((str(v["sku"]).upper(), int(v["cantidad"] or 0), e) for v in datos["ventas_vivo"]
         if v["sku"] is not None and (e := envio_real(v)) is not None),
        PARAMS["min_piezas_promedio_envio"])

    def envio_estimado(sku: str, cant: int, ing: float) -> tuple[float, str]:
        if sku in promedio:
            return promedio[sku] * cant, "promedio_real"
        return envio_estimado_ml(medidas.get(sku), cant, ing, PARAMS["densidad_max_kg_l"])

    ventas: list[dict] = []
    for v in datos["ventas_vivo"]:
        if v["sku"] is None:
            continue
        sku = str(v["sku"]).upper()
        cant = int(v["cantidad"] or 0)
        ing = _f(v["precio_unitario"]) * cant
        com, com_e = comision_de(ing, v["comision"], tasa.get((v["canal"], v["cuenta"], v["item_id"])), est)
        if v["canal"] == "mercado_libre":
            env = envio_real(v)
            env_e = "real"
            if env is None:
                env, env_e = envio_estimado(sku, cant, ing)
        else:
            env, env_e = None, "sin_dato"
        ventas.append({"fecha": v["fecha"], "canal": v["canal"], "cuenta": v["cuenta"],
                       "item_id": v["item_id"], "sku": str(v["sku"]).upper(), "unidades": cant,
                       "ingreso": ing, "comision": com, "comision_estado": com_e,
                       "envio": env, "envio_estado": env_e, "fuente": "vivo",
                       "contribucion": contribucion(ing, com, env, iva)})
    for h in datos["archivo"]:
        u = int(h["units_sold"] or 0)
        ing = _f(h["revenue"])
        com, com_e = comision_de(ing, h["sale_fee"], tasa.get(("mercado_libre", h["cuenta"], h["item_id"])), est)
        if u > 0:
            unit, env_e = envio_estimado(str(h["sku"]).upper(), 1, ing / u)   # una pieza por pedido
            env = unit * u
        else:
            env, env_e = 0.0, "tabla"
        ventas.append({"fecha": h["fecha"], "canal": "mercado_libre", "cuenta": h["cuenta"],
                       "item_id": h["item_id"], "sku": str(h["sku"]).upper(), "unidades": u,
                       "ingreso": ing, "comision": com, "comision_estado": com_e,
                       "envio": env, "envio_estado": env_e, "fuente": "archivo",
                       "contribucion": contribucion(ing, com, env, iva)})
    return ventas


def _stock_y_canales(datos: dict[str, list[dict]]) -> tuple[dict, dict]:
    """stock por SKU (Woo manda; espejo con max(), nunca sum()) y presencia por
    canal. La presencia hereda la del PADRE de Woo: la publicación puede vivir en
    el padre mientras las ventas y el stock llegan en el hijo."""
    stock: dict[str, dict] = defaultdict(lambda: {"woo": None, "espejo": None, "full": 0, "fba": 0})
    pres: dict[str, set[str]] = defaultdict(set)
    for l in datos["listings"]:
        s = str(l["sku"]).upper()
        c = l["canal"]
        st = stock[s]
        if l["stock_own"] is not None:
            if c == "general":
                st["woo"] = max(st["woo"] or 0, int(l["stock_own"]))
            else:
                st["espejo"] = max(st["espejo"] or 0, int(l["stock_own"]))
        if c == "mercado_libre":
            st["full"] += int(l["stock_full"] or 0)
            if (l["situacion"] or "").lower() == "active":
                pres[s].add("ml_activa")
            if l["is_fulfillment"] or (l["logistic_type"] or "") == "fulfillment":
                pres[s].add("ml_full")
        elif c == "amazon":
            st["fba"] += int(l["stock_fba"] or 0)
        elif c == "tiktok" and (l["status"] or "").upper() == "ACTIVATE":
            pres[s].add("tiktok")
        elif c == "temu" and (l["status"] or "") in TEMU_VENDIBLES:
            pres[s].add("temu")
        elif c == "walmart" and (l["status"] or "").upper() == WALMART_VIVO:
            pres[s].add("walmart")
    por_wc = {p["wc_id"]: str(p["sku"]).upper() for p in datos["productos"] if p["wc_id"]}
    padre = {str(p["sku"]).upper(): por_wc.get(p["wc_parent_id"])
             for p in datos["productos"] if p["wc_parent_id"]}
    presencia = {}
    for s in set(pres) | set(padre):
        presencia[s] = pres.get(s, set()) | pres.get(padre.get(s) or "", set())
    total = {}
    for s, st in stock.items():
        propio = st["woo"] if st["woo"] is not None else (st["espejo"] or 0)
        total[s] = {"propio": propio, "full": st["full"], "fba": st["fba"],
                    "total": propio + st["full"] + st["fba"]}
    return total, presencia


def armar(datos: dict[str, list[dict]], x_por_cont: dict[int, dict], generado: dt.datetime) -> dict:
    ventas = ventas_con_contribucion(datos)
    atrib = atribuir(datos["sku_contenedor"])
    stock, presencia = _stock_y_canales(datos)
    titulo = {str(p["sku"]).upper(): p["name"] for p in datos["productos"]}

    conts_de: dict[str, list[int]] = defaultdict(list)
    codigos: dict[int, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    skus_de: dict[int, set[str]] = defaultdict(set)
    for f in datos["sku_contenedor"]:
        s, n = str(f["sku"]).upper(), int(f["numero"])
        conts_de[s].append(n)
        skus_de[n].add(s)
        if f["codigo"]:
            codigos[n][f["codigo"]] += 1

    vivas = [v for v in ventas if v["fuente"] == "vivo"]
    corte = max(v["fecha"] for v in vivas)
    inicio = min(v["fecha"] for v in ventas)
    d90 = corte - dt.timedelta(days=PARAMS["dias_velocidad"] - 1)

    # Series por cubeta y agregados por SKU.
    serie: dict[Any, list] = defaultdict(list)
    por_sku: dict[str, dict] = defaultdict(lambda: {"u90": 0, "c90": 0.0, "u": 0, "c": 0.0, "ultima": None})
    cobertura_envio: dict[Any, list[int]] = defaultdict(lambda: [0, 0, 0])   # piezas, envío real, comisión real
    for v in ventas:
        cub = cubeta_de(v["sku"], atrib)
        serie[cub].append((v["fecha"], v["contribucion"], v["unidades"], v["fuente"], v["sku"]))
        a = por_sku[v["sku"]]
        a["u"] += v["unidades"]
        a["c"] += v["contribucion"]
        if v["fecha"] >= d90:
            a["u90"] += v["unidades"]
            a["c90"] += v["contribucion"]
        a["ultima"] = max(a["ultima"] or v["fecha"], v["fecha"])

    def info_sku(s: str) -> dict:
        a = por_sku.get(s, {"u90": 0, "c90": 0.0, "u": 0, "c": 0.0, "ultima": None})
        st = stock.get(s, {"propio": 0, "full": 0, "fba": 0, "total": 0})
        vd = a["u90"] / PARAMS["dias_velocidad"]
        cob = cobertura_dias(st["total"], vd)
        media = (a["c90"] / a["u90"]) if a["u90"] > 0 else ((a["c"] / a["u"]) if a["u"] > 0 else None)
        return {**st, "u90": a["u90"], "ventas_dia": vd, "cobertura": cob, "ultima": a["ultima"],
                "contrib_media": media,
                "clase": clasificar(s, st["total"], a["u90"], cob, PARAMS["horizonte_exceso_dias"])}

    sku_info = {s: info_sku(s) for s in conts_de}

    filas_cont: list[dict] = []
    numeros = sorted(set(skus_de) | set(x_por_cont))
    for n in numeros:
        xi = x_por_cont.get(n, {})
        x = xi.get("x")
        desde, truncada = desde_efectivo(xi.get("fecha_liberacion"), inicio)
        s_n = [(d, c) for d, c, *_ in serie.get(n, [])]
        rec = recuperado(s_n, desde)
        rec_arch = recuperado([(d, c) for d, c, _, fu, _ in serie.get(n, []) if fu == "archivo"], desde)
        piezas = sum(u for d, _, u, *_ in serie.get(n, []) if d >= desde)
        skus_venta = {sk for d, *_, sk in serie.get(n, []) if d >= desde}
        ritmo = ritmo_semanal(s_n, corte, PARAMS["semanas_ritmo"], desde)
        falta = falta_de(x, rec)
        eta, semanas = eta_de(falta, ritmo, corte)
        solos = [s for s in skus_de.get(n, set()) if atrib.get(s) == n]
        multis = [s for s in skus_de.get(n, set()) if atrib.get(s) == MULTI]
        stock_n = sum(sku_info[s]["total"] for s in solos)
        potencial = sum(sku_info[s]["total"] * sku_info[s]["contrib_media"] for s in solos
                        if sku_info[s]["total"] > 0 and sku_info[s]["contrib_media"] is not None)
        sin_hist = sum(1 for s in solos if sku_info[s]["total"] > 0 and sku_info[s]["contrib_media"] is None)
        if x is None:
            alcanza = "sin X"
        elif falta is not None and falta <= 0:
            alcanza = "ya recuperado"
        else:
            alcanza = "sí" if potencial >= (falta or 0) else "no"
        filas_cont.append({
            "contenedor": n,
            "codigo": max(codigos[n], key=codigos[n].get) if codigos.get(n) else "",
            "skus": len(skus_de.get(n, set())), "skus_multi": len(multis),
            "skus_con_venta": len(skus_venta), "piezas": piezas,
            "x": x, "fecha_liberacion": xi.get("fecha_liberacion"),
            "desde": desde, "truncada": truncada,
            "recuperado": rec, "recuperado_archivo": rec_arch,
            "falta": falta, "ritmo": ritmo, "eta": eta, "semanas": semanas,
            "estado": estado_contenedor(x, falta, ritmo, semanas, PARAMS["horizonte_lento_semanas"]),
            "stock": stock_n, "potencial": potencial, "sin_hist": sin_hist, "alcanza": alcanza,
            "multi_compartida": sum(por_sku[s]["c"] for s in multis if s in por_sku),
            "fuente_x": xi.get("fuente", ""),
        })

    def cubeta_fila(nombre: Any, etiqueta: str, n_skus: int) -> dict:
        s_b = [(d, c) for d, c, *_ in serie.get(nombre, [])]
        rec = recuperado(s_b, inicio)
        return {"contenedor": etiqueta, "codigo": "", "skus": n_skus, "skus_multi": n_skus if nombre == MULTI else 0,
                "skus_con_venta": len({sk for *_, sk in serie.get(nombre, [])}),
                "piezas": sum(u for _, _, u, *_ in serie.get(nombre, [])),
                "x": None, "fecha_liberacion": None, "desde": inicio, "truncada": True,
                "recuperado": rec,
                "recuperado_archivo": recuperado([(d, c) for d, c, _, fu, _ in serie.get(nombre, []) if fu == "archivo"], inicio),
                "falta": None, "ritmo": ritmo_semanal(s_b, corte, PARAMS["semanas_ritmo"]),
                "eta": None, "semanas": None,
                "estado": "sin asignar: FIFO imposible" if nombre == MULTI else "fuera de costing.sku_contenedor",
                "stock": None, "potencial": None, "sin_hist": None, "alcanza": "no aplica",
                "multi_compartida": None, "fuente_x": ""}

    n_multi = sum(1 for v in atrib.values() if v == MULTI)
    filas_cont.append(cubeta_fila(MULTI, MULTI, n_multi))
    filas_cont.append(cubeta_fila(SIN_CONTENEDOR, SIN_CONTENEDOR,
                                  len({sk for *_, sk in serie.get(SIN_CONTENEDOR, [])})))

    # ── SKUs que frenan ──
    def etiqueta(s: str) -> tuple[Any, str]:
        a = atrib.get(s)
        if a == MULTI:
            return MULTI, ", ".join(str(n) for n in sorted(set(conts_de[s])))
        cod = codigos.get(a, {})
        return a, (max(cod, key=cod.get) if cod else "")

    frenan: list[dict] = []
    for s, i in sku_info.items():
        if i["total"] <= 0:
            continue
        lenta = i["cobertura"] is not None and i["cobertura"] > PARAMS["cobertura_frena_dias"]
        if i["clase"] != "exceso" and not lenta:
            continue
        if i["u90"] <= 0:
            motivo = "0 ventas en 90 d con stock"
        elif i["cobertura"] > PARAMS["horizonte_exceso_dias"]:
            motivo = f"cobertura {i['cobertura']:,.0f} d > {PARAMS['horizonte_exceso_dias']}"
        else:
            motivo = f"cobertura {i['cobertura']:,.0f} d > {PARAMS['cobertura_frena_dias']}"
        cont, cod = etiqueta(s)
        frenan.append({"contenedor": cont, "codigo": cod, "sku": s, "titulo": titulo.get(s) or "",
                       "motivo": motivo, **i,
                       "potencial": (i["total"] * i["contrib_media"]) if i["contrib_media"] is not None else None})
    frenan.sort(key=lambda r: (str(r["contenedor"]).zfill(6) if isinstance(r["contenedor"], int) else "~" + str(r["contenedor"]),
                               -(r["potencial"] if r["potencial"] is not None else -1e18), -r["total"]))

    # ── Palancas sin precio ──
    pubs: dict[tuple, dict] = {}
    for p in datos["publicaciones"]:
        k = (p["cuenta"], p["ml_item_id"])
        per = p["periodo"] or dt.date.min
        if k not in pubs or per > (pubs[k]["periodo"] or dt.date.min):
            pubs[k] = p
    palancas: list[dict] = []
    for s, i in sku_info.items():
        if i["total"] <= 0:
            continue
        cont, cod = etiqueta(s)
        base = {"contenedor": cont, "sku": s, "titulo": titulo.get(s) or "", "stock": i["total"],
                "ventas_dia": i["ventas_dia"], "clase": i["clase"],
                "cuenta": None, "item": None, "visitas": None, "unidades": None, "conversion": None, "precio": None}
        pr = presencia.get(s, set())
        faltan = [nom for key, nom in (("tiktok", "TikTok"), ("temu", "Temu"), ("walmart", "Walmart")) if key not in pr]
        if faltan:
            palancas.append({**base, "palanca": "sin publicación activa en otro canal",
                             "detalle": "faltan: " + ", ".join(faltan)})
        if i["ventas_dia"] >= PARAMS["rotacion_alta_dia"] and "ml_full" not in pr:
            palancas.append({**base, "palanca": "rotación alta sin Full",
                             "detalle": f"{i['ventas_dia']:.2f} piezas/día, sin publicación Full en ML"})
    for p in pubs.values():
        s = str(p["sku"] or "").upper()
        if s not in sku_info or sku_info[s]["total"] <= 0:
            continue
        vis, uni = int(p["visitas_30d"] or 0), int(p["unidades_30d"] or 0)
        if vis >= PARAMS["visitas_min"] and uni / vis < PARAMS["conversion_baja"]:
            cont, _ = etiqueta(s)
            i = sku_info[s]
            palancas.append({"contenedor": cont, "sku": s, "titulo": titulo.get(s) or "",
                             "stock": i["total"], "ventas_dia": i["ventas_dia"], "clase": i["clase"],
                             "palanca": "muchas visitas, poca conversión",
                             "detalle": f"{vis:,} visitas, {uni} u en 30 d ({uni / vis:.2%}); estado {p['estado']}",
                             "cuenta": p["cuenta"], "item": p["ml_item_id"], "visitas": vis,
                             "unidades": uni, "conversion": uni / vis, "precio": _f(p["precio"]) or None})
    orden_pal = {"muchas visitas, poca conversión": 0, "rotación alta sin Full": 1,
                 "sin publicación activa en otro canal": 2}
    palancas.sort(key=lambda r: (orden_pal[r["palanca"]],
                                 str(r["contenedor"]).zfill(6) if isinstance(r["contenedor"], int) else "~",
                                 -r["stock"]))

    # ── Cobertura de datos (para Notas) ──
    piezas_tot = sum(v["unidades"] for v in ventas)
    def pct(pred) -> float:
        return (sum(v["unidades"] for v in ventas if pred(v)) / piezas_tot) if piezas_tot else 0.0
    por_cubeta: dict[Any, list[dict]] = defaultdict(list)
    for v in ventas:
        por_cubeta[cubeta_de(v["sku"], atrib)].append(v)
    for f in filas_cont:
        vs = [v for v in por_cubeta.get(f["contenedor"], []) if v["fecha"] >= f["desde"]]
        pz = sum(v["unidades"] for v in vs)
        f["pct_envio_real"] = (sum(v["unidades"] for v in vs if v["envio_estado"] == "real") / pz) if pz else None
        f["pct_envio_promedio"] = (sum(v["unidades"] for v in vs if v["envio_estado"] == "promedio_real") / pz) if pz else None
        f["pct_comision_real"] = (sum(v["unidades"] for v in vs if v["comision_estado"] == "real") / pz) if pz else None
    envios_hasta = max((e["consultado"] for e in datos["envios"]), default=None)
    cobertura = {
        "piezas": piezas_tot,
        "pct_comision_real": pct(lambda v: v["comision_estado"] == "real"),
        "pct_envio_real": pct(lambda v: v["envio_estado"] == "real"),
        "pct_envio_promedio": pct(lambda v: v["envio_estado"] == "promedio_real"),
        "pct_envio_tabla": pct(lambda v: v["envio_estado"] == "tabla"),
        "pct_envio_peso_dudoso": pct(lambda v: v["envio_estado"] == "tabla_peso_dudoso"),
        "pct_envio_sin_medidas": pct(lambda v: v["envio_estado"] == "tabla_sin_medidas"),
        "pct_envio_sin_dato": pct(lambda v: v["envio_estado"] == "sin_dato"),
        # Piezas cuyo envío pesa más que su precio sin IVA: casi nunca pasa con el
        # real (~1 %); si en la tabla es mucho más, hay pesos malos sin atrapar.
        "envio_mayor_precio": {g: (sum(v["unidades"] for v in ventas if pred(v) and v["ingreso"] > 0
                                       and (v["envio"] or 0) > v["ingreso"] / (1 + PARAMS["iva"]))
                                   / max(sum(v["unidades"] for v in ventas if pred(v)), 1))
                               for g, pred in (("real", lambda v: v["envio_estado"] == "real"),
                                               ("tabla", lambda v: v["envio_estado"].startswith("tabla")))},
        "negativas": (sum(1 for v in ventas if v["contribucion"] < 0),
                      sum(v["contribucion"] for v in ventas if v["contribucion"] < 0)),
        "pct_archivo": pct(lambda v: v["fuente"] == "archivo"),
        "estados": {e: sum(v["unidades"] for v in ventas if estado_contribucion(v["comision_estado"], v["envio_estado"]) == e)
                    for e in ("real", "parcial", "estimado")},
        "sin_sku": datos["vivo_sin_sku"][0] if datos["vivo_sin_sku"] else {"lineas": 0, "piezas": 0},
        "envios_consultados_hasta": envios_hasta,
        "por_canal": {c: sum(v["contribucion"] for v in ventas if v["canal"] == c)
                      for c in sorted({v["canal"] for v in ventas})},
    }
    return {"contenedores": filas_cont, "frenan": frenan, "palancas": palancas,
            "corte": corte, "inicio": inicio, "cobertura": cobertura, "generado": generado,
            "n_skus": len(conts_de), "n_multi": n_multi}


# ══════════════════════════════════════════════════════════════════════════════
# Excel
# ══════════════════════════════════════════════════════════════════════════════

_MONEDA = '"$"#,##0;[Red]-"$"#,##0'
_FECHA = "yyyy-mm-dd"


def _hoja(wb, titulo: str, columnas: list[tuple[str, str, str | None, int]], filas: list[dict]):
    """columnas = [(encabezado, clave, formato, ancho)]. None → «sin dato»
    solo donde el encabezado lo pide (claves con «?» al final)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(titulo)
    ws.append([c[0] for c in columnas])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F2937")
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    for f in filas:
        fila = []
        for _, clave, _, _ in columnas:
            sin = clave.endswith("?")
            v = f.get(clave.rstrip("?"))
            fila.append(SIN_DATO if (v is None and sin) else v)
        ws.append(fila)
    for j, (_, _, fmt, ancho) in enumerate(columnas, start=1):
        letra = get_column_letter(j)
        ws.column_dimensions[letra].width = ancho
        if fmt:
            for celda in ws[letra][1:]:
                if isinstance(celda.value, (int, float)):
                    celda.number_format = fmt
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 45
    return ws


def escribir_excel(r: dict, ruta: Path, x_csv: Path, plantilla_creada: bool) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for f in r["contenedores"]:
        f["desde_txt"] = f"{f['desde']:%Y-%m-%d}" + (" (inicio de los datos)" if f["truncada"] else " (liberación)")
        f["semanas_txt"] = round(f["semanas"], 1) if f["semanas"] is not None else None
    _hoja(wb, "Contenedores", [
        ("Contenedor", "contenedor", None, 14), ("Código", "codigo", None, 15),
        ("SKUs", "skus", "#,##0", 7), ("SKUs multi", "skus_multi", "#,##0", 8),
        ("SKUs con venta", "skus_con_venta", "#,##0", 8), ("Piezas vendidas", "piezas", "#,##0", 9),
        ("X (MXN)", "x?", _MONEDA, 13), ("Fecha de liberación", "fecha_liberacion?", _FECHA, 12),
        ("Recuperado desde", "desde_txt", None, 26),
        ("Recuperado (cota superior)", "recuperado", _MONEDA, 14),
        ("de ello: archivo ≤15-jul", "recuperado_archivo", _MONEDA, 13),
        ("Falta", "falta?", _MONEDA, 13), ("Ritmo semanal (4 sem)", "ritmo", _MONEDA, 12),
        ("ETA", "eta?", _FECHA, 12), ("Semanas a ETA", "semanas_txt?", "0.0", 9),
        ("Estado", "estado", None, 16),
        ("Stock restante (pzs, SKUs de un solo contenedor)", "stock", "#,##0", 12),
        ("Potencial del stock (cota superior)", "potencial", _MONEDA, 14),
        ("SKUs con stock sin historial (no suman al potencial)", "sin_hist", "#,##0", 12),
        ("¿Alcanza con el stock?", "alcanza", None, 12),
        ("Contribución de sus SKUs multi (compartida, NO sumada)", "multi_compartida", _MONEDA, 15),
        ("% piezas con envío real", "pct_envio_real?", "0%", 9),
        ("% piezas con envío = promedio real del SKU", "pct_envio_promedio?", "0%", 10),
        ("% piezas con comisión real", "pct_comision_real?", "0%", 9),
        ("Fuente de X", "fuente_x", None, 14),
    ], r["contenedores"])
    for f in r["frenan"]:
        f["cobertura_txt"] = round(f["cobertura"]) if f["cobertura"] is not None else "sin ventas"
    _hoja(wb, "SKUs que frenan", [
        ("Contenedor", "contenedor", None, 14), ("Código(s)", "codigo", None, 15),
        ("SKU", "sku", None, 22), ("Título", "titulo", None, 40), ("Clase", "clase", None, 9),
        ("Motivo", "motivo", None, 26), ("Stock propio (Woo)", "propio", "#,##0", 9),
        ("Full ML", "full", "#,##0", 8), ("FBA", "fba", "#,##0", 7), ("Stock total", "total", "#,##0", 9),
        ("Unidades 90 d", "u90", "#,##0", 9), ("Ventas/día", "ventas_dia", "0.00", 9),
        ("Cobertura (días)", "cobertura_txt", "#,##0", 10),
        ("Contribución media/pieza (cota sup.)", "contrib_media?", _MONEDA, 13),
        ("Valor detenido (stock × contribución)", "potencial?", _MONEDA, 14),
        ("Última venta", "ultima?", _FECHA, 12),
    ], r["frenan"])
    _hoja(wb, "Palancas sin precio", [
        ("Contenedor", "contenedor", None, 14), ("SKU", "sku", None, 22), ("Título", "titulo", None, 40),
        ("Palanca", "palanca", None, 30), ("Detalle", "detalle", None, 44), ("Clase", "clase", None, 9),
        ("Stock total", "stock", "#,##0", 9), ("Ventas/día", "ventas_dia", "0.00", 9),
        ("Cuenta ML", "cuenta", None, 14), ("Item ML", "item", None, 15),
        ("Visitas 30 d", "visitas", "#,##0", 9), ("Unidades 30 d", "unidades", "#,##0", 9),
        ("Conversión", "conversion", "0.00%", 9), ("Precio ML", "precio", _MONEDA, 10),
    ], r["palancas"])

    c = r["cobertura"]
    ws = wb.create_sheet("Notas")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 110
    lineas: list[tuple[str, Any]] = [
        ("RADAR DE PRECIOS · F1 · ESTADO POR CONTENEDOR", ""),
        ("Generado (UTC)", f"{r['generado']:%Y-%m-%d %H:%M}"),
        ("Ambiente", "SANDBOX (yvootpbz), solo lectura. NO es producción: el sandbox es un clon con fecha."),
        ("Corte (último día con ventas vivas)", f"{r['corte']:%Y-%m-%d}"),
        ("Inicio de los datos", f"{r['inicio']:%Y-%m-%d} (archivo dailytrack hasta el 15-jul; order_items desde el 16-jul)"),
        ("Archivo de X", f"{x_csv} " + ("(plantilla CREADA en esta corrida: X vacío)" if plantilla_creada else "(leído, no modificado)")),
        ("", ""),
        ("QUÉ ES COTA SUPERIOR", "Toda contribución: no resta Full (almacenaje), publicidad ni devoluciones (sin dato). "
                                 "Los canales distintos de ML tampoco restan envío (sin dato)."),
        ("QUÉ ES COTA INFERIOR", "El recuperado de un contenedor liberado antes del inicio de los datos: lo vendido antes no se ve "
                                 "(columna «Recuperado desde» dice «inicio de los datos»)."),
        ("SKU en varios contenedores", "Cubeta «multi: sin asignar». No se reparte: no hay piezas por contenedor ni fechas confiables "
                                       "para un FIFO. Su contribución aparece aparte por contenedor como «compartida, NO sumada»."),
        ("SKU sin contenedor", "Ventas de SKUs que no están en costing.sku_contenedor: su propia cubeta, fuera de todo X."),
        ("Ventas que entran", "El filtro de channel.sales_daily: fuera cancelled/invalid/canceled sin importar la caja; "
                              "partially_refunded se queda (decisión abierta, igual que allá)."),
        ("Comisión", "La real de channel.order_items (sale_fee de ML: % sobre el precio con IVA; se resta tal cual). "
                     "Comisión 0 = aún no llega o el canal no la reporta → se estima con la tasa real media de la "
                     "publicación, o con comision_estimada. Amazon/TikTok/Temu: siempre estimada."),
        ("Envío (ML)", "1) El real de enrich.order_shipping_cost (por pedido, repartido por importe). 2) Si ese pedido no "
                       "lo tiene: el promedio real por pieza del mismo SKU (≥ min_piezas_promedio_envio piezas medidas). "
                       "3) Si no: la tabla oficial de services/costos.py con el peso efectivo de la pieza "
                       "(costos_validados), un envío por línea; con densidad > densidad_max_kg_l el peso es de caja "
                       "master y se usa solo el volumétrico; sin medidas, 0.5 kg por defecto del motor."),
        ("Stock", "Woo (canal general) manda; si no hay fila de Woo, el máximo del espejo (nunca la suma). "
                  "Más Full de ML y FBA de Amazon. La presencia en canales hereda la del padre de Woo."),
        ("Clase", "exceso = stock con 0 ventas en 90 d, o cobertura > horizonte_exceso_dias; recompra = lista del "
                  "negocio (vacía); si no, normal. «Frena» = exceso o cobertura > cobertura_frena_dias."),
        ("Palancas", "Sin publicación activa en TikTok (ACTIVATE) / Temu (" + ", ".join(sorted(TEMU_VENDIBLES)) + ") / "
                     f"Walmart ({WALMART_VIVO}); rotación ≥ rotacion_alta_dia sin publicación Full en ML; ≥ visitas_min "
                     "visitas en 30 d con conversión < conversion_baja (última foto de market_publicaciones_v)."),
        ("", ""),
        ("PARÁMETROS (decisión, no medición)", ""),
        *[(f"  {k}", v) for k, v in PARAMS.items()],
        ("  RECOMPRA", "vacía (lo decide el negocio)"),
        ("", ""),
        ("COBERTURA DE LOS DATOS", ""),
        ("  Piezas vendidas en la serie", c["piezas"]),
        ("  % piezas con comisión real", f"{c['pct_comision_real']:.1%}"),
        ("  % piezas con envío real", f"{c['pct_envio_real']:.1%} (order_shipping_cost consultado hasta "
                                      f"{c['envios_consultados_hasta'] or 'nunca'}; después, estimado)"),
        ("  % piezas con envío = promedio real del SKU", f"{c['pct_envio_promedio']:.1%}"),
        ("  % piezas con envío de tabla", f"{c['pct_envio_tabla']:.1%}"),
        ("  % piezas con envío de tabla, peso de caja master (volumétrico)", f"{c['pct_envio_peso_dudoso']:.1%}"),
        ("  % piezas con envío de tabla sin medidas (0.5 kg)", f"{c['pct_envio_sin_medidas']:.1%}"),
        ("  % piezas sin envío (canal no ML)", f"{c['pct_envio_sin_dato']:.1%}"),
        ("  % piezas del archivo (≤15-jul)", f"{c['pct_archivo']:.1%}"),
        ("  Piezas con envío > precio sin IVA", f"real {c['envio_mayor_precio']['real']:.1%} · tabla "
                                                f"{c['envio_mayor_precio']['tabla']:.1%} (la diferencia son pesos "
                                                "malos que la regla de densidad no atrapa: ahí la contribución "
                                                "puede quedar CORTA, no larga)"),
        ("  Ventas con contribución negativa", f"{c['negativas'][0]:,} por ${c['negativas'][1]:,.0f} (se suman tal cual)"),
        ("  Piezas por estado de contribución", ", ".join(f"{k} {v:,}" for k, v in c["estados"].items())),
        ("  Líneas vivas sin SKU (fuera)", f"{c['sin_sku']['lineas']:,} líneas / {c['sin_sku']['piezas']:,} piezas"),
        ("  Contribución por canal", ", ".join(f"{k} ${v:,.0f}" for k, v in c["por_canal"].items())),
    ]
    for a, b in lineas:
        ws.append([a, b])
        if a.strip() and a == a.upper():
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    wb.save(ruta)


# ══════════════════════════════════════════════════════════════════════════════

def leer_env(ruta: Path) -> dict[str, str]:
    vals: dict[str, str] = {}
    if not ruta.exists():
        return vals
    for linea in ruta.read_text(encoding="utf-8", errors="ignore").splitlines():
        s = linea.strip()
        if s and not s.startswith("#") and "=" in s:
            k, _, v = s.partition("=")
            vals[k.strip()] = v.strip().strip('"').strip("'")
    return vals


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--env", type=Path, default=ROOT / "env.staging")
    ap.add_argument("--salida", type=Path, default=ROOT.parent / "_radar_f0")
    ap.add_argument("--x-csv", type=Path, default=None)
    a = ap.parse_args(argv)

    salida = a.salida.resolve()
    x_csv = (a.x_csv or salida / "x_contenedores.csv").resolve()
    for ruta in (salida, x_csv):
        if dentro_de(ruta, ROOT):
            raise SystemExit(f"ABORT: {ruta} está dentro del repo (público). X y el Excel van FUERA.")
    salida.mkdir(parents=True, exist_ok=True)

    dsn = validar_dsn(leer_env(a.env).get("SUPABASE_DB_URL", ""))
    print(f"Leyendo el sandbox ({REF_SANDBOX}), solo lectura…")
    datos = leer_sandbox(dsn)

    numeros = sorted({int(f["numero"]) for f in datos["sku_contenedor"]})
    creada = crear_plantilla_x(x_csv, numeros)
    x_por_cont = leer_x(x_csv)

    r = armar(datos, x_por_cont, dt.datetime.now(dt.timezone.utc))
    ruta = salida / f"estado_contenedores_{dt.date.today():%Y-%m-%d}.xlsx"
    escribir_excel(r, ruta, x_csv, creada)

    conts = [f for f in r["contenedores"] if isinstance(f["contenedor"], int)]
    multi = next(f for f in r["contenedores"] if f["contenedor"] == MULTI)
    sinc = next(f for f in r["contenedores"] if f["contenedor"] == SIN_CONTENEDOR)
    print(f"Plantilla de X: {'CREADA' if creada else 'ya existía (no se tocó)'} → {x_csv}")
    print(f"Corte {r['corte']:%Y-%m-%d} · datos desde {r['inicio']:%Y-%m-%d}")
    print(f"Contenedores: {len(conts)} · SKUs ubicados: {r['n_skus']:,} (multi: {r['n_multi']})")
    print(f"Recuperado en contenedores (cota superior): ${sum(f['recuperado'] for f in conts):,.0f}")
    print(f"  multi sin asignar: ${multi['recuperado']:,.0f} · sin contenedor: ${sinc['recuperado']:,.0f}")
    print(f"Sin X: {sum(1 for f in conts if f['x'] is None)} de {len(conts)}")
    print(f"SKUs que frenan: {len(r['frenan']):,} · palancas: {len(r['palancas']):,}")
    print(f"Excel: {ruta}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
