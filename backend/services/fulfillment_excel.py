"""
fulfillment_excel.py — La planeación semanal en EXCEL, con la forma de salida del
prompt estándar: A) una hoja por tienda, B) totales en el resumen, C) ganadores
agotados con su reemplazo, D) la lista de SKUs separada por coma, E) alertas.

Lo arma el backend con openpyxl (Brandon, 24-sep: "descargar la lista sin
problemas"): un .xlsx abre igual en cualquier Excel, sin el lío de acentos y
separadores de un CSV. Recibe la planeación TAL COMO la ve la persona —con sus
cambios— y no lee ni escribe nada más.
"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_INDIGO = PatternFill("solid", fgColor="4F46E5")
_BLANCO = Font(color="FFFFFF", bold=True)
_NEGRITA = Font(bold=True)

COLUMNAS = [
    ("sku", "SKU", 20), ("nombre", "Nombre (Omnicanal)", 44), ("destino", "Destino", 12),
    ("precio", "Precio (MXN)", 12),
    ("vv", "Vendió (ventana)", 12), ("v7", "Vendió 7 d", 10), ("stock", "En almacén hoy", 13),
    ("en_camino", "En camino", 10), ("borrador", "En borradores", 12), ("libre", "Libre Odoo", 11),
    ("pidio", "Pidió (faltante)", 13), ("bodega_puede", "Bodega puede", 12), ("propuesta", "Propuesta", 11),
    ("a_mandar", "A mandar", 10), ("estado", "Estado", 12), ("caja", "Piezas por caja", 12),
    ("listing_id", "Publicación", 16), ("reemplazo_de", "Reemplazo de", 16),
]


def _encabezado(ws, titulos: list[tuple[str, int]]) -> None:
    for i, (t, ancho) in enumerate(titulos, start=1):
        c = ws.cell(row=1, column=i, value=t)
        c.fill, c.font = _INDIGO, _BLANCO
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"


def armar(plan: dict[str, Any]) -> bytes:
    wb = Workbook()
    res = wb.active
    res.title = "Resumen"
    res.column_dimensions["A"].width = 34
    for col in "BCDEFG":
        res.column_dimensions[col].width = 16
    filas_res: list[list[Any]] = [
        ["Planeación semanal de FULL — Kubera / Omnicanal"],
        ["Semana", plan.get("semana") or ""],
        ["Ventana de ventas", plan.get("ventana") or ""],
        ["Parámetros", plan.get("parametros_texto") or ""],
        ["Datos", plan.get("en_vivo") or ""],
        [],
        ["Tienda", "Renglones", "Piezas pedidas", "Piezas propuestas", "A mandar", "Tasa de validado",
         "% final vs pedido"],
    ]
    for t in plan.get("tiendas") or []:
        tot = t.get("totales") or {}
        filas_res.append([t.get("nombre"), tot.get("renglones"), tot.get("pedidas"), tot.get("propuestas"),
                          tot.get("a_mandar"), tot.get("tasa_validado"), tot.get("final_vs_pedido")])
    for fila in filas_res:
        res.append(fila)
    res["A1"].font = Font(bold=True, size=13)
    for c in res[7]:
        c.fill, c.font = _INDIGO, _BLANCO

    for t in plan.get("tiendas") or []:
        ws = wb.create_sheet((t.get("nombre") or t.get("tienda") or "Tienda")[:31])
        _encabezado(ws, [(titulo, ancho) for _, titulo, ancho in COLUMNAS])
        for r in t.get("renglones") or []:
            ws.append([r.get(k) for k, _, _ in COLUMNAS])

    g = wb.create_sheet("Ganadores agotados")
    _encabezado(g, [("Tienda", 16), ("SKU agotado", 20), ("Nombre", 40), ("Vendió", 10),
                    ("Reemplazo", 20), ("Nombre del reemplazo", 40), ("Match", 18), ("Libre", 10),
                    ("Precio del reemplazo", 14)])
    for x in plan.get("ganadores") or []:
        cand = x.get("candidatos") or [{}]
        for c in cand:
            g.append([x.get("tienda"), x.get("sku"), x.get("nombre"), x.get("vv"),
                      c.get("sku"), c.get("nombre"), c.get("tipo"), c.get("libre"), c.get("precio")])

    s = wb.create_sheet("SKUs")
    _encabezado(s, [("Tienda", 16), ("SKUs a enviar, separados por coma", 120)])
    for t in plan.get("tiendas") or []:
        skus = [r["sku"] for r in t.get("renglones") or [] if (r.get("a_mandar") or 0) > 0]
        s.append([t.get("nombre"), ", ".join(skus)])
        s.cell(row=s.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    a = wb.create_sheet("Alertas")
    _encabezado(a, [("Tienda", 16), ("SKU", 20), ("Alerta", 22), ("Detalle", 90)])
    for x in plan.get("alertas") or []:
        a.append([x.get("tienda"), x.get("sku"), x.get("tipo"), x.get("detalle")])

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
