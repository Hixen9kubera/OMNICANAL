"""
checklist.py — INVENTARIO · Checklist: la validación de ALMACÉN, SKU por SKU.

Brandon, 24-sep-2026: *«crear una tab en inventario llamado CHECKLIST,
únicamente para almacén, ya sea mediante CSV, Excel o MCP; seleccionar que de
las opcionales se puedan poner en obligatorios y crear matriz. Debemos pedir
cajas con piezas. Esta tab es la que va a determinar si cada SKU tiene sus
atributos correctos que Mercado Libre exige, además de pedir dimensiones, cajas
y piezas que almacén tiene. Cada semana se seleccionan ~100 SKUs; al
seleccionarlos se podrá descargar un Excel con las características necesarias
y, una vez llenado, cargarlo con los campos llenos».*

QUÉ CONTESTA
------------
Por cada SKU del LOTE de la semana: ¿está completo para Mercado Libre y para
almacén? «Completo» es la suma de dos listas:

  1. ATRIBUTOS DE ML que se EXIGEN: los obligatorios del propio ML
     (`required` / `catalog_required` de su API pública) MÁS los opcionales que
     el equipo subió a obligatorios en la MATRIZ (filas `fuente = 'manual'`
     de `channel.field_requirements`).
  2. LO QUE ALMACÉN MIDE Y CUENTA: largo, ancho, alto y peso del producto
     EMPACADO, cajas y piezas por caja (columnas `almacen_*` de `core.products`).

DE DÓNDE SALE CADA COSA — solo kubera y la API pública de ML (regla de
Brandon del 17-sep: nada de WordPress)
  · Categoría ML del SKU ... `specs._categorias` (channel.product_category con
                             la precedencia panel > real > predictor, y si no
                             hay, la de la publicación viva).
  · Qué pide la categoría .. `specs_editor._campos_ml` (API pública, 1 h de caché).
  · Lo ya capturado ....... `enrich.channel_content` (cuenta ''), el MISMO sitio
                             que el Publicador y el editor de specs del cajón.
  · Nombre de la categoría  `channel.categories`.
  · Título del SKU ........ `core.products.name`.

DÓNDE SE GUARDA LO QUE SUBE ALMACÉN
  · Atributos → `specs_editor.guardar_sync`, que fusiona por campo y escribe
    la lista COMPLETA (los tres cuidados de su cabecera).
  · Medidas, cajas, piezas → columnas `almacen_*` de `core.products`
    (migración 0058). Es la caja «Bodega» del cotejo del Catálogo Maestro, que
    hasta hoy era NULL porque nadie la medía.
  · La MATRIZ → `channel.field_requirements` con fuente = 'manual': la misma
    lista de requisitos por categoría que ya existe, sin tabla aparte.

NORMALIZACIÓN (Brandon, 24-sep: «buscarlo en la base de datos para no repetir»)
Las columnas largo/ancho/alto/peso/cajas/piezas_por_caja de
`costing.costos_validados` se REVISARON y no son el mismo dato:
  · sus medidas son el volumen de FLETE reconstruido del CBM (solo el 26% cuadra
    con el flete; ACC-0313-NEG dice 60×51×51 cm y 0.281 kg = la caja máster), y
    costos.py las usa para flete y comisión de envío: escribir ahí la medida real
    cambiaría precios en el siguiente recálculo;
  · sus cajas y piezas por caja son las del PACKING LIST (las escribe el
    Resolver), la otra mitad de la comparación que pidió Brandon el 8-sep.
Por eso lo de almacén va en columnas propias, en la fila del producto, y lo de
costos_validados se enseña al lado como REFERENCIA («en sistema»). Y no en
costos_validados: Costos decide «sin costo» por la existencia de la fila, y
crearle fila a un SKU sin costo lo sacaría de la lista de trabajo de los KAM.

LA SEMANA es la ISO, la que el equipo llama «Week 39». La lista semanal llega
como Excel (una hoja «Week NN» con columna SKU, con o sin corchetes) y se carga
tal cual con `lista_sync`.

DOS DECISIONES QUE CONVIENE SABER
  · UNA CELDA VACÍA NO BORRA NADA. El Excel sale pre-llenado con lo que ya hay;
    si almacén deja algo en blanco se entiende «no lo sé», no «quítalo». Para
    borrar un atributo está el editor del cajón en el Catálogo Maestro.
  · Las medidas NO se escriben como atributos SELLER_PACKAGE_* de ML. Esos
    campos deciden la TARIFA DE ENVÍO que cobra ML: mandarlos es cambiar
    dinero en publicaciones vivas, y eso espera su dale (regla 3).

SI LA MIGRACIÓN 0058 NO ESTÁ APLICADA, la pestaña lo dice en vez de tronar:
todo lo que toca la tabla o las columnas nuevas levanta `FaltaMigracion` y el
router lo devuelve como `falta_migracion: true`.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import logging
import re
from concurrent.futures import ThreadPoolExecutor
from typing import Any, Iterable
from zoneinfo import ZoneInfo

import psycopg2

from core import actor
from services import specs, specs_editor
from services import supabase_db as sdb

log = logging.getLogger("omnicanal.services.checklist")

CANAL = "mercado_libre"
_ZONA = ZoneInfo("America/Mexico_City")

# Lo que almacén mide y cuenta. (clave, etiqueta, tipo). La clave va en
# minúsculas a propósito: los ids de ML son MAYÚSCULAS, así que en el Excel y
# en el CSV las dos familias nunca chocan.
LOGISTICA: tuple[tuple[str, str, str], ...] = (
    ("largo_cm", "Largo (cm)", "decimal"),
    ("ancho_cm", "Ancho (cm)", "decimal"),
    ("alto_cm", "Alto (cm)", "decimal"),
    ("peso_kg", "Peso (kg)", "decimal"),
    ("cajas", "Cajas", "entero"),
    ("piezas_por_caja", "Piezas por caja", "entero"),
)
_LOG_CLAVES = tuple(k for k, _, _ in LOGISTICA)
_LOG_ETIQUETA = {k: e for k, e, _ in LOGISTICA}
_LOG_TIPO = {k: t for k, _, t in LOGISTICA}

# El orden en que se enseñan y se exigen los atributos.
NIVELES = ("ml", "matriz", "auto", "principal", "secundario")
_EXIGIDOS = frozenset({"ml", "matriz"})

_MAX_SKUS_LOTE = 500      # por llamada; la semana típica son ~100
_HILOS_ML = 8             # categorías pedidas a la API de ML en paralelo
_HILOS_GUARDAR = 4        # SKUs guardados en paralelo (el pool de kubera es chico)


class FaltaMigracion(RuntimeError):
    """`ops.checklist_lote` o las columnas `almacen_*` no existen (migración 0058)."""


_MSG_MIGRACION = ("Falta aplicar la migración 0058 (ops.checklist_lote y las "
                  "columnas almacen_* de core.products) en kubera.")


def _sin_tabla(exc: Exception) -> bool:
    """42P01 = tabla que no existe; 42703 = columna que no existe.

    El CÓDIGO manda. El texto solo se mira si el error no trae código: un
    `savepoint "…" does not exist` (3B001) también dice «does not exist», y
    confundirlo con la falta de la 0058 escondía una conexión muerta."""
    codigo = getattr(exc, "pgcode", None)
    if codigo:
        return codigo in ("42P01", "42703")
    return ("does not exist" in str(exc)
            and ("checklist_" in str(exc) or "almacen_" in str(exc)))


def _q(sql: str, params: Any = None) -> list[dict[str, Any]]:
    try:
        return sdb.fetch_all(sql, params)
    except Exception as exc:  # noqa: BLE001
        if _sin_tabla(exc):
            raise FaltaMigracion(_MSG_MIGRACION) from exc
        raise


# ══════════════════════════════════════════════════════════════════════════════
# La semana
# ══════════════════════════════════════════════════════════════════════════════

def hoy() -> dt.date:
    return dt.datetime.now(_ZONA).date()


def lunes(fecha: dt.date | None = None) -> dt.date:
    f = fecha or hoy()
    return f - dt.timedelta(days=f.weekday())


def etiqueta(semana: dt.date) -> str:
    """El lunes → «Week 39», como el equipo nombra sus hojas."""
    return f"Week {semana.isocalendar()[1]}"


def lunes_iso(anio: int, numero: int) -> dt.date | None:
    try:
        return dt.date.fromisocalendar(anio, numero, 1)
    except ValueError:
        return None


def semana_de(texto: str | None) -> dt.date:
    """'2026-09-24' → el lunes de esa semana. Sin nada, la semana de hoy."""
    if not texto:
        return lunes()
    try:
        return lunes(dt.date.fromisoformat(texto.strip()[:10]))
    except ValueError:
        return lunes()


def _sin_comillas_envolventes(p: str) -> str:
    """Quita SOLO las comillas que ENVUELVEN el texto. Las de adentro o las del
    final son parte del SKU: hay 12 así en core.products (HERR-0032-ROJ-16",
    ACC-0665-NEG-7"…), y quitarlas los volvía «desconocidos». Excel, al copiar
    una celda con comillas, la envuelve y DUPLICA las de adentro
    ("HERR-0032-ROJ-16""): eso también se deshace aquí."""
    for q in ('"', "'"):
        if len(p) >= 2 and p[0] == q and p[-1] == q:
            return p[1:-1].replace(q + q, q)
    return p


def limpiar_skus(crudo: Iterable[str] | str | None) -> list[str]:
    """Acepta lista o texto pegado de Excel (renglones, comas, tabs, espacios).
    Quita duplicados conservando el orden."""
    if crudo is None:
        return []
    if isinstance(crudo, str):
        # Espacios, renglones, tabs y ; siempre separan. La COMA también, salvo
        # ENTRE DOS DÍGITOS: hay SKUs con coma decimal (TEC-1660-NEG-SONIC-1,6L),
        # y en «TEC-0370-NEG,ORG-0863-ROS» la coma sí separa. (Decidirlo por si
        # el texto trae renglones rompía «A-1,B-2» con un salto al final.)
        partes = re.split(r"[\s;]+|(?<!\d),|,(?!\d)", crudo)
    else:
        # Una LISTA ya viene separada (celdas del Excel, la hoja semanal, la
        # selección de la pantalla): cada elemento es UN SKU y no se parte.
        partes = [str(x or "").strip() for x in crudo]
    salida: list[str] = []
    vistos: set[str] = set()
    for p in partes:
        p = _sin_comillas_envolventes(p.strip())
        # Ningún SKU EMPIEZA con comilla (medido: 0 en core.products). La que
        # queda al inicio es de una celda multilínea que Excel envolvió; la
        # del final puede ser parte del SKU y la decide _canonicos.
        p = p.lstrip('"\'').replace('""', '"')
        if not p.strip('"\''):
            continue
        if p and p.upper() not in vistos:
            vistos.add(p.upper())
            salida.append(p)
    return salida


# ══════════════════════════════════════════════════════════════════════════════
# Lecturas de kubera (sin las tablas nuevas)
# ══════════════════════════════════════════════════════════════════════════════

def _canonicos(skus: list[str]) -> tuple[dict[str, str], list[str]]:
    """Empata lo que se pegó con el SKU real (sin importar mayúsculas).

    Vale el SKU que esté en `core.products`: es la llave foránea del lote y la
    fila donde se guarda lo de almacén. Devuelve ({pegado_mayus: real},
    desconocidos).
    """
    if not skus:
        return {}, []
    mayus = [s.upper() for s in skus]
    reales: dict[str, str] = {}
    # Con y sin la comilla FINAL: 'HERR-0032-ROJ-16"' existe así, y
    # 'ABC-0002-NEG"' es un SKU normal con la comilla de la celda de Excel.
    sin_final = {s.upper().rstrip('"'): s.upper() for s in skus if s.endswith('"')}
    for r in sdb.fetch_all("select sku::text as sku from core.products "
                           "where upper(sku::text) = any(%s)",
                           (mayus + list(sin_final),)):
        reales.setdefault(r["sku"].upper(), r["sku"])
    for corto, pegado in sin_final.items():
        if pegado not in reales and corto in reales:
            reales[pegado] = reales[corto]
    desconocidos = [s for s in skus if s.upper() not in reales]
    return reales, desconocidos


def _titulos(skus: list[str]) -> dict[str, str]:
    if not skus:
        return {}
    try:
        return {r["sku"]: r["name"] for r in sdb.fetch_all(
            "select sku::text as sku, name from core.products "
            "where sku = any(%s::citext[])", (skus,))}
    except Exception as exc:  # noqa: BLE001
        log.warning("checklist: títulos no disponibles: %s", exc)
        return {}


def _urls_ml(skus: list[str]) -> dict[str, str]:
    """Una publicación viva de ML por SKU, para que almacén vea el producto."""
    if not skus:
        return {}
    salida: dict[str, str] = {}
    try:
        for r in sdb.fetch_all(
                "select sku, url from channel.listings "
                "where canal = %s and sku = any(%s) and url is not null "
                "  and coalesce(status, '') <> 'error' "
                # En channel.listings de ML el estado vivo se llama
                # 'published' (medido el 24-sep: published 3,580, NULL 1,771,
                # error 275; 'active' no existe).
                # coalesce: en DESC Postgres pone los NULL PRIMERO, y hay
                # 1,751 filas de ML con status NULL y url.
                "order by coalesce(status = 'published', false) desc, updated_at desc",
                (CANAL, skus)):
            salida.setdefault(r["sku"], r["url"])
    except Exception as exc:  # noqa: BLE001
        log.warning("checklist: publicaciones no disponibles: %s", exc)
    return salida


def _nombres_categoria(cats: Iterable[str]) -> dict[str, dict[str, str]]:
    cats = sorted({c for c in cats if c})
    if not cats:
        return {}
    try:
        return {r["category_id"]: {"nombre": r["name"], "ruta": r["path"]}
                for r in sdb.fetch_all(
                    "select category_id, name, path from channel.categories "
                    "where channel_id = %s and category_id = any(%s)", (CANAL, cats))}
    except Exception as exc:  # noqa: BLE001
        log.warning("checklist: nombres de categoría no disponibles: %s", exc)
        return {}


def _valores_ml(skus: list[str]) -> dict[str, dict[str, str]]:
    """Lo ya capturado de ML por SKU: {sku: {ATRIBUTO: valor}}. Una consulta."""
    if not skus:
        return {}
    salida: dict[str, dict[str, str]] = {}
    for r in sdb.fetch_all(
            "select sku, contenido->'atributos' as atributos "
            "from enrich.channel_content "
            "where canal = %s and cuenta = '' and sku = any(%s)", (CANAL, skus)):
        attrs = r.get("atributos")
        if not isinstance(attrs, list):
            continue
        salida[r["sku"]] = {
            specs_editor._clave(a): specs_editor._valor(a)
            for a in attrs if isinstance(a, dict) and specs_editor._clave(a)}
    return salida


def _campos_por_categoria(cats: Iterable[str]) -> dict[str, list[dict[str, Any]]]:
    """Los atributos de cada categoría, pedidos a ML en paralelo (con caché)."""
    cats = sorted({c for c in cats if c})
    if not cats:
        return {}
    with ThreadPoolExecutor(max_workers=min(_HILOS_ML, len(cats))) as ex:
        return dict(zip(cats, ex.map(specs_editor._campos_ml, cats)))


# ══════════════════════════════════════════════════════════════════════════════
# Lo de almacén (core.products.almacen_*, 0058) y la matriz (field_requirements)
# ══════════════════════════════════════════════════════════════════════════════

# La clave de la pantalla/Excel → la columna de core.products.
_COLUMNA = {
    "largo_cm": "almacen_largo_cm", "ancho_cm": "almacen_ancho_cm",
    "alto_cm": "almacen_alto_cm", "peso_kg": "almacen_peso_kg",
    "cajas": "almacen_cajas", "piezas_por_caja": "almacen_piezas_por_caja",
}


def _num_o_none(v: Any, entero: bool = False) -> Any:
    if v is None:
        return None
    return int(v) if entero else float(v)


def _capturado(skus: list[str]) -> dict[str, dict[str, Any]]:
    """Lo que almacén midió y contó, de core.products. Solo lo capturado en el
    Checklist: si nadie lo ha medido, todo viene en None."""
    if not skus:
        return {}
    cols = ", ".join(_COLUMNA.values())
    salida: dict[str, dict[str, Any]] = {}
    for r in _q(f"select sku::text as sku, {cols}, almacen_por, almacen_en "
                f"from core.products where sku = any(%s::citext[])", (skus,)):
        d = {k: _num_o_none(r.get(c), _LOG_TIPO[k] == "entero")
             for k, c in _COLUMNA.items()}
        d["capturado_por"] = r.get("almacen_por")
        d["capturado_en"] = r["almacen_en"].isoformat() if r.get("almacen_en") else None
        salida[r["sku"]] = d
    return salida


def _sistema(skus: list[str]) -> dict[str, dict[str, Any]]:
    """Lo que YA dice el sistema, como REFERENCIA: costos_validados. No son
    medidas de almacén (ver la cabecera); se enseñan al lado para comparar."""
    if not skus:
        return {}
    salida: dict[str, dict[str, Any]] = {}
    try:
        for r in sdb.fetch_all(
                "select sku::text as sku, largo, ancho, alto, peso, cajas, "
                "       piezas_por_caja "
                "from costing.costos_validados where sku = any(%s::citext[])",
                (skus,)):
            # En MAYÚSCULAS: costos_validados guarda 'ROP-0695-BEI-m' y
            # core.products 'ROP-0695-BEI-M'. citext los empata; el dict no.
            salida[r["sku"].upper()] = {
                "largo": _num_o_none(r["largo"]), "ancho": _num_o_none(r["ancho"]),
                "alto": _num_o_none(r["alto"]), "peso": _num_o_none(r["peso"]),
                "cajas_pl": _num_o_none(r["cajas"]),
                "piezas_por_caja_pl": _num_o_none(r["piezas_por_caja"]),
            }
    except Exception as exc:  # noqa: BLE001
        log.warning("checklist: costos_validados no disponible: %s", exc)
    return salida


def _almacen(skus: list[str]) -> dict[str, dict[str, Any]]:
    capturado = _capturado(skus)
    sistema = _sistema(skus)
    salida = {}
    for s in skus:
        d = dict(capturado.get(s) or {**{k: None for k in _LOG_CLAVES},
                                      "capturado_por": None, "capturado_en": None})
        d["sistema"] = sistema.get(s.upper())
        salida[s] = d
    return salida


def almacen_de(skus: list[str]) -> dict[str, dict[str, Any]]:
    """Para el Catálogo Maestro: lo capturado, o {} si la 0058 no existe aún.
    NUNCA truena — el cotejo de cajas no puede caerse por unas columnas nuevas."""
    pedidos = {s.upper(): s for s in skus}
    try:
        # Con la grafía que PIDIÓ el Catálogo: core.products puede escribirlo
        # distinto (citext los empata; el dict no).
        return {pedidos.get(s.upper(), s): d for s, d in _capturado(skus).items()
                if d.get("capturado_en")}
    except Exception as exc:  # noqa: BLE001
        if not isinstance(exc, FaltaMigracion):
            log.warning("checklist: almacén no disponible: %s", exc)
        return {}


def _matriz(cats: Iterable[str]) -> dict[str, dict[str, bool]]:
    """Las promociones del equipo: filas `fuente = 'manual'` de ML en
    channel.field_requirements. Las de la API (`fuente = 'api'`) son los
    obligatorios del propio ML y NO son matriz."""
    cats = sorted({c for c in cats if c})
    if not cats:
        return {}
    salida: dict[str, dict[str, bool]] = {}
    for r in sdb.fetch_all(
            "select categoria_id, campo from channel.field_requirements "
            "where canal = %s and fuente = 'manual' and obligatorio "
            "  and categoria_id = any(%s)", (CANAL, cats)):
        salida.setdefault(r["categoria_id"], {})[r["campo"]] = True
    return salida


def _escribir_matriz(cat: str, promover: dict[str, str | None],
                     quitar: Iterable[str]) -> int:
    """Sube (inserta la fila manual) o baja (borra la fila manual). NUNCA toca
    una fila de la API: la guarda `fuente = 'manual'` vive en el SQL del upsert
    y del delete, no en quien llama."""
    quitar = list(quitar)

    def _hacer() -> int:
        n = 0
        with sdb.get_cursor() as cur:
            for campo, tipo in promover.items():
                cur.execute(
                    # campo_canonico = 'atributos', como las 2,753 filas de
                    # la API: channel_content.faltantes() revisa el atributo
                    # dentro de contenido->'atributos' SOLO con ese canónico;
                    # con NULL lo daba por faltante para siempre en el
                    # Publicador aunque Bodega ya lo hubiera llenado.
                    "insert into channel.field_requirements as fr "
                    "  (canal, categoria_id, campo, campo_canonico, obligatorio, "
                    "   tipo, fuente, leido_at) "
                    "values (%s, %s, %s, 'atributos', true, %s, 'manual', now()) "
                    "on conflict (canal, categoria_id, campo) do update set "
                    "  obligatorio = true, updated_at = now(), "
                    "  campo_canonico = coalesce(fr.campo_canonico, 'atributos') "
                    "where fr.fuente = 'manual'", (CANAL, cat, campo, tipo))
                n += cur.rowcount
            for campo in quitar:
                cur.execute(
                    "delete from channel.field_requirements "
                    "where canal = %s and categoria_id = %s and campo = %s "
                    "  and fuente = 'manual'", (CANAL, cat, campo))
                n += cur.rowcount
        return n
    return sdb.reintentar_transitorio(_hacer)


def _guardar_almacen(filas: dict[str, dict[str, Any]],
                     por: str) -> tuple[int, list[dict[str, str]]]:
    """UPDATE de core.products por SKU. Solo pisa lo que trae valor (coalesce):
    una captura parcial no borra lo que ya se había medido. La fila siempre
    existe: el SKU se validó contra core.products.

    UNA TRANSACCIÓN CORTA POR SKU. Antes iban todos en una y un solo valor que
    la base rechazara revertía las ~100 medidas de la semana sin decir cuál.
    Ahora un error del VALOR (22xxx desbordamiento, 23xxx CHECK) se queda en ese
    SKU y los demás se guardan.

    Por qué NO savepoints (se probaron y se quitaron): DBUtils (SteadyDB)
    re-ejecuta la sentencia en una conexión NUEVA, fuera de la transacción,
    cuando choca con el candado de solo-lectura heredado del pooler o con una
    conexión muerta; el savepoint se queda en la vieja y el RELEASE truena con
    3B001, que nadie reconoce como transitorio. Con una transacción por SKU ese
    cambio de conexión es inofensivo y reintentar_transitorio cura cada SKU por
    su cuenta (regla 13). Devuelve (guardados, fallidos)."""
    if not filas:
        return 0, []
    sets = ", ".join(f"{c} = coalesce(%s, {c})" for c in _COLUMNA.values())
    sql = (f"update core.products set {sets}, almacen_por = %s, "
           f"almacen_en = now() where sku = %s")
    n, fallidos = 0, []
    pendientes = list(filas.items())
    for i, (sku, vals) in enumerate(pendientes):
        def _uno(sku: str = sku, vals: dict[str, Any] = vals) -> int:
            with sdb.get_cursor() as cur:
                cur.execute(sql, (*[vals.get(k) for k in _COLUMNA], por or None, sku))
                return cur.rowcount
        try:
            n += sdb.reintentar_transitorio(_uno)
        except (psycopg2.DataError, psycopg2.IntegrityError) as exc:
            fallidos.append({"sku": sku, "motivo": f"medidas/cajas: {exc}".strip()})
        except Exception as exc:  # noqa: BLE001
            if _sin_tabla(exc):
                raise FaltaMigracion(_MSG_MIGRACION) from exc
            # La base no contesta (y el reintento no la curó): se para aquí y
            # se dice cuáles quedaron sin guardar, en vez de esperar un error
            # por cada uno de los que faltan.
            log.warning("checklist: almacén se detuvo en %s: %s", sku, exc)
            fallidos.extend({"sku": s, "motivo": f"no se guardó: {exc}".strip()}
                            for s, _ in pendientes[i:])
            break
    return n, fallidos


# ══════════════════════════════════════════════════════════════════════════════
# Los campos de una categoría, ya con su NIVEL
# ══════════════════════════════════════════════════════════════════════════════

def _automaticos(cats: Iterable[str]) -> dict[str, dict[str, str]]:
    """{categoría: {campo: valor por omisión}} de ML, de field_requirements
    (la categoría y el comodín '*'). Es lo que el PUBLICADOR llena solo si se
    deja vacío —BRAND = 'Ferrahome' en todas—, y por eso NO se le pide a
    almacén. Mismo criterio que specs._estado usa en el Catálogo Maestro."""
    cats = sorted({c for c in cats if c})
    if not cats:
        return {}
    try:
        reqs = specs._requisitos({(CANAL, c) for c in cats})
    except Exception as exc:  # noqa: BLE001
        log.warning("checklist: valores por omisión no disponibles: %s", exc)
        return {}
    salida: dict[str, dict[str, str]] = {}
    for c in cats:
        for r in reqs.get((CANAL, c)) or []:
            d = r.get("default")
            if d is not None and d != "":
                salida.setdefault(c, {})[r["campo"]] = (
                    d if isinstance(d, str) else json.dumps(d, ensure_ascii=False))
    return salida


def _nivel(c: dict[str, Any], promovidos: dict[str, bool],
           automaticos: dict[str, str] | None = None) -> str:
    """ml > matriz > principal > secundario.

    SECUNDARIO = jerarquía `ITEM` de ML: clave SAT, unidad de medida SAT, IVA,
    IEPS, número de pedimento, nombre en factura. Son datos FISCALES de la
    venta, no del producto: almacén no los conoce y se pliegan. Todo lo demás
    (llaves del catálogo, GTIN, características de familia) es PRINCIPAL.
    Medido el 24-sep: `relevance` vale 1 en TODOS los visibles, no separa nada.
    """
    if (automaticos or {}).get(c["campo"]):
        return "auto"
    if c.get("obligatorio"):
        return "ml"
    if promovidos.get(c["campo"]):
        return "matriz"
    if c.get("jerarquia") == "ITEM":
        return "secundario"
    return "principal"


def campos_de(cat: str, crudos: list[dict[str, Any]],
              promovidos: dict[str, bool],
              automaticos: dict[str, str] | None = None) -> list[dict[str, Any]]:
    """COPIAS de los campos de la caché de ML, con `nivel` y `exigido`.
    Copias: la lista de `_campos_ml` es compartida entre peticiones."""
    salida = []
    for c in crudos:
        d = {k: c.get(k) for k in ("campo", "etiqueta", "tipo", "jerarquia",
                                   "unidad_default")}
        d["valores"] = list(c.get("valores") or [])
        d["unidades"] = list(c.get("unidades") or [])
        d["nivel"] = _nivel(c, promovidos, automaticos)
        d["exigido"] = d["nivel"] in _EXIGIDOS
        d["por_omision"] = (automaticos or {}).get(c["campo"])
        salida.append(d)
    salida.sort(key=lambda x: NIVELES.index(x["nivel"]))
    return salida


# ══════════════════════════════════════════════════════════════════════════════
# El tablero
# ══════════════════════════════════════════════════════════════════════════════

def _vacio(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _contexto(skus: list[str]) -> dict[str, Any]:
    """Todo lo que hace falta para evaluar y exportar un grupo de SKUs, en
    pocas consultas: categorías, campos, matriz, valores y almacén."""
    info = specs._categorias(skus) if skus else {}
    cats_por_sku = {s: (info.get(s) or {}).get(CANAL) or {} for s in skus}
    cats = {v.get("categoria") for v in cats_por_sku.values() if v.get("categoria")}
    crudos = _campos_por_categoria(cats)
    matriz = _matriz(cats)
    auto = _automaticos(cats)
    campos = {c: campos_de(c, crudos.get(c) or [], matriz.get(c) or {}, auto.get(c) or {})
              for c in cats}
    return {
        "cats_por_sku": cats_por_sku,
        "campos": campos,
        "matriz": matriz,
        "nombres": _nombres_categoria(cats),
        "valores": _valores_ml(skus),
        "almacen": _almacen(skus),
        "titulos": _titulos(skus),
    }


def _evaluar(sku: str, ctx: dict[str, Any]) -> dict[str, Any]:
    cat_info = ctx["cats_por_sku"].get(sku) or {}
    cat = cat_info.get("categoria")
    campos = (ctx["campos"].get(cat) or []) if cat else []
    valores = ctx["valores"].get(sku) or {}
    alm = ctx["almacen"].get(sku) or {}
    nombre = (ctx["nombres"].get(cat) or {}) if cat else {}

    exigidos = [c for c in campos if c["exigido"]]
    faltan_ml = [{"campo": c["campo"], "etiqueta": c["etiqueta"], "nivel": c["nivel"]}
                 for c in exigidos if _vacio(valores.get(c["campo"]))]
    # Lo automático (BRAND → Ferrahome) ni se exige ni es opcional de almacén.
    opcionales = [c for c in campos if not c["exigido"] and c["nivel"] != "auto"]
    faltan_alm = [k for k in _LOG_CLAVES if alm.get(k) is None]

    if not cat:
        estado = "sin_categoria"
    elif not campos:
        estado = "sin_lista"
    elif not faltan_ml and not faltan_alm:
        estado = "completo"
    else:
        estado = "incompleto"

    cajas, ppc = alm.get("cajas"), alm.get("piezas_por_caja")
    return {
        "sku": sku,
        "titulo": ctx["titulos"].get(sku),
        "categoria": cat,
        "categoria_fuente": cat_info.get("fuente"),
        "categoria_nombre": nombre.get("nombre"),
        "categoria_ruta": nombre.get("ruta"),
        "exigidos_total": len(exigidos),
        "exigidos_llenos": len(exigidos) - len(faltan_ml),
        "faltan_ml": faltan_ml,
        "opcionales_total": len(opcionales),
        "opcionales_llenos": sum(1 for c in opcionales
                                 if not _vacio(valores.get(c["campo"]))),
        "almacen": {k: alm.get(k) for k in (*_LOG_CLAVES, "capturado_por",
                                             "capturado_en", "sistema")},
        "faltan_almacen": faltan_alm,
        "piezas_total": cajas * ppc if cajas is not None and ppc is not None else None,
        "estado": estado,
    }


def tablero_sync(semana_txt: str | None) -> dict[str, Any]:
    semana = semana_de(semana_txt)
    base = {"semana": semana.isoformat(),
            "semana_fin": (semana + dt.timedelta(days=6)).isoformat(),
            "etiqueta": etiqueta(semana),
            "campos_almacen": [{"campo": k, "etiqueta": e} for k, e, _ in LOGISTICA]}
    try:
        semanas = [{"semana": r["semana"].isoformat(), "etiqueta": etiqueta(r["semana"]),
                    "skus": r["n"]} for r in _q(
            "select semana, count(*) n from ops.checklist_lote "
            "group by semana order by semana desc limit 26")]
        # El orden de la LISTA: agregado_en usa clock_timestamp() renglón por
        # renglón, así la tabla sale en el orden en que almacén armó su hoja.
        lote = _q("select sku::text as sku, comentario, agregado_por, agregado_en "
                  "from ops.checklist_lote where semana = %s order by agregado_en, sku",
                  (semana,))
        skus = [r["sku"] for r in lote]
        ctx = _contexto(skus)
    except FaltaMigracion as exc:
        return {**base, "ok": False, "falta_migracion": True, "motivo": str(exc),
                "semanas": [], "filas": [], "categorias": [],
                "resumen": _resumen([])}

    urls = _urls_ml(skus)
    filas = []
    for r in lote:
        f = _evaluar(r["sku"], ctx)
        f["url_ml"] = urls.get(r["sku"])
        f["comentario"] = r.get("comentario")
        f["agregado_por"] = r["agregado_por"]
        f["agregado_en"] = r["agregado_en"].isoformat() if r["agregado_en"] else None
        filas.append(f)

    por_cat: dict[str, int] = {}
    for f in filas:
        if f["categoria"]:
            por_cat[f["categoria"]] = por_cat.get(f["categoria"], 0) + 1
    categorias = sorted((
        {"categoria": c, "nombre": (ctx["nombres"].get(c) or {}).get("nombre"),
         "skus": n,
         "obligatorios_ml": sum(1 for x in ctx["campos"].get(c) or [] if x["nivel"] == "ml"),
         "promovidos": sum(1 for x in ctx["campos"].get(c) or [] if x["nivel"] == "matriz")}
        for c, n in por_cat.items()), key=lambda x: (-x["skus"], x["nombre"] or ""))

    return {**base, "ok": True, "falta_migracion": False, "motivo": None,
            "semanas": semanas, "filas": filas, "categorias": categorias,
            "resumen": _resumen(filas)}


def _resumen(filas: list[dict[str, Any]]) -> dict[str, int]:
    return {
        "total": len(filas),
        "completos": sum(1 for f in filas if f["estado"] == "completo"),
        "incompletos": sum(1 for f in filas if f["estado"] == "incompleto"),
        "sin_categoria": sum(1 for f in filas if f["estado"] in ("sin_categoria", "sin_lista")),
        "faltan_ml": sum(1 for f in filas if f["faltan_ml"]),
        "faltan_almacen": sum(1 for f in filas if f["faltan_almacen"]),
    }


# ══════════════════════════════════════════════════════════════════════════════
# El lote de la semana
# ══════════════════════════════════════════════════════════════════════════════

def agregar_sync(semana_txt: str | None, crudo: Iterable[str] | str,
                 comentarios: dict[str, str] | None = None) -> dict[str, Any]:
    """Agrega SKUs al lote. `comentarios` va por SKU en MAYÚSCULAS (la columna
    «Comentarios» de la lista semanal); a un SKU que ya estaba solo se le
    actualiza el comentario."""
    semana = semana_de(semana_txt)
    pegados = limpiar_skus(crudo)
    if len(pegados) > _MAX_SKUS_LOTE:
        return {"ok": False, "motivo": f"Son {len(pegados)} SKUs; el máximo por "
                                       f"carga es {_MAX_SKUS_LOTE}."}
    reales, desconocidos = _canonicos(pegados)
    skus = [reales[s.upper()] for s in pegados if s.upper() in reales]
    comentarios = {k.upper(): v for k, v in (comentarios or {}).items() if v}
    por = actor.actual() or None
    nuevos = 0
    if skus:
        try:
            ya = {r["sku"].upper() for r in _q(
                "select sku::text as sku from ops.checklist_lote "
                "where semana = %s and sku = any(%s::citext[])", (semana, skus))}
        except FaltaMigracion as exc:
            return {"ok": False, "falta_migracion": True, "motivo": str(exc)}

        def _hacer() -> int:
            n = 0
            with sdb.get_cursor() as cur:
                for s in skus:
                    com = comentarios.get(s.upper())
                    if s.upper() not in ya:
                        cur.execute(
                            "insert into ops.checklist_lote (semana, sku, comentario, "
                            "agregado_por, agregado_en) values (%s, %s, %s, %s, "
                            "clock_timestamp()) on conflict (semana, sku) do nothing",
                            (semana, s, com, por))
                        n += cur.rowcount
                    elif com:
                        cur.execute("update ops.checklist_lote set comentario = %s "
                                    "where semana = %s and sku = %s", (com, semana, s))
            return n
        nuevos = sdb.reintentar_transitorio(_hacer)
    return {"ok": True, "semana": semana.isoformat(), "etiqueta": etiqueta(semana),
            "agregados": nuevos, "ya_estaban": len(skus) - nuevos,
            "desconocidos": desconocidos}


def quitar_sync(semana_txt: str | None, crudo: Iterable[str] | str) -> dict[str, Any]:
    semana = semana_de(semana_txt)
    skus = limpiar_skus(crudo)
    if not skus:
        return {"ok": True, "quitados": 0}
    try:
        n = sdb.execute("delete from ops.checklist_lote "
                        "where semana = %s and sku = any(%s::citext[])",
                        (semana, skus))
    except Exception as exc:  # noqa: BLE001
        if _sin_tabla(exc):
            return {"ok": False, "falta_migracion": True, "motivo": _MSG_MIGRACION}
        raise
    return {"ok": True, "quitados": n}


# ══════════════════════════════════════════════════════════════════════════════
# La matriz
# ══════════════════════════════════════════════════════════════════════════════

def matriz_sync(cat: str) -> dict[str, Any]:
    crudos = specs_editor._campos_ml(cat)
    nombre = _nombres_categoria([cat]).get(cat) or {}
    base = {"categoria": cat, "nombre": nombre.get("nombre"), "ruta": nombre.get("ruta")}
    campos = campos_de(cat, crudos, _matriz([cat]).get(cat) or {},
                       _automaticos([cat]).get(cat) or {})
    return {**base, "ok": True, "falta_migracion": False,
            "motivo": None if campos else
            "Mercado Libre no contestó qué pide esta categoría; intenta en un momento.",
            "campos": campos}


def guardar_matriz_sync(cat: str, cambios: dict[str, bool]) -> dict[str, Any]:
    """Sube (o baja) opcionales a obligatorios. Los obligatorios de ML no se
    tocan: bajarlos no serviría de nada, ML los rechazaría igual."""
    crudos = {c["campo"]: c for c in specs_editor._campos_ml(cat)}
    if not crudos:
        return {"ok": False, "motivo": "Mercado Libre no contestó qué pide esta "
                                       "categoría; no se guarda a ciegas."}
    validos = {k: bool(v) for k, v in (cambios or {}).items()
               if k in crudos and not crudos[k].get("obligatorio")}
    if not validos:
        return {"ok": True, "guardados": 0}
    promover = {k: crudos[k].get("tipo") for k, v in validos.items() if v}
    quitar = [k for k, v in validos.items() if not v]
    n = _escribir_matriz(cat, promover, quitar)
    # field_requirements no tiene columna de autor: queda en el log.
    log.info("checklist: %s cambió la matriz de %s: +%s -%s", actor.actual() or "?",
             cat, sorted(promover), sorted(quitar))
    return {"ok": True, "guardados": n}


# ══════════════════════════════════════════════════════════════════════════════
# Normalizar lo que escribe almacén
# ══════════════════════════════════════════════════════════════════════════════

_NUM = re.compile(r"^\s*(-?\d+(?:[.,]\d+)?)\s*([^\d\s].*?)?\s*$")
_SI = {"si", "sí", "s", "yes", "y", "true", "verdadero", "1", "x"}
_NO = {"no", "n", "false", "falso", "0"}


def _texto(v: Any) -> str:
    """Lo que trae una celda, como texto. 12.0 → '12'."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:g}"
    return str(v).strip()


# Unidades que se aceptan en lo de almacén, convertidas a la de la columna.
# Antes «250 g» se guardaba como 250 KG: la unidad escrita se ignoraba.
_UNIDADES_CM = {
    "cm": 1.0, "cms": 1.0, "centimetro": 1.0, "centimetros": 1.0,
    "centímetro": 1.0, "centímetros": 1.0,
    "mm": 0.1, "milimetro": 0.1, "milimetros": 0.1, "milímetro": 0.1, "milímetros": 0.1,
    "m": 100.0, "mt": 100.0, "mts": 100.0, "metro": 100.0, "metros": 100.0,
}
_UNIDADES_KG = {
    "kg": 1.0, "kgs": 1.0, "kilo": 1.0, "kilos": 1.0, "kilogramo": 1.0, "kilogramos": 1.0,
    "g": 0.001, "gr": 0.001, "grs": 0.001, "gramo": 0.001, "gramos": 0.001,
}
_PALABRAS_ENTERO = {"", "cajas", "caja", "cj", "cjs", "pzs", "pz", "pza", "pzas",
                    "piezas", "pieza", "pcs", "pc", "unidad", "unidades", "u"}
# Los topes del tipo de cada columna de core.products (numeric(8,2),
# numeric(9,3), integer). Pasarse daba 22003 al guardar, no en la vista previa.
_TOPE = {"decimal_cm": 999999.99, "decimal_kg": 999999.999, "entero": 2147483647}


def _normalizar_logistica(campo: str, crudo: str) -> tuple[Any, str | None]:
    """(valor, error). Acepta '12,5', '12.5 cm', '120 mm', '250 g', '3 kg'.

    Todo lo que la base fuera a rechazar se rechaza AQUÍ, en la vista previa,
    con un motivo legible: número fuera del tipo, algo que al redondear a la
    escala de la columna queda en 0, o una unidad que no es de esa medida."""
    m = _NUM.match(crudo)  # el \s de Python ya cubre el espacio duro de Excel
    if not m:
        return None, f"«{crudo}» no es un número"
    n = float(m.group(1).replace(",", "."))
    unidad = (m.group(2) or "").strip().lower().rstrip(".")
    if _LOG_TIPO[campo] == "entero":
        if unidad not in _PALABRAS_ENTERO:
            return None, f"«{crudo}»: escribe solo el número"
        if not n.is_integer():
            return None, f"«{crudo}» debe ser un número entero"
        n = int(n)
        if n > _TOPE["entero"]:
            return None, f"«{crudo}» es demasiado grande"
        if campo == "piezas_por_caja" and n <= 0:
            return None, "las piezas por caja deben ser más de 0"
        if n < 0:
            return None, "no puede ser negativo"
        return n, None
    es_peso = campo == "peso_kg"
    tabla = _UNIDADES_KG if es_peso else _UNIDADES_CM
    if unidad and unidad not in tabla:
        return None, (f"unidad «{unidad}» no válida; usa "
                      f"{'kg o g' if es_peso else 'cm, mm o m'}")
    # «1,200 g» ¿es 1.2 g o 1200 g? Con una unidad chica (g, mm) y justo tres
    # dígitos tras el separador, la coma casi siempre es de MILES. Adivinar da
    # un error de 1000×, así que se pide que se escriba sin ambigüedad.
    if tabla.get(unidad, 1.0) < 1.0 and re.fullmatch(r"[1-9]\d{0,2}[.,]\d{3}", m.group(1)):
        return None, (f"«{crudo}» es ambiguo: escribe "
                      f"{'1200 g o 1.2 kg' if es_peso else '1500 mm o 150 cm'}")
    n *= tabla.get(unidad, 1.0)
    escala, tope = (3, _TOPE["decimal_kg"]) if es_peso else (2, _TOPE["decimal_cm"])
    n = round(n, escala)
    if n <= 0:
        return None, ("debe ser mayor que 0" if float(m.group(1).replace(",", ".")) <= 0
                      else f"«{crudo}» es tan chico que se redondea a 0")
    if n > tope:
        return None, f"«{crudo}» es demasiado grande"
    return n, None


def _normalizar_ml(c: dict[str, Any], crudo: str) -> tuple[str | None, str | None, str | None]:
    """(valor, error, aviso) según el tipo que declara ML."""
    tipo = c.get("tipo") or "string"
    if tipo == "boolean":
        t = crudo.strip().lower()
        if t in _SI:
            return "Sí", None, None
        if t in _NO:
            return "No", None, None
        return None, f"«{crudo}» debe ser Sí o No", None
    if tipo == "number":
        m = _NUM.match(crudo)
        if not m or m.group(2):
            return None, f"«{crudo}» debe ser un número", None
        return m.group(1).replace(",", "."), None, None
    if tipo == "number_unit":
        m = _NUM.match(crudo)
        if not m:
            return None, f"«{crudo}» debe ser número y unidad (p. ej. 12 cm)", None
        numero = m.group(1).replace(",", ".")
        unidad = (m.group(2) or "").strip()
        permitidas = c.get("unidades") or []
        if not unidad:
            # Sin unidad se usa la que ML mismo asume (`default_unit`), o la
            # única permitida. NUNCA «la primera de la lista»: en Peso máximo
            # soportado es `g`, y «20» de un corral son 20 kg, no 20 g.
            por_omision = c.get("unidad_default") or (
                permitidas[0] if len(permitidas) == 1 else None)
            if not permitidas:
                return numero, None, None
            if not por_omision:
                return None, (f"falta la unidad ({', '.join(permitidas)})"), None
            return (f"{numero} {por_omision}", None,
                    f"sin unidad: se usó {por_omision}, la que Mercado Libre asume")
        if permitidas:
            igual = next((u for u in permitidas if u.lower() == unidad.lower()), None)
            if not igual:
                return None, (f"unidad «{unidad}» no válida; ML acepta "
                              f"{', '.join(permitidas)}"), None
            unidad = igual
        return f"{numero} {unidad}", None, None
    valores = c.get("valores") or []
    if valores:
        igual = next((v for v in valores if v.lower() == crudo.lower()), None)
        if igual:
            return igual, None, None
        if tipo == "list":
            return crudo, None, "no está en la lista de Mercado Libre; puede rechazarlo"
    return crudo, None, None


# ══════════════════════════════════════════════════════════════════════════════
# La descarga: Excel y CSV
# ══════════════════════════════════════════════════════════════════════════════

# Los colores del canal: el amarillo de Mercado Libre para lo que ML exige.
_COLOR = {
    "ml": "FFE600", "matriz": "FDBA74", "auto": "D1FAE5", "principal": "E2E8F0",
    "secundario": "F1F5F9", "almacen": "BFDBFE", "id": "E5E7EB",
}
_HUECO = {"ml": "FFF9C4", "matriz": "FFEDD5", "almacen": "DBEAFE"}
_NIVEL_TXT = {"ml": "Obligatorio ML", "matriz": "Obligatorio (matriz)",
              "auto": "Automático",
              "principal": "Opcional", "secundario": "Opcional · facturación"}


def _skus_export(semana_txt: str | None, sel: str | None) -> tuple[dt.date, list[str]]:
    semana = semana_de(semana_txt)
    elegidos = limpiar_skus(sel)
    if elegidos:
        return semana, elegidos
    return semana, [r["sku"] for r in _q(
        "select sku::text as sku from ops.checklist_lote where semana = %s "
        "order by agregado_en, sku", (semana,))]


def _texto_sistema(ref: dict[str, Any] | None) -> str:
    """La referencia de costos_validados en una línea: lo que HOY dice el
    sistema, que almacén NO debe copiar (son flete y packing list, no medidas)."""
    if not ref:
        return "sin datos en el sistema"
    partes = []
    if all(ref.get(k) for k in ("largo", "ancho", "alto")):
        partes.append(f"{ref['largo']:g}×{ref['ancho']:g}×{ref['alto']:g} cm")
    if ref.get("peso"):
        partes.append(f"{ref['peso']:g} kg")
    if ref.get("cajas_pl") is not None or ref.get("piezas_por_caja_pl") is not None:
        partes.append(f"PL {_texto(ref.get('cajas_pl')) or '—'} cajas × "
                      f"{_texto(ref.get('piezas_por_caja_pl')) or '—'} pzs")
    return " · ".join(partes) or "sin datos en el sistema"


def _nombre_archivo(semana: dt.date, n: int, ext: str) -> str:
    anio, num, _ = semana.isocalendar()
    return f"checklist_week{num}_{anio}_{n}skus.{ext}"


def _grupos(skus: list[str], ctx: dict[str, Any]) -> list[tuple[str | None, list[str]]]:
    """SKUs agrupados por categoría ML (una hoja por categoría)."""
    por: dict[str | None, list[str]] = {}
    for s in skus:
        por.setdefault((ctx["cats_por_sku"].get(s) or {}).get("categoria"), []).append(s)

    def llave(item):
        c = item[0]
        return (c is None, ((ctx["nombres"].get(c) or {}).get("nombre") or c or ""))
    return sorted(por.items(), key=llave)


def _nombre_hoja(texto: str, usados: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", texto or "Hoja").strip()[:28] or "Hoja"
    nombre, i = base, 2
    while nombre.lower() in usados:
        nombre = f"{base[:25]} ({i})"
        i += 1
    usados.add(nombre.lower())
    return nombre


def _pista(c: dict[str, Any]) -> str:
    partes = [_NIVEL_TXT[c["nivel"]]]
    if c["nivel"] == "auto":
        partes.append(f"si lo dejas vacío: {c.get('por_omision')}")
    elif c.get("tipo") == "boolean":
        partes.append("Sí / No")
    elif c.get("tipo") == "number_unit" and c.get("unidades"):
        partes.append("número + " + "/".join(c["unidades"][:4])
                      + (f" (sin unidad = {c['unidad_default']})"
                         if c.get("unidad_default") else ""))
    elif c.get("tipo") == "number":
        partes.append("número")
    elif c.get("valores"):
        partes.append("elige de la lista")
    return " · ".join(partes)


def excel_sync(semana_txt: str | None, sel: str | None) -> tuple[bytes, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    semana, skus = _skus_export(semana_txt, sel)
    ctx = _contexto(skus)

    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"
    listas = wb.create_sheet("_listas")
    listas.sheet_state = "hidden"
    col_lista = [0]

    fill = lambda c: PatternFill("solid", start_color=c, end_color=c)  # noqa: E731
    fino = Side(style="thin", color="CBD5E1")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    negrita = Font(bold=True, color="0F172A")

    def lista_ref(valores: list[str]) -> str:
        col_lista[0] += 1
        letra = get_column_letter(col_lista[0])
        for i, v in enumerate(valores[:500], start=1):
            listas.cell(row=i, column=col_lista[0], value=v)
        return f"'_listas'!${letra}$1:${letra}${min(len(valores), 500)}"

    usados = {"instrucciones", "_listas"}
    indice = []
    for cat, grupo in _grupos(skus, ctx):
        campos = (ctx["campos"].get(cat) or []) if cat else []
        nombre_cat = ((ctx["nombres"].get(cat) or {}).get("nombre") or cat) if cat else None
        ws = wb.create_sheet(_nombre_hoja(nombre_cat or "Sin categoría ML", usados))
        indice.append((ws.title, nombre_cat or "— sin categoría ML —", cat or "",
                       len(grupo), sum(1 for c in campos if c["exigido"])))

        columnas: list[dict[str, Any]] = [
            {"clave": "sku", "titulo": "SKU", "pista": "no lo cambies", "nivel": "id", "ancho": 20},
            {"clave": "titulo", "titulo": "Producto", "pista": "solo referencia",
             "nivel": "id", "ancho": 42},
            {"clave": "sistema", "titulo": "En sistema (NO es medida)",
             "pista": "flete y packing list: no lo copies, mide", "nivel": "id",
             "ancho": 34},
        ]
        columnas += [{"clave": k, "titulo": e, "pista": "Almacén · " + (
            "entero" if t == "entero" else "número"), "nivel": "almacen", "ancho": 13,
            "tipo_log": t} for k, e, t in LOGISTICA]
        columnas += [{"clave": c["campo"], "titulo": c["etiqueta"], "pista": _pista(c),
                      "nivel": c["nivel"], "ancho": max(14, min(34, len(c["etiqueta"] or "") + 4)),
                      "campo": c} for c in campos]

        for j, col in enumerate(columnas, start=1):
            ws.cell(row=1, column=j, value=col["clave"])
            h = ws.cell(row=2, column=j, value=col["titulo"])
            h.font = Font(bold=True, color="0F172A" if col["nivel"] != "ml" else "1E1B4B")
            h.fill = fill(_COLOR[col["nivel"]])
            h.alignment = Alignment(wrap_text=True, vertical="center")
            h.border = borde
            p = ws.cell(row=3, column=j, value=col["pista"])
            p.font = Font(italic=True, size=9, color="475569")
            p.fill = fill(_COLOR[col["nivel"]])
            p.alignment = Alignment(wrap_text=True, vertical="top")
            p.border = borde
            ws.column_dimensions[get_column_letter(j)].width = col["ancho"]
        ws.row_dimensions[1].hidden = True
        ws.row_dimensions[2].height = 34
        ws.row_dimensions[3].height = 30
        ws.freeze_panes = "C4"

        ultima = 3 + len(grupo)
        for i, sku in enumerate(grupo, start=4):
            valores = ctx["valores"].get(sku) or {}
            alm = ctx["almacen"].get(sku) or {}
            for j, col in enumerate(columnas, start=1):
                clave = col["clave"]
                if clave == "sku":
                    v: Any = sku
                elif clave == "titulo":
                    v = ctx["titulos"].get(sku) or ""
                elif clave == "sistema":
                    v = _texto_sistema(alm.get("sistema"))
                elif col["nivel"] == "almacen":
                    v = alm.get(clave)
                else:
                    v = valores.get(clave) or None
                celda = ws.cell(row=i, column=j, value=v)
                celda.border = borde
                if clave == "sku":
                    celda.font = negrita
                if _vacio(v) and col["nivel"] in _HUECO:
                    celda.fill = fill(_HUECO[col["nivel"]])

        # Validaciones: sugerencias con AVISO (ML acepta texto libre en casi
        # todo) y números estrictos en lo de almacén.
        for j, col in enumerate(columnas, start=1):
            letra = get_column_letter(j)
            rango = f"{letra}4:{letra}{max(ultima, 4)}"
            if col["nivel"] == "almacen":
                if col["tipo_log"] == "entero":
                    dv = DataValidation(type="whole", operator="greaterThanOrEqual",
                                        formula1="0", allow_blank=True,
                                        error="Escribe un número entero", errorTitle="Almacén")
                else:
                    dv = DataValidation(type="decimal", operator="greaterThan",
                                        formula1="0", allow_blank=True,
                                        error="Escribe un número mayor que 0",
                                        errorTitle="Almacén")
            elif "campo" in col and col["campo"].get("tipo") == "boolean":
                dv = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
            elif "campo" in col and col["campo"].get("valores"):
                dv = DataValidation(type="list", formula1=lista_ref(col["campo"]["valores"]),
                                    allow_blank=True, errorStyle="warning",
                                    error="No está en la lista de Mercado Libre. "
                                          "¿Seguro?", errorTitle="Mercado Libre")
            else:
                continue
            ws.add_data_validation(dv)
            dv.add(rango)

        # Los de facturación se agrupan PLEGADOS: están para quien los sepa,
        # sin tapar lo que sí se exige.
        sec = [j for j, col in enumerate(columnas, start=1) if col["nivel"] == "secundario"]
        if sec:
            ws.column_dimensions.group(get_column_letter(sec[0]), get_column_letter(sec[-1]),
                                       hidden=True, outline_level=1)
        ws.sheet_properties.outlinePr.summaryRight = False

    _instrucciones(ins, semana, len(skus), indice, fill, negrita)
    if not col_lista[0]:
        listas.cell(row=1, column=1, value="")
    # La hoja de listas al final: oculta, pero que no quede entre las de trabajo.
    wb.move_sheet(listas, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("_listas"))

    buf = io.BytesIO()
    wb.save(buf)
    nombre = _nombre_archivo(semana, len(skus), "xlsx")
    return buf.getvalue(), nombre


def _instrucciones(ws, semana: dt.date, n: int, indice: list, fill, negrita) -> None:
    from openpyxl.styles import Alignment, Font

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 14
    ws["A1"] = "CHECKLIST DE ALMACÉN — Mercado Libre"
    ws["A1"].font = Font(bold=True, size=16, color="1E1B4B")
    fin = semana + dt.timedelta(days=6)
    ws["A2"] = (f"{etiqueta(semana)} · del {semana.strftime('%d/%m')} al "
                f"{fin.strftime('%d/%m/%Y')} · {n} SKUs · generado "
                f"{dt.datetime.now(_ZONA).strftime('%d/%m/%Y %H:%M')}"
                + (f" por {actor.actual()}" if actor.actual() else ""))
    ws["A2"].font = Font(color="475569")

    pasos = [
        "1. Cada hoja es una categoría de Mercado Libre. Llena las celdas de color.",
        "2. AZUL = almacén: MIDE el producto EMPACADO (cm y kg), cuenta las CAJAS "
        "y cuántas PIEZAS trae cada caja.",
        "   La columna gris «En sistema» es lo que hoy dice el sistema (flete y "
        "packing list). NO la copies: es la referencia contra la que se compara.",
        "3. AMARILLO = lo que Mercado Libre exige. NARANJA = lo que el equipo "
        "decidió exigir (matriz).",
        "4. Las columnas grises son opcionales. Las de FACTURACIÓN (clave SAT, "
        "IVA, IEPS, pedimento) vienen plegadas: no son de almacén.",
        "5. NO cambies la columna SKU ni borres el renglón oculto 1: así se sabe "
        "qué es cada cosa al cargarlo.",
        "6. Una celda vacía NO borra nada. Lo que almacén ya capturó antes viene "
        "pre-llenado: corrígelo si está mal.",
        "7. Guarda y cárgalo en Omnicanal → Inventario → Checklist → «Cargar "
        "Excel». Antes de guardar te enseña qué va a cambiar.",
    ]
    for i, t in enumerate(pasos, start=4):
        ws.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=False)

    fila = 4 + len(pasos) + 1
    ws.cell(row=fila, column=1, value="Colores").font = negrita
    for k, t in (("almacen", "Almacén (medidas, cajas, piezas)"),
                 ("ml", "Obligatorio de Mercado Libre"),
                 ("matriz", "Obligatorio por la matriz del equipo"),
                 ("auto", "Automático: el publicador lo llena si lo dejas vacío"),
                 ("principal", "Opcional del producto"),
                 ("secundario", "Opcional de facturación (plegado)")):
        fila += 1
        ws.cell(row=fila, column=1, value="").fill = fill(_COLOR[k])
        ws.cell(row=fila, column=2, value=t)

    fila += 2
    for j, t in enumerate(("Hoja", "Categoría de Mercado Libre", "ID", "SKUs",
                           "Exigidos"), start=1):
        ws.cell(row=fila, column=j, value=t).font = negrita
    for hoja, nombre, cid, cuantos, exig in indice:
        fila += 1
        for j, v in enumerate((hoja, nombre, cid, cuantos, exig), start=1):
            ws.cell(row=fila, column=j, value=v)


def csv_sync(semana_txt: str | None, sel: str | None) -> tuple[bytes, str]:
    """Formato LARGO: un renglón por (SKU, campo). Es el que sirve para
    automatizar (un script o un agente lo llena sin conocer las hojas)."""
    semana, skus = _skus_export(semana_txt, sel)
    ctx = _contexto(skus)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["sku", "producto", "categoria_id", "categoria", "campo", "etiqueta",
                "nivel", "tipo", "unidades", "en_sistema", "valor"])
    # La referencia de costos_validados que corresponde a cada campo de almacén.
    ref_de = {"largo_cm": "largo", "ancho_cm": "ancho", "alto_cm": "alto",
              "peso_kg": "peso", "cajas": "cajas_pl",
              "piezas_por_caja": "piezas_por_caja_pl"}
    for sku in skus:
        cat = (ctx["cats_por_sku"].get(sku) or {}).get("categoria")
        nombre = ((ctx["nombres"].get(cat) or {}).get("nombre") or "") if cat else ""
        titulo = ctx["titulos"].get(sku) or ""
        alm = ctx["almacen"].get(sku) or {}
        for k, e, t in LOGISTICA:
            v = alm.get(k)
            ref = (alm.get("sistema") or {}).get(ref_de[k])
            w.writerow([sku, titulo, cat or "", nombre, k, e, "almacen", t, "",
                        _texto(ref), "" if v is None else _texto(v)])
        valores = ctx["valores"].get(sku) or {}
        for c in ((ctx["campos"].get(cat) or []) if cat else []):
            w.writerow([sku, titulo, cat, nombre, c["campo"], c["etiqueta"], c["nivel"],
                        c.get("tipo") or "", "/".join(c.get("unidades") or []), "",
                        valores.get(c["campo"], "")])
    nombre = _nombre_archivo(semana, len(skus), "csv")
    # BOM: sin él, Excel en español abre los acentos como basura.
    return ("﻿" + buf.getvalue()).encode("utf-8"), nombre


# ══════════════════════════════════════════════════════════════════════════════
# La carga
# ══════════════════════════════════════════════════════════════════════════════

def _celdas_xlsx(datos: bytes) -> tuple[list[tuple[str, str, str, str, int]], list[dict]]:
    """[(sku, campo, valor, hoja, fila)], errores de estructura."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
    celdas, errores = [], []
    for ws in wb.worksheets:
        if ws.title.lower() in ("instrucciones", "_listas"):
            continue
        filas = ws.iter_rows(values_only=True)
        claves: list[str] | None = None
        n_fila = 0
        for n_fila, fila in enumerate(filas, start=1):
            if claves is None:
                if fila and _texto(fila[0]).lower() == "sku" and n_fila <= 5:
                    claves = [_texto(x) for x in fila]
                    inicio = n_fila + 3   # claves, títulos, pistas
                elif n_fila > 5:
                    break
                continue
            if n_fila < inicio or not fila:
                continue
            sku = _texto(fila[0]) if fila else ""
            if not sku:
                continue
            for j, clave in enumerate(claves):
                if j == 0 or not clave or clave in ("titulo", "sistema"):
                    continue
                v = _texto(fila[j]) if j < len(fila) else ""
                if v:
                    celdas.append((sku, clave, v, ws.title, n_fila))
        if claves is None:
            errores.append({"sku": None, "campo": None, "hoja": ws.title, "fila": None,
                            "motivo": "La hoja no trae el renglón de claves (el oculto "
                                      "1). Descarga la plantilla desde el panel."})
    wb.close()
    return celdas, errores


def _celdas_csv(datos: bytes) -> tuple[list[tuple[str, str, str, str, int]], list[dict]]:
    texto = datos.decode("utf-8-sig", errors="replace")
    muestra = texto[:4096]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t")
        # El sniff mira solo 4 KB: si ahí no hay un SKU con comillas, deduce
        # doublequote=False y lee 'HERR-0032-ROJ-16""'. Excel y el panel
        # siempre escriben las comillas DOBLADAS.
        dialecto.doublequote = True
    except csv.Error:
        dialecto = csv.excel
    lector = csv.DictReader(io.StringIO(texto), dialect=dialecto)
    cab = {(c or "").strip().lower() for c in (lector.fieldnames or [])}
    if not {"sku", "campo", "valor"} <= cab:
        return [], [{"sku": None, "campo": None, "hoja": "CSV", "fila": None,
                     "motivo": "El CSV necesita las columnas sku, campo y valor "
                               "(formato largo, el que descarga el panel)."}]
    celdas = []
    for n, r in enumerate(lector, start=2):
        r = {(k or "").strip().lower(): (v or "") for k, v in r.items()}
        sku, campo, valor = r["sku"].strip(), r["campo"].strip(), r["valor"].strip()
        if sku and campo and valor:
            celdas.append((sku, campo, valor, "CSV", n))
    return celdas, []


def importar_sync(datos: bytes, nombre_archivo: str, aplicar: bool) -> dict[str, Any]:
    ext = (nombre_archivo or "").lower().rsplit(".", 1)[-1]
    try:
        if ext in ("xlsx", "xlsm"):
            celdas, errores = _celdas_xlsx(datos)
        elif ext == "csv":
            celdas, errores = _celdas_csv(datos)
        else:
            return {"ok": False, "motivo": "Sube el Excel (.xlsx) o el CSV que "
                                           "descargaste del panel."}
    except Exception as exc:  # noqa: BLE001
        log.warning("checklist: no se pudo leer %s: %s", nombre_archivo, exc)
        return {"ok": False, "motivo": f"No se pudo leer el archivo: {exc}"}

    avisos: list[dict[str, Any]] = []
    # Cada celda SKU se limpia UNA vez y se busca ya limpia: con la celda
    # cruda, un SKU envuelto en comillas («"MIC-0001-GRI"») o leído del CSV con
    # comillas dobles se encontraba en kubera y aun así se tiraba en silencio.
    limpio = {p: (limpiar_skus([p]) or [p])[0] for p in {c[0] for c in celdas}}
    pegados = limpiar_skus([limpio[c[0]] for c in celdas])
    reales, desconocidos = _canonicos(pegados)
    for s in desconocidos:
        errores.append({"sku": s, "campo": None, "hoja": None, "fila": None,
                        "motivo": "SKU que kubera no conoce; se ignoró."})
    skus = [reales[s.upper()] for s in pegados if s.upper() in reales]
    try:
        ctx = _contexto(skus)
    except FaltaMigracion as exc:
        return {"ok": False, "falta_migracion": True, "motivo": str(exc)}

    ml: dict[str, dict[str, str]] = {}
    etiquetas: dict[str, dict[str, str]] = {}
    alm: dict[str, dict[str, Any]] = {}
    cambios: list[dict[str, Any]] = []
    sin_cambios = 0
    vistos: set[tuple[str, str]] = set()

    for pegado, campo, crudo, hoja, fila in celdas:
        sku = reales.get(limpio.get(pegado, pegado).upper())
        if not sku:
            continue
        donde = {"sku": sku, "campo": campo, "hoja": hoja, "fila": fila}
        if (sku, campo) in vistos:
            avisos.append({**donde, "motivo": "el campo viene repetido; se usó el primero"})
            continue
        vistos.add((sku, campo))

        if campo in _LOG_CLAVES:
            valor, error = _normalizar_logistica(campo, crudo)
            if error:
                errores.append({**donde, "motivo": f"{_LOG_ETIQUETA[campo]}: {error}"})
                continue
            antes = (ctx["almacen"].get(sku) or {}).get(campo)
            if antes is not None and float(antes) == float(valor):
                sin_cambios += 1
                continue
            alm.setdefault(sku, {})[campo] = valor
            cambios.append({"sku": sku, "campo": campo, "etiqueta": _LOG_ETIQUETA[campo],
                            "antes": _texto(antes), "despues": _texto(valor),
                            "tipo": "almacen"})
            continue

        cat = (ctx["cats_por_sku"].get(sku) or {}).get("categoria")
        if not cat:
            errores.append({**donde, "motivo": "el SKU no tiene categoría de Mercado "
                                               "Libre; sus atributos no se guardan."})
            continue
        c = next((x for x in ctx["campos"].get(cat) or [] if x["campo"] == campo), None)
        if not c:
            avisos.append({**donde, "motivo": f"{campo} ya no está en la categoría "
                                              f"{cat}; se ignoró."})
            continue
        antes = (ctx["valores"].get(sku) or {}).get(campo, "")
        # La celda que nadie tocó (viene pre-llenada) NO se normaliza: si el
        # valor guardado no trae unidad, «arreglarlo» aquí sería inventar un
        # cambio que almacén no hizo.
        if crudo == antes:
            sin_cambios += 1
            continue
        valor, error, aviso = _normalizar_ml(c, crudo)
        if error:
            errores.append({**donde, "motivo": f"{c['etiqueta']}: {error}"})
            continue
        if aviso:
            avisos.append({**donde, "motivo": f"{c['etiqueta']}: {aviso}"})
        if antes == valor:
            sin_cambios += 1
            continue
        ml.setdefault(sku, {})[campo] = valor
        etiquetas.setdefault(sku, {})[campo] = c["etiqueta"]
        cambios.append({"sku": sku, "campo": campo, "etiqueta": c["etiqueta"],
                        "antes": antes, "despues": valor, "tipo": "ml"})

    salida = {"ok": True, "aplicado": False, "archivo": nombre_archivo,
              "skus": len({c["sku"] for c in cambios}), "cambios": cambios,
              "errores": errores, "avisos": avisos, "sin_cambios": sin_cambios}
    if not aplicar or not cambios:
        return salida

    fallidos: list[dict[str, str]] = []
    guardados_ml = 0

    def _uno(sku: str) -> tuple[str, dict[str, Any]]:
        return sku, specs_editor.guardar_sync(sku, CANAL, ml[sku], etiquetas.get(sku))

    if ml:
        with ThreadPoolExecutor(max_workers=min(_HILOS_GUARDAR, len(ml))) as ex:
            for sku, res in ex.map(_uno, list(ml)):
                if res.get("ok"):
                    guardados_ml += 1
                else:
                    fallidos.append({"sku": sku, "motivo": res.get("motivo") or "?"})
    guardados_alm = 0
    if alm:
        try:
            guardados_alm, fallidos_alm = _guardar_almacen(alm, actor.actual())
            fallidos.extend(fallidos_alm)
        except FaltaMigracion as exc:
            fallidos.append({"sku": "—", "motivo": str(exc)})
        except Exception as exc:  # noqa: BLE001
            log.warning("checklist: almacén no se guardó: %s", exc)
            fallidos.append({"sku": "—", "motivo": f"medidas/cajas: {exc}"})
    log.info("checklist: %s cargó %s — %d SKUs ML, %d de almacén, %d fallidos",
             actor.actual() or "?", nombre_archivo, guardados_ml, guardados_alm,
             len(fallidos))
    return {**salida, "aplicado": True,
            "guardados": {"ml": guardados_ml, "almacen": guardados_alm,
                          "fallidos": fallidos}}


def guardar_almacen_sync(sku: str, valores: dict[str, Any]) -> dict[str, Any]:
    """Captura desde la pantalla: los seis campos de almacén de UN SKU."""
    reales, _ = _canonicos([sku])
    real = reales.get(sku.upper())
    if not real:
        return {"ok": False, "motivo": f"{sku} no existe en kubera."}
    fila: dict[str, Any] = {}
    for k in _LOG_CLAVES:
        crudo = _texto((valores or {}).get(k))
        if not crudo:
            continue
        v, error = _normalizar_logistica(k, crudo)
        if error:
            return {"ok": False, "motivo": f"{_LOG_ETIQUETA[k]}: {error}"}
        fila[k] = v
    if not fila:
        return {"ok": True, "guardados": 0}
    try:
        _, fallidos = _guardar_almacen({real: fila}, actor.actual())
    except FaltaMigracion as exc:
        return {"ok": False, "falta_migracion": True, "motivo": str(exc)}
    except Exception as exc:  # noqa: BLE001
        log.warning("checklist: almacén de %s no se guardó: %s", real, exc)
        return {"ok": False, "motivo": f"No se pudo guardar: {exc}"}
    if fallidos:
        return {"ok": False, "motivo": fallidos[0]["motivo"]}
    return {"ok": True, "guardados": len(fila)}


# ══════════════════════════════════════════════════════════════════════════════
# La lista de la semana — el Excel que arma el equipo («Week 39»)
# ══════════════════════════════════════════════════════════════════════════════
#
# Cada semana llegaba con otras columnas (Week 36: SKU y contenedor sin
# encabezado; Week 37: SKU, stock y almacén; Week 39: sku entre corchetes,
# contenedor, tarima, comentarios). El ESTÁNDAR que se acepta es el mínimo
# común: una hoja «Week NN» con una columna «SKU» (con o sin corchetes) y, si
# hay, una «Comentarios». Lo demás se ignora: contenedor ya vive en
# costos_validados y la tarima es ubicación de Odoo, no se repiten aquí.

_HOJA_SEMANA = re.compile(r"(?i)\b(?:week|wk|semana|sem)\s*[-_.]?\s*(\d{1,2})\b")
_ENCABEZADO_SKU = re.compile(r"^(c[oó]digo|clave|id)\s*(de\s*)?sku")
# Con comillas: hay SKUs reales como HERR-0032-ROJ-16" (12 en core.products).
# Prefijo de letras (con Ñ: BAÑ-0486-EST), guion y lo que sea sin espacios:
# hay SKUs reales con * ´ ° > , y " (MUE-0445-VER-180*300, TEC-1813-NEG-19´…).
# La EXISTENCIA la decide _canonicos, no esta forma.
_PARECE_SKU = re.compile(r"^[A-Za-zÑñ]{2,6}-\S+$")


def _sku_de_lista(v: Any) -> str:
    """'[MIC-0001-GRI]' → 'MIC-0001-GRI'. Las listas copian el SKU como lo
    escribe Ferraforme, entre corchetes."""
    return re.sub(r"[\[\]\s]", "", _texto(v))


def _hojas_lista(datos: bytes, nombre: str,
                 conocidos=None) -> list[dict[str, Any]]:
    """Cada hoja que trae SKUs: {hoja, semana_iso, skus, comentarios, descartados}.

    `conocidos(lista) -> set(MAYÚSCULAS)` dice cuáles existen en kubera; con él
    se elige la columna SKU. Sin él (pruebas), se puntúa por la forma."""
    tablas: list[tuple[str, list[list[Any]]]] = []
    if nombre.lower().endswith(".csv"):
        texto = datos.decode("utf-8-sig", errors="replace")
        try:
            dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
            dialecto.doublequote = True   # ver _celdas_csv
        except csv.Error:
            dialecto = csv.excel
        tablas.append((nombre.rsplit(".", 1)[0],
                       list(csv.reader(io.StringIO(texto), dialecto))))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
        for ws in wb.worksheets:
            filas = []
            for i, fila in enumerate(ws.iter_rows(values_only=True)):
                if i >= 5000:
                    break
                filas.append(list(fila))
            tablas.append((ws.title, filas))
        wb.close()

    salida = []
    for titulo, filas in tablas:
        col_sku, col_com, inicio = None, None, 0
        # El ENCABEZADO y la COLUMNA se eligen juntos: entre los primeros 5
        # renglones, cada celda que diga «sku» es candidata, y gana la que más
        # SKUs de kubera tenga debajo. Así un renglón de título («Lista de SKUs
        # semana 39») no se confunde con el encabezado, y de un packing list con
        # 'SKU ODOO' y 'SKU' (referencias del proveedor) gana la de Odoo.
        mejor = (0, 0, -1)
        for i, fila in enumerate(filas[:5]):
            textos = [_texto(x).strip().lower() for x in fila]
            for j, x in enumerate(textos):
                # Un ENCABEZADO de SKU: empieza con «sku» ('SKU', 'SKU ODOO',
                # 'SKUs') o es «código/clave SKU». Por subcadena, un renglón de
                # datos como «crear SKU» se volvía encabezado y se comía los SKUs
                # de arriba.
                if not (x.startswith("sku") or _ENCABEZADO_SKU.match(x)):
                    continue
                debajo = [_sku_de_lista(f[j]) for f in filas[i + 1:] if j < len(f)]
                debajo = [s for s in debajo if s and _PARECE_SKU.match(s)]
                reales = len(conocidos(debajo)) if conocidos and debajo else 0
                # En EMPATE gana el renglón más bajo, el pegado a los datos: un
                # título «SKUs semana 40» arriba del encabezado real ve los
                # mismos SKUs debajo, y si ganaba se perdían los comentarios.
                puntos = (reales, len(debajo), i)
                if puntos[:2] > (0, 0) and puntos > mejor:
                    mejor = puntos
                    col_sku, inicio = j, i + 1
                    col_com = next((k for k, y in enumerate(textos)
                                    if y.startswith("comentario")), None)
        if col_sku is None:
            col_sku = 0   # sin encabezado (así venía la Week 36): la primera columna
        skus: list[str] = []
        comentarios: dict[str, str] = {}
        descartados: list[str] = []
        vistos: set[str] = set()
        for fila in filas[inicio:]:
            if col_sku >= len(fila):
                continue
            s = _sku_de_lista(fila[col_sku])
            if s and not _PARECE_SKU.match(s) and s not in descartados:
                descartados.append(s)
            if not s or not _PARECE_SKU.match(s) or s.upper() in vistos:
                continue
            vistos.add(s.upper())
            skus.append(s)
            if col_com is not None and col_com < len(fila):
                c = _texto(fila[col_com]).strip()
                if c:
                    comentarios[s.upper()] = c
        if skus:
            m = _HOJA_SEMANA.search(titulo)
            salida.append({"hoja": titulo, "semana_iso": int(m.group(1)) if m else None,
                           "skus": skus, "comentarios": comentarios,
                           "descartados": descartados})
    return salida


def lista_sync(datos: bytes, nombre: str, semana_txt: str | None,
               hoja: str | None, aplicar: bool) -> dict[str, Any]:
    """Carga la lista semanal al lote. Con `aplicar=false` dice qué hojas vio,
    cuál tomaría y a qué semana; con `true` la agrega.

    La hoja se elige así: la que se pidió; si no, la «Week NN» de la semana que
    se está viendo; si no, la de número más alto. La semana sale del NOMBRE de
    la hoja («Week 39» → lunes 21-sep-2026), en el año ISO de la semana vista.
    """
    def conocidos(lista: list[str]) -> set[str]:
        reales, _ = _canonicos(lista)
        return {s.upper() for s in lista if s.upper() in reales}

    try:
        hojas = _hojas_lista(datos, nombre or "", conocidos)
    except Exception as exc:  # noqa: BLE001
        log.warning("checklist: no se pudo leer la lista %s: %s", nombre, exc)
        return {"ok": False, "motivo": f"No se pudo leer el archivo: {exc}"}
    if not hojas:
        return {"ok": False, "motivo": "No encontré ninguna hoja con una columna SKU."}

    vista = semana_de(semana_txt)
    anio, num_vista, _ = vista.isocalendar()

    def lunes_de(h: dict[str, Any]) -> dt.date:
        return (lunes_iso(anio, h["semana_iso"]) if h["semana_iso"] else None) or vista

    elegida = next((h for h in hojas if hoja and h["hoja"] == hoja), None)
    if elegida is None:
        con_numero = [h for h in hojas if h["semana_iso"]]
        elegida = (next((h for h in con_numero if h["semana_iso"] == num_vista), None)
                   or max(con_numero, key=lambda h: h["semana_iso"], default=None)
                   or hojas[0])
    semana = lunes_de(elegida)
    _, desconocidos = _canonicos(elegida["skus"])
    base = {
        "ok": True,
        "hojas": [{"hoja": h["hoja"], "semana_iso": h["semana_iso"],
                   "skus": len(h["skus"]), "semana": lunes_de(h).isoformat(),
                   "etiqueta": etiqueta(lunes_de(h))} for h in hojas],
        "elegida": elegida["hoja"], "semana": semana.isoformat(),
        "etiqueta": etiqueta(semana), "skus": len(elegida["skus"]),
        "comentarios": len(elegida["comentarios"]), "desconocidos": desconocidos,
        # Lo que venía en la columna SKU y no parece SKU: se enseña, no se
        # tira en silencio.
        "descartados": elegida["descartados"][:30],
        "descartados_total": len(elegida["descartados"]),
    }
    if not aplicar:
        return {**base, "aplicado": False}
    res = agregar_sync(semana.isoformat(), elegida["skus"], elegida["comentarios"])
    if not res.get("ok"):
        return {**base, **res, "aplicado": False}
    log.info("checklist: %s cargó la lista %s/%s → %s: %d nuevos, %d ya estaban",
             actor.actual() or "?", nombre, elegida["hoja"], etiqueta(semana),
             res.get("agregados", 0), res.get("ya_estaban", 0))
    return {**base, **res, "aplicado": True}
