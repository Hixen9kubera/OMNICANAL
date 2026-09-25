"""
ubicar_contenedores_xlsx.py — La lista para Brandon: lo que la carga de
`costing.sku_contenedor` (0060) NO pudo decidir sola, en un Excel.

Sale de la MISMA corrida que carga la tabla (`scripts/ubicar_skus_contenedor.py
--excel`), no de un análisis aparte: si una regla cambia, la tabla y la lista
cambian juntas. Puro: recibe lo que la corrida ya calculó y arma el libro; no
habla con la BD, Odoo ni Drive (lo fijan `tests/test_ubicar_contenedores_xlsx.py`).

HOJAS
  Léeme                         qué es cada hoja, la fecha y cuántos renglones trae.
  1 Conflictos                  SKUs con varias N y alguna sin documento: la N
                                y el texto que da cada fuente, y qué decidir.
  2 Refutados                   Odoo es la única fuente y el Ferraforme de esa
                                N no trae el SKU: ¿falta en el Ferraforme o
                                Odoo está mal?
  3 Provisionales               los «NNNN-NNNN» por homologar, por contenedor.
                                Ni kubera ni Odoo guardan su descripción: la
                                pista es el costo y las medidas capturadas.
  4 Errores en Odoo             Odoo escribe una N que los documentos
                                contradicen (el lote del 12 capturado como 34,
                                «OOLU9155398 - cont 98» que es el 94). Se
                                corrigen EN Odoo: nosotros no escribimos ahí.
  5 Contenedores sin documento  N con SKUs reales que ningún Ferraforme
                                indexado de esa N trae, incluidos los Ferraforme
                                que están en packing_archivos y no se indexaron.
                                Un archivo cuyo nombre no dice «contenedor N»
                                (los originales vienen solo con el código) se
                                asigna por sus códigos (`mapa_codigos`), no por
                                `contenedor_base`, que está mal en varios.

Sin fórmulas: todo es valor. Un texto que empieza con «=» se guarda como texto.
"""
from __future__ import annotations

import collections
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services import ubicar_contenedores as U

_CAB_FILL = PatternFill("solid", fgColor="1F3864")
_CONOCIDO_FILL = PatternFill("solid", fgColor="FFF2CC")
_INT = "#,##0"
_DEC = "#,##0.00"

HOJAS = ("1 Conflictos", "2 Refutados", "3 Provisionales", "4 Errores en Odoo",
         "5 Contenedores sin documento")

# Cómo se nombra cada fuente interna en el Excel.
QUIEN = {U.FERRA_ALINEADO: "el Ferraforme", U.FERRA_NO_ALINEADO: "el Ferraforme",
         U.COSTOS: "costos (tu captura)", U.CAJA: "la caja compartida",
         U.ODOO_CAMPO: "el campo de Odoo", U.ODOO_PELON: "el código de Odoo (sin número)",
         U.ODOO_AMBIGUO: "Odoo (el código y el número no cuadran)", U.OC_RECIBIDA: "una OC recibida"}

# Columnas por fuente de la hoja de conflictos: (título, fuentes internas).
GRUPOS = (("Ferraforme", (U.FERRA_ALINEADO, U.FERRA_NO_ALINEADO)),
          ("Costos", (U.COSTOS, U.CAJA)),
          ("Odoo", (U.ODOO_CAMPO, U.ODOO_AMBIGUO, U.ODOO_PELON)),
          ("OC recibida", (U.OC_RECIBIDA,)))

ODOO_EXPLICITO = frozenset({U.ODOO_CAMPO, U.ODOO_AMBIGUO})

# Nombre corto de los documentos primarios (dentro de un paréntesis).
CORTO = {U.FERRA_ALINEADO: "Ferraforme", U.FERRA_NO_ALINEADO: "Ferraforme", U.COSTOS: "costos",
         U.OC_RECIBIDA: "OC recibida"}


# ── Utilidades ────────────────────────────────────────────────────────────────

MESES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
         "octubre", "noviembre", "diciembre")


def fecha_larga(dt) -> str:
    """datetime → «25 de septiembre de 2026, 14:05»."""
    return f"{dt.day} de {MESES[dt.month - 1]} de {dt.year}, {dt:%H:%M}"


def _f(**kw) -> Font:
    return Font(name="Arial", **{"size": 10, **kw})


def _ns(ns: Iterable[Any]) -> str:
    return ", ".join(str(n) for n in sorted({int(n) for n in ns}))


def _num(valor: Any) -> float | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, Decimal):
        return float(valor)
    try:
        return float(valor)
    except (TypeError, ValueError):
        return None


def _g(x: Any) -> str:
    """12.0 → «12»; 2.5 → «2.5»."""
    v = _num(x)
    return "" if v is None else f"{v:g}"


def texto_ref(fuente: str, r: dict) -> str:
    """Una referencia de evidencia → texto legible para Brandon."""
    if fuente in (U.FERRA_ALINEADO, U.FERRA_NO_ALINEADO):
        t = f"{r.get('archivo') or '?'}, fila {r.get('fila')}"
        return t + (" (cotejado)" if fuente == U.FERRA_ALINEADO else "")
    if fuente == U.COSTOS:
        return f"«{r.get('valor')}»"
    if fuente == U.CAJA:
        return f"caja compartida «{r.get('contenedor')}» ({r.get('archivo') or 'sin archivo'})"
    if fuente == U.ODOO_CAMPO:
        return f"«{r.get('texto')}»" + ("" if r.get("activo", True) else " (producto inactivo)")
    if fuente == U.ODOO_PELON:
        return f"solo código «{r.get('texto')}»"
    if fuente == U.ODOO_AMBIGUO:
        return (f"«{r.get('texto')}»: el código es del {r.get('codigo_dice')}, "
                f"el texto dice {r.get('texto_dice')}")
    if fuente == U.OC_RECIBIDA:
        return f"{r.get('oc')} ({_g(r.get('recibido'))} de {_g(r.get('pedido'))} recibidas)"
    return str(r)


def _detalle(ev: dict, fuentes: Iterable[str]) -> tuple[list[int], str]:
    """(N que dan esas fuentes, «N: texto | N: texto») sin repetir textos."""
    ns: set[int] = set()
    partes: list[str] = []
    for f in fuentes:
        for n, refs in sorted(ev.get(f, {}).items(), key=lambda kv: int(kv[0])):
            ns.add(int(n))
            for r in refs:
                t = texto_ref(f, r)
                # El ambiguo trae la misma ref en sus dos N: se escribe una vez.
                linea = t if f == U.ODOO_AMBIGUO else f"{n}: {t}"
                if linea not in partes:
                    partes.append(linea)
    return sorted(ns), " | ".join(partes)


def _quien_dice(ev: dict, n: int) -> list[str]:
    return [f for f in U.ORDEN_FUENTES if str(n) in ev.get(f, {})]


def _ambiguo(ev: dict) -> dict | None:
    for refs in ev.get(U.ODOO_AMBIGUO, {}).values():
        if refs:
            return refs[0]
    return None


def que_decidir(e: dict) -> str:
    """La pregunta concreta de un conflicto, en una o dos frases."""
    ev = e.get("evidencia") or {}
    sin = [int(n) for n in (e.get("sin_documento") or [])]
    con = [int(n) for n in e.get("ns") or [] if int(n) not in sin]
    amb = _ambiguo(ev)
    frases = []
    oc_no = e.get("oc_no_documenta") or {}
    if con:
        for n in sin:
            quien = sorted({QUIEN[f] for f in _quien_dice(ev, n)})
            donde = ("Odoo" if any(f.startswith("odoo") or f == U.OC_RECIBIDA for f in _quien_dice(ev, n))
                     else " y ".join(quien))
            razon = oc_no.get(n) or oc_no.get(str(n))
            nota = f" (la OC no cuenta como documento: {'; '.join(razon)})" if razon else ""
            frases.append(f"El {n} solo lo dice {' y '.join(quien)}{nota}, sin documento del {n}: si no llegó "
                          f"ahí, corregir {donde}; si también llegó en el {n}, mandar el documento del {n}.")
        frases.insert(0, f"Los documentos lo ponen en el {_ns(con)}.")
    else:
        frases.append(f"Ninguna de sus N ({_ns(sin)}) tiene documento (Ferraforme, costos u OC recibida): "
                      f"confirmar en cuál llegó.")
        for n in sin:
            razon = oc_no.get(n) or oc_no.get(str(n))
            if razon:
                frases.append(f"La OC del {n} no cuenta como documento: {'; '.join(razon)}.")
    if amb:
        frases.append(f"En Odoo el código es del {amb.get('codigo_dice')} y el texto dice "
                      f"{amb.get('texto_dice')}.")
    return " ".join(frases)


def es_error_odoo(e: dict) -> bool:
    """¿Todas las N sin documento las dice SOLO Odoo, y con número escrito
    (campo o código/número que no cuadran)? El pelón no entra: su N es
    nuestra lectura del código, no algo que Odoo escribió."""
    ev = e.get("evidencia") or {}
    sin = [int(n) for n in (e.get("sin_documento") or [])]
    if e.get("motivo") != "conflicto" or not sin:
        return False
    for n in sin:
        quien = set(_quien_dice(ev, n))
        if not quien or not quien <= {U.ODOO_CAMPO, U.ODOO_AMBIGUO, U.ODOO_PELON} or not quien & ODOO_EXPLICITO:
            return False
    return True


def patron_odoo(e: dict) -> tuple[str, str, str]:
    """(patrón, N correcta probable y por qué, qué corregir en Odoo)."""
    ev = e.get("evidencia") or {}
    sin = [int(n) for n in (e.get("sin_documento") or [])]
    con = [int(n) for n in e.get("ns") or [] if int(n) not in sin]
    amb = _ambiguo(ev)
    if amb and not con:
        cd, td = amb.get("codigo_dice"), amb.get("texto_dice")
        return (f"código del {cd}, texto dice {td}",
                f"{cd} (lo dice el código)",
                f"Cambiar el número del texto «{amb.get('texto')}» a {cd}, si el código es el correcto.")
    dice = _ns(n for n in sin)
    fuentes_con = sorted({CORTO[f] for n in con for f in _quien_dice(ev, n) if f in U.PRIMARIAS})
    correcta = f"{_ns(con)} ({', '.join(fuentes_con)})" if con else ""
    textos = sorted({r.get("texto") for f in (U.ODOO_CAMPO, U.ODOO_AMBIGUO)
                     for n in sin for r in ev.get(f, {}).get(str(n), []) if r.get("texto")})
    if con:
        que = (f"Cambiar «{' / '.join(textos)}» para que diga el {_ns(con)} "
               f"(si llegó en los dos, poner ambos).")
    else:
        que = f"Ningún documento respalda el {dice}: confirmar en cuál llegó y corregir «{' / '.join(textos)}»."
    return f"documentos dicen {_ns(con) or '—'}, Odoo dice {dice}", correcta, que


# Errores de lote que el análisis del 25-sep ya confirmó (se resaltan).
def es_error_conocido(patron: str, e: dict) -> bool:
    amb = _ambiguo(e.get("evidencia") or {})
    if amb and "OOLU9155398" in str(amb.get("texto") or "").upper():
        return True
    return patron == "documentos dicen 12, Odoo dice 34"


# ── Escritura de hojas ────────────────────────────────────────────────────────

def _valor(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def _hoja(wb: Workbook, titulo: str, cabs: list[str], filas: list[list], *,
          anchos_max: dict[int, int] | None = None, formatos: dict[int, str] | None = None,
          resaltar: Iterable[int] = ()) -> int:
    """Encabezado en la fila 1 (congelado, con filtro) y un renglón por fila."""
    ws = wb.create_sheet(titulo)
    for c, t in enumerate(cabs, 1):
        cel = ws.cell(1, c, t)
        cel.font = _f(bold=True, color="FFFFFF")
        cel.fill = _CAB_FILL
        cel.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    resaltar = set(resaltar)
    formatos = formatos or {}
    for i, fila in enumerate(filas, 2):
        for c, v in enumerate(fila, 1):
            cel = ws.cell(i, c, _valor(v))
            if isinstance(cel.value, str) and cel.value.startswith("="):
                cel.data_type = "s"             # texto, nunca fórmula
            cel.font = _f()
            cel.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
            if c in formatos and isinstance(cel.value, (int, float)):
                cel.number_format = formatos[c]
            if i - 2 in resaltar:
                cel.fill = _CONOCIDO_FILL
    ultima = get_column_letter(len(cabs))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ultima}{max(1, len(filas) + 1)}"
    anchos_max = anchos_max or {}
    for c, t in enumerate(cabs, 1):
        largos = sorted(len(str(f[c - 1])) for f in filas if c - 1 < len(f) and f[c - 1] not in (None, ""))
        p90 = largos[int(len(largos) * 0.9)] if largos else 0
        ancho = max(min(len(t), 28) + 2, p90 + 2, 9)
        ws.column_dimensions[get_column_letter(c)].width = min(ancho, anchos_max.get(c, 60))
    return len(filas)


def _nombre(nombres: dict[str, str], sku: str) -> str:
    return nombres.get(U.clave(sku), "")


def _conflictos(excluidos: list[dict], nombres: dict[str, str], hoja_odoo: set[str]) -> tuple[list[str], list[list]]:
    cabs = ["SKU", "Nombre (catálogo)", "N candidatas", "N con documento", "N sin documento"]
    for titulo, _ in GRUPOS:
        cabs += [f"{titulo}: N", f"{titulo}: archivo o texto"]
    cabs += ["Qué decidir", "También en hoja 4"]
    filas = []
    for e in excluidos:
        if e["motivo"] != "conflicto":
            continue
        sin = [int(n) for n in e.get("sin_documento") or []]
        fila = [e["sku"], _nombre(nombres, e["sku"]), _ns(e["ns"]),
                _ns(n for n in e["ns"] if int(n) not in sin), _ns(sin)]
        for _, fuentes in GRUPOS:
            ns, det = _detalle(e["evidencia"], fuentes)
            fila += [_ns(ns), det]
        fila += [que_decidir(e), "sí" if U.clave(e["sku"]) in hoja_odoo else ""]
        filas.append(fila)
    return cabs, filas


def _refutados(excluidos: list[dict], nombres: dict[str, str],
               ferra_por_n: dict[int, dict]) -> tuple[list[str], list[list]]:
    cabs = ["SKU", "Nombre (catálogo)", "Nombre en Odoo", "N que da Odoo", "Texto del campo en Odoo",
            "Ferraforme de esa N (archivo)", "SKUs en ese Ferraforme", "Otras pistas (código sin número, OC)",
            "Qué confirmar"]
    filas = []
    for e in excluidos:
        if e["motivo"] != "refutado":
            continue
        ev = e["evidencia"]
        ns, texto = _detalle(ev, (U.ODOO_CAMPO,))
        n = ns[0] if ns else None
        fe = ferra_por_n.get(n, {})
        _, otras = _detalle(ev, (U.ODOO_PELON, U.OC_RECIBIDA))
        # Si el nombre de Odoo es otro producto, el SKU choca en Odoo: no falta en el Ferraforme.
        n_odoo = next((r.get("nombre_odoo") for refs in ev.get(U.ODOO_CAMPO, {}).values() for r in refs
                       if r.get("nombre_odoo")), "")
        filas.append([e["sku"], _nombre(nombres, e["sku"]), n_odoo, n,
                      texto.split(": ", 1)[-1] if len(ns) == 1 else texto,
                      " | ".join(sorted(fe.get("archivos", ()))), len(fe.get("skus", ())), otras,
                      f"El Ferraforme del {n} no lo trae. Si sí llegó en el {n}, falta en el Ferraforme; "
                      f"si no, corregir el campo de Odoo (o el SKU, si el nombre en Odoo es otro producto)."])
    return cabs, filas


def _provisionales(excluidos: list[dict], costos_por_sku: dict[str, dict], codigo_por_n: dict[int, str],
                   cargados_por_n: dict[int, set]) -> tuple[list[str], list[list]]:
    cabs = ["Identificador", "Contenedor (N)", "Código del contenedor", "Valor capturado en costos",
            "Provisionales en ese contenedor", "SKUs reales cargados en ese contenedor", "Costo producto",
            "Largo", "Ancho", "Alto", "Peso", "Piezas por caja", "Cajas"]
    prov = [e for e in excluidos if e["motivo"] == "provisional"]

    def grupo(e: dict) -> tuple:
        # Sin N, el contenedor es el valor capturado tal cual («ONEU10522791»).
        return tuple(e["ns"]) or ("", U._txt(costos_por_sku.get(U.clave(e["sku"]), {}).get("contenedor")))

    por_n = collections.Counter(grupo(e) for e in prov)
    filas = []
    for e in sorted(prov, key=lambda e: (not e["ns"], tuple(e["ns"]), grupo(e), U.clave(e["sku"]))):
        c = costos_por_sku.get(U.clave(e["sku"]), {})
        n = e["ns"][0] if len(e["ns"]) == 1 else (_ns(e["ns"]) or None)
        cod = codigo_por_n.get(n, "") if isinstance(n, int) else ""
        filas.append([e["sku"], n, cod, U._txt(c.get("contenedor")), por_n[grupo(e)],
                      len(cargados_por_n.get(n, ())) if isinstance(n, int) else None,
                      _num(c.get("costo_producto")), _num(c.get("largo")), _num(c.get("ancho")),
                      _num(c.get("alto")), _num(c.get("peso")), _num(c.get("piezas_por_caja")),
                      _num(c.get("cajas"))])
    return cabs, filas


def _errores_odoo(excluidos: list[dict], nombres: dict[str, str]) -> tuple[list[str], list[list], list[int]]:
    cabs = ["SKU", "Nombre (catálogo)", "Qué dice Odoo (texto del campo)", "N que dice Odoo",
            "N correcta probable (según)", "Patrón", "SKUs con el mismo patrón", "Ya confirmado (25-sep)",
            "Qué corregir en Odoo"]
    sel = [e for e in excluidos if es_error_odoo(e)]
    pat = {U.clave(e["sku"]): patron_odoo(e) for e in sel}
    cuenta = collections.Counter(p[0] for p in pat.values())
    filas, resaltar = [], []
    orden = sorted(sel, key=lambda e: (-cuenta[pat[U.clave(e["sku"])][0]], pat[U.clave(e["sku"])][0],
                                       U.clave(e["sku"])))
    for i, e in enumerate(orden):
        patron, correcta, que = pat[U.clave(e["sku"])]
        sin = [int(n) for n in e.get("sin_documento") or []]
        _, texto = _detalle({f: e["evidencia"].get(f, {}) for f in (U.ODOO_CAMPO, U.ODOO_AMBIGUO)},
                            (U.ODOO_CAMPO, U.ODOO_AMBIGUO))
        conocido = es_error_conocido(patron, e)
        if conocido:
            resaltar.append(i)
        filas.append([e["sku"], _nombre(nombres, e["sku"]), texto, _ns(sin), correcta, patron,
                      cuenta[patron], "sí" if conocido else "por revisar", que])
    return cabs, filas, resaltar


def _sin_documento(*, evidencias: list[dict], asignaciones: list[dict], ns_con_ferraforme: set[int],
                   codigo_por_n: dict[int, str], archivos: list[dict], shas_indexados: set[str],
                   catalogo: dict[str, str], nombres: dict[str, str], ferra_por_n: dict[int, dict],
                   mapa_codigos: dict[str, set[int]] | None = None) -> tuple[list[str], list[list]]:
    cabs = ["Contenedor (N)", "Código", "Situación", "Archivos Ferraforme en packing_archivos",
            "…de esos, sin indexar", "SKUs indexados del Ferraforme", "Packing list original en packing_archivos",
            "SKUs reales que lo nombran", "…que no están en su Ferraforme", "…cargados en la tabla",
            "…por costos", "…por el campo de Odoo",
            "…por el código de Odoo (sin número)", "…por OC recibida", "Provisionales en ese contenedor",
            "Ejemplos fuera del Ferraforme (SKU — nombre)", "Qué pedir"]
    reales: dict[int, dict[str, set]] = collections.defaultdict(lambda: collections.defaultdict(set))
    prov: dict[int, set] = collections.defaultdict(set)
    for e in evidencias:
        k = U.clave(e["sku"])
        if e["fuente"] == U.ODOO_AMBIGUO:
            continue                            # código y número no cuadran: no nombra un contenedor
        if U.es_provisional(k):
            prov[e["numero"]].add(k)
        elif k in catalogo:
            reales[e["numero"]][e["fuente"]].add(k)
    cargados: dict[int, set] = collections.defaultdict(set)
    for a in asignaciones:
        cargados[a["numero"]].add(U.clave(a["sku"]))
    archivos_n: dict[int, dict[str, list]] = collections.defaultdict(lambda: collections.defaultdict(list))
    for a in archivos:
        n = U.numero_de_archivo(a.get("nombre"), a.get("miembro"))
        if n is None and mapa_codigos:
            # «EISU8559654 Lista de empaque.xlsx»: el original solo trae el código.
            n, _ = U.n_por_codigos(f"{U._txt(a.get('nombre'))} {U._txt(a.get('miembro'))}", mapa_codigos)
        if n is not None:
            archivos_n[n][a.get("tipo") or ""].append(a)
    filas = []
    for n in sorted(set(reales) | {n for n, t in archivos_n.items() if t.get("ferraforme")}):
        if n in ns_con_ferraforme:
            continue
        f = reales.get(n, {})
        todos = set().union(*f.values()) if f else set()
        if not todos:
            continue
        ferras = archivos_n.get(n, {}).get("ferraforme", [])
        sin_indexar = [a for a in ferras if (a.get("sha256") or "") not in shas_indexados]
        en_ferra = ferra_por_n.get(n, {}).get("skus", set())
        k_ferra = len(en_ferra)
        fuera = todos - en_ferra
        if not fuera:
            # Contenedor chico con su Ferraforme completo (el 4, el 5…): tiene
            # documento aunque no llegue al umbral para refutar.
            continue
        originales = len(archivos_n.get(n, {}).get("original", []))
        if ferras and sin_indexar and not k_ferra:
            situacion, pedir = ("Hay Ferraforme en packing_archivos pero no se indexó",
                                f"Revisar por qué no se indexó el Ferraforme del {n} y volver a indexarlo.")
        elif k_ferra:
            situacion, pedir = (f"Ferraforme con solo {k_ferra} SKUs; {len(fuera)} que lo nombran no están en él",
                                f"Confirmar si el Ferraforme del {n} está completo.")
        elif originales:
            situacion, pedir = ("Solo packing list original, sin Ferraforme",
                                f"Conseguir el Ferraforme del {n}.")
        else:
            situacion, pedir = ("Sin Ferraforme ni packing list en kubera",
                                f"Conseguir el Ferraforme (o la packing list) del {n}.")
        ejemplos = "; ".join(f"{catalogo.get(k, k)} — {nombres.get(k, '')}".rstrip(" —")
                             for k in sorted(fuera)[:5])
        filas.append([n, codigo_por_n.get(n, ""), situacion, len(ferras), len(sin_indexar), k_ferra, originales,
                      len(todos), len(fuera), len(cargados.get(n, set()) & todos),
                      len(f.get(U.COSTOS, set()) | f.get(U.CAJA, set())), len(f.get(U.ODOO_CAMPO, set())),
                      len(f.get(U.ODOO_PELON, set())), len(f.get(U.OC_RECIBIDA, set())), len(prov.get(n, ())),
                      ejemplos, pedir])
    return cabs, filas


def _leeme(wb: Workbook, conteos: dict[str, int], generado: str, resumen: dict) -> None:
    ws = wb.active
    ws.title = "Léeme"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    ws["A1"] = "Contenedores por SKU: lo que falta decidir"
    ws["A1"].font = _f(bold=True, size=14)
    ws["A2"] = f"Fecha: {generado}"
    ws["A2"].font = _f(italic=True)
    frases = [
        "Este libro sale del mismo proceso que cargó la tabla de SKU por contenedor en el sandbox; aquí va "
        f"solo lo que ese proceso no pudo decidir solo (los {resumen.get('skus_cargables', 0):,} SKUs que sí "
        "quedaron ubicados no se repiten aquí, salvo como ejemplo en la hoja 5).",
        "«1 Conflictos» son SKUs que aparecen en dos o más contenedores y al menos uno no tiene documento "
        "(Ferraforme, costos u OC recibida); cada renglón trae lo que dice cada fuente y qué decidir. Una OC "
        "no cuenta como documento si su contenedor se dedujo de los otros SKUs de la orden, si el Ferraforme "
        "de ese contenedor no trae el SKU o si recibió lo mismo que otra OC de otro contenedor.",
        "«2 Refutados» son SKUs que solo Odoo pone en un contenedor cuyo Ferraforme no los trae: o faltan "
        "en ese Ferraforme o el campo de Odoo está mal (si el nombre en Odoo es otro producto, el SKU choca "
        "en Odoo).",
        "«3 Provisionales» son los identificadores NNNN-NNNN que siguen sin homologar a un SKU real; ni el "
        "catálogo ni Odoo guardan su descripción, así que la pista es el costo y las medidas que se "
        "capturaron en costos.",
        "«4 Errores en Odoo» son SKUs donde Odoo escribe un contenedor que los documentos contradicen, "
        "agrupados por patrón (en amarillo los lotes ya confirmados: el 12 capturado como 34 y "
        "«OOLU9155398 - cont 98», que es el 94); se corrigen en Odoo, porque nosotros no escribimos ahí.",
        "«5 Contenedores sin documento» son los números con SKUs reales que no están en ningún Ferraforme "
        "indexado de ese contenedor (incluye los Ferraforme que están en kubera y no se indexaron); con el "
        "Ferraforme de cada uno se resuelven de golpe muchos renglones de las otras hojas.",
        "Lo de las hojas 1 a 4 no se cargó en la tabla: al corregir la fuente (Odoo, costos o el Ferraforme), "
        "la siguiente corrida del proceso lo toma sola.",
    ]
    for i, t in enumerate(frases, 4):
        ws.cell(i, 1, t).font = _f()
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 40
    r = 4 + len(frases) + 1
    for c, t in enumerate(("Hoja", "Renglones"), 1):
        cel = ws.cell(r, c, t)
        cel.font = _f(bold=True, color="FFFFFF")
        cel.fill = _CAB_FILL
    for i, h in enumerate(HOJAS, r + 1):
        ws.cell(i, 1, h).font = _f()
        ws.cell(i, 2, conteos.get(h, 0)).font = _f()
        ws.cell(i, 2).number_format = _INT
        ws.cell(i, 2).alignment = Alignment(horizontal="left")


# ── Libro completo ────────────────────────────────────────────────────────────

def ferraforme_por_numero(evidencias: Iterable[dict]) -> dict[int, dict]:
    """N → {archivos, skus} del Ferraforme indexado."""
    salida: dict[int, dict] = collections.defaultdict(lambda: {"archivos": set(), "skus": set()})
    for e in evidencias:
        if e["fuente"] in (U.FERRA_ALINEADO, U.FERRA_NO_ALINEADO):
            salida[e["numero"]]["archivos"].add(U._txt((e.get("ref") or {}).get("archivo")))
            salida[e["numero"]]["skus"].add(U.clave(e["sku"]))
    return dict(salida)


def armar_libro(*, excluidos: list[dict], asignaciones: list[dict], evidencias: list[dict],
                ns_con_ferraforme: set[int], codigo_por_numero: dict[int, str], archivos: list[dict],
                shas_indexados: set[str], catalogo: dict[str, str], nombres: dict[str, str],
                costos: list[dict], generado: str, resumen: dict | None = None,
                mapa_codigos: dict[str, set[int]] | None = None) -> tuple[Workbook, dict[str, int]]:
    """Todo lo que la corrida ya calculó → (libro, renglones por hoja)."""
    ferra_por_n = ferraforme_por_numero(evidencias)
    costos_por_sku = {U.clave(c.get("sku")): c for c in costos}
    cargados_por_n: dict[int, set] = collections.defaultdict(set)
    for a in asignaciones:
        cargados_por_n[a["numero"]].add(U.clave(a["sku"]))
    wb = Workbook()
    conteos: dict[str, int] = {}
    hoja_odoo = {U.clave(e["sku"]) for e in excluidos if es_error_odoo(e)}

    cabs, filas = _conflictos(excluidos, nombres, hoja_odoo)
    anchos = {2: 40, 3: 14, 4: 14, 5: 14, 6: 12, 7: 50, 8: 12, 9: 40, 10: 12, 11: 50, 12: 12, 13: 40, 14: 70}
    conteos[HOJAS[0]] = _hoja(wb, HOJAS[0], cabs, filas, anchos_max=anchos)
    cabs, filas = _refutados(excluidos, nombres, ferra_por_n)
    conteos[HOJAS[1]] = _hoja(wb, HOJAS[1], cabs, filas, anchos_max={2: 40, 3: 40, 5: 40, 6: 50, 8: 40, 9: 70},
                              formatos={4: "0", 7: _INT})
    cabs, filas = _provisionales(excluidos, costos_por_sku, codigo_por_numero, cargados_por_n)
    conteos[HOJAS[2]] = _hoja(wb, HOJAS[2], cabs, filas,
                              formatos={2: "0", 5: _INT, 6: _INT, 7: _DEC, 8: _DEC, 9: _DEC, 10: _DEC,
                                        11: _DEC, 12: _INT, 13: _INT})
    cabs, filas, resaltar = _errores_odoo(excluidos, nombres)
    conteos[HOJAS[3]] = _hoja(wb, HOJAS[3], cabs, filas, anchos_max={2: 40, 3: 50, 5: 40, 6: 40, 9: 70},
                              formatos={7: _INT}, resaltar=resaltar)
    cabs, filas = _sin_documento(evidencias=evidencias, asignaciones=asignaciones,
                                 ns_con_ferraforme=set(ns_con_ferraforme), codigo_por_n=codigo_por_numero,
                                 archivos=archivos, shas_indexados=shas_indexados, catalogo=catalogo,
                                 nombres=nombres, ferra_por_n=ferra_por_n, mapa_codigos=mapa_codigos)
    conteos[HOJAS[4]] = _hoja(wb, HOJAS[4], cabs, filas, anchos_max={3: 45, 16: 80, 17: 60},
                              formatos={1: "0", **{c: _INT for c in range(4, 16)}})
    _leeme(wb, conteos, generado, resumen or {})
    return wb, conteos


def escribir_libro(ruta: Path | str, **kw) -> dict[str, int]:
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb, conteos = armar_libro(**kw)
    wb.save(ruta)
    return conteos


__all__ = ["armar_libro", "escribir_libro", "fecha_larga", "que_decidir", "es_error_odoo", "patron_odoo", "texto_ref",
           "ferraforme_por_numero", "HOJAS"]
