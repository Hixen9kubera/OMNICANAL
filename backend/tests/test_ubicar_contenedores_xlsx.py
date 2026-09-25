"""Pruebas de la lista para Brandon (`services/ubicar_contenedores_xlsx.py`).

── QUÉ FIJAN ───────────────────────────────────────────────────────────────────
1. **Las hojas**: Léeme primero y las cinco del spec, cada una con el
   encabezado congelado, filtro y un renglón por caso.
2. **Qué va a cada hoja**: el lote del 12 capturado como 34 es conflicto Y
   error de Odoo (confirmado); «OOLU9155398 - cont 98» igual; Odoo solo contra
   su Ferraforme es refutado; «NNNN-NNNN» es provisional con su costo; un
   contenedor sin documento y un Ferraforme que no se indexó salen en la 5.
3. **Nada de fórmulas**: un nombre que empieza con «=» queda como texto, y el
   libro se vuelve a abrir con openpyxl.

Sin red: la evidencia se arma a mano, como la entregaría `armar_evidencia`.

    cd backend && python -m unittest tests.test_ubicar_contenedores_xlsx -v
"""
from __future__ import annotations

import datetime
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from services import ubicar_contenedores as U  # noqa: E402
from services import ubicar_contenedores_xlsx as X  # noqa: E402

F12 = "Contenedor 12 TIIU6522619.xlsx"


def _ev(sku, fuente, n, **ref):
    return {"sku": sku, "fuente": fuente, "numero": n, "ref": ref}


class Libro(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ev = [_ev(f"FER-{i:02d}", U.FERRA_ALINEADO, 12, archivo=F12, fila=i, alineado=True) for i in range(20)]
        ev += [
            # El lote del 12 capturado como 34 en Odoo.
            _ev("LOTE-1", U.FERRA_ALINEADO, 12, archivo=F12, fila=30, alineado=True),
            _ev("LOTE-1", U.ODOO_CAMPO, 34, texto="EGSU1664119 - 34", via="codigo+numero", activo=True),
            # «OOLU9155398 - cont 98»: el código es del 94.
            _ev("OOLU-1", U.ODOO_AMBIGUO, 94, texto="OOLU9155398 - cont 98", via="conflicto", activo=True,
                codigo_dice=94, texto_dice=98),
            _ev("OOLU-1", U.ODOO_AMBIGUO, 98, texto="OOLU9155398 - cont 98", via="conflicto", activo=True,
                codigo_dice=94, texto_dice=98),
            # Odoo solo, y el Ferraforme del 12 no lo trae.
            _ev("SOLO-ODOO", U.ODOO_CAMPO, 12, texto="Contenedor 12", via="solo_numero", activo=True,
                nombre_odoo="Correas tobillo"),
            # Provisional con costo capturado.
            _ev("4814-0001", U.COSTOS, 91, valor="TRHU6540031 - 91", via="sufijo"),
            # Solo costos en el 97 (sin Ferraforme) y Odoo en el 84 (Ferraforme sin indexar).
            _ev("COSTOS-97", U.COSTOS, 97, valor="XXXU1234567 - 97", via="sufijo"),
            _ev("ODOO-84", U.ODOO_CAMPO, 84, texto="Contenedor 84", via="solo_numero", activo=True),
            # Un pelón contra costos: conflicto, pero NO error de Odoo.
            _ev("PELON-1", U.COSTOS, 97, valor="XXXU1234567 - 97", via="sufijo"),
            _ev("PELON-1", U.ODOO_PELON, 85, texto="MRSU7008648", via="codigo", activo=True),
            # Un contenedor chico con su Ferraforme completo (no sale en la 5) y otro
            # cuyo Ferraforme chico no trae a un SKU que Odoo pone ahí (sí sale).
            _ev("CHICO-1", U.FERRA_NO_ALINEADO, 4, archivo="Contenedor 4.xlsx", fila=1, alineado=False),
            _ev("CHICO-2", U.FERRA_NO_ALINEADO, 4, archivo="Contenedor 4.xlsx", fila=2, alineado=False),
            _ev("CHICO-2", U.COSTOS, 4, valor="WHSU6763243 - 4", via="sufijo"),
            _ev("MEDIO-1", U.FERRA_NO_ALINEADO, 88, archivo="Contenedor 88.xlsx", fila=1, alineado=False),
            _ev("FUERA-88", U.ODOO_CAMPO, 88, texto="Contenedor 88", via="solo_numero", activo=True),
        ]
        skus = {U.clave(e["sku"]) for e in ev} - {"4814-0001"}
        catalogo = {s: s for s in skus}
        nombres = {s: f"Producto {s}" for s in skus}
        nombres["PELON-1"] = "=SUMA(A1)"             # no debe volverse fórmula
        ns_ferra = U.numeros_con_ferraforme(ev)
        res = U.ubicar(ev, catalogo, set(), ns_ferra, {12: "TIIU6522619"})
        archivos = [{"id": 1, "tipo": "ferraforme", "sha256": "f12", "nombre": F12},
                    {"id": 2, "tipo": "ferraforme", "sha256": "f84", "nombre": "Contenedor 84 ABCU7654321.xlsx"}]
        costos = [{"sku": "4814-0001", "contenedor": "TRHU6540031 - 91", "costo_producto": 57,
                   "largo": 10, "ancho": 5, "alto": 2, "peso": 0.4, "piezas_por_caja": 24, "cajas": 3}]
        cls.res = res
        wb, cls.conteos = X.armar_libro(
            excluidos=res["excluidos"], asignaciones=res["asignaciones"], evidencias=ev,
            ns_con_ferraforme=ns_ferra, codigo_por_numero={12: "TIIU6522619"}, archivos=archivos,
            shas_indexados={"f12"}, catalogo=catalogo, nombres=nombres, costos=costos,
            generado=X.fecha_larga(datetime.datetime(2026, 9, 25, 14, 5)), resumen=res["resumen"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        cls.wb = load_workbook(buf)

    def _filas(self, hoja):
        ws = self.wb[hoja]
        cabs = [c.value for c in ws[1]]
        return [dict(zip(cabs, [c.value for c in fila])) for fila in ws.iter_rows(min_row=2)]

    def test_hojas_y_formato(self):
        self.assertEqual(self.wb.sheetnames, ["Léeme", *X.HOJAS])
        for h in X.HOJAS:
            ws = self.wb[h]
            self.assertEqual(ws.freeze_panes, "A2", h)
            self.assertTrue(ws.auto_filter.ref.startswith("A1:"), h)
            self.assertTrue(ws["A1"].font.bold, h)
        leeme = self.wb["Léeme"]
        self.assertIn("25 de septiembre de 2026", leeme["A2"].value)

    def test_conteos(self):
        self.assertEqual(self.conteos, {"1 Conflictos": 3, "2 Refutados": 1, "3 Provisionales": 1,
                                        "4 Errores en Odoo": 2, "5 Contenedores sin documento": 5})
        for h, n in self.conteos.items():
            self.assertEqual(self.wb[h].max_row - 1, n, h)

    def test_conflicto_lote_12_contra_34(self):
        f = {r["SKU"]: r for r in self._filas("1 Conflictos")}
        lote = f["LOTE-1"]
        self.assertEqual((lote["N con documento"], lote["N sin documento"]), ("12", "34"))
        self.assertIn(F12, lote["Ferraforme: archivo o texto"])
        self.assertIn("Los documentos lo ponen en el 12", lote["Qué decidir"])
        self.assertIn("corregir Odoo", lote["Qué decidir"])
        self.assertEqual(lote["También en hoja 4"], "sí")
        self.assertEqual(f["PELON-1"]["También en hoja 4"], None)
        self.assertEqual(f["PELON-1"]["Nombre (catálogo)"], "=SUMA(A1)")
        self.assertIn("Ninguna de sus N (94, 98)", f["OOLU-1"]["Qué decidir"])

    def test_errores_en_odoo(self):
        f = {r["SKU"]: r for r in self._filas("4 Errores en Odoo")}
        self.assertEqual(set(f), {"LOTE-1", "OOLU-1"})
        self.assertEqual(f["LOTE-1"]["Ya confirmado (25-sep)"], "sí")
        self.assertEqual(f["LOTE-1"]["N correcta probable (según)"], "12 (Ferraforme)")
        self.assertEqual(f["OOLU-1"]["N correcta probable (según)"], "94 (lo dice el código)")
        self.assertEqual(f["OOLU-1"]["Ya confirmado (25-sep)"], "sí")

    def test_refutado(self):
        (r,) = self._filas("2 Refutados")
        self.assertEqual((r["SKU"], r["N que da Odoo"], r["SKUs en ese Ferraforme"]), ("SOLO-ODOO", 12, 21))
        self.assertEqual(r["Ferraforme de esa N (archivo)"], F12)
        self.assertEqual(r["Nombre en Odoo"], "Correas tobillo")      # para ver una colisión de SKU

    def test_provisional(self):
        (r,) = self._filas("3 Provisionales")
        self.assertEqual((r["Identificador"], r["Contenedor (N)"], r["Provisionales en ese contenedor"]),
                         ("4814-0001", 91, 1))
        self.assertEqual((r["Costo producto"], r["Piezas por caja"]), (57, 24))

    def test_contenedores_sin_documento(self):
        f = {r["Contenedor (N)"]: r for r in self._filas("5 Contenedores sin documento")}
        # El 34 (lo dice el campo de Odoo) y el 85 (el pelón) también cuentan; el
        # ambiguo 94/98 no nombra un contenedor.
        self.assertEqual(set(f), {34, 84, 85, 88, 97})
        self.assertEqual((f[88]["SKUs reales que lo nombran"], f[88]["…que no están en su Ferraforme"]), (2, 1))
        self.assertTrue(f[88]["Situación"].startswith("Ferraforme con solo 1 SKUs"))
        self.assertEqual(f[84]["Situación"], "Hay Ferraforme en packing_archivos pero no se indexó")
        self.assertEqual(f[97]["SKUs reales que lo nombran"], 2)
        self.assertEqual(f[97]["…cargados en la tabla"], 1)
        self.assertIn("COSTOS-97 — Producto COSTOS-97", f[97]["Ejemplos fuera del Ferraforme (SKU — nombre)"])

    def test_original_sin_n_en_el_nombre_se_asigna_por_codigo(self):
        # «EISU8559654 Lista de empaque.xlsx» es el original del 90: no dice
        # «contenedor 90», pero su código sí está en el mapa.
        ev = [_ev("C-90", U.COSTOS, 90, valor="EISU8559654 - 90", via="sufijo")]
        archivos = [{"id": 108, "tipo": "original", "sha256": "o90", "nombre": "EISU8559654 Lista de empaque.xlsx",
                     "contenedor_base": "EISU8559654"}]
        kw = dict(evidencias=ev, asignaciones=[], ns_con_ferraforme=set(), codigo_por_n={90: "EISU8559654"},
                  archivos=archivos, shas_indexados=set(), catalogo={"C-90": "C-90"}, nombres={}, ferra_por_n={})
        cabs, filas = X._sin_documento(**kw)
        f = dict(zip(cabs, filas[0]))
        self.assertEqual((f["Packing list original en packing_archivos"], f["Situación"]),
                         (0, "Sin Ferraforme ni packing list en kubera"))
        cabs, filas = X._sin_documento(**kw, mapa_codigos={"EISU8559654": {90}})
        f = dict(zip(cabs, filas[0]))
        self.assertEqual((f["Packing list original en packing_archivos"], f["Situación"]),
                         (1, "Solo packing list original, sin Ferraforme"))

    def test_que_decidir_dice_por_que_la_oc_no_cuenta(self):
        e = {"motivo": "conflicto", "ns": [47, 52], "sin_documento": [52],
             "oc_no_documenta": {52: ["el Ferraforme del 52 no lo trae"]},
             "evidencia": {U.FERRA_ALINEADO: {"47": [{}]}, U.OC_RECIBIDA: {"52": [{"oc": "P02543"}]}}}
        t = X.que_decidir(e)
        self.assertIn("la OC no cuenta como documento: el Ferraforme del 52 no lo trae", t)
        self.assertIn("corregir Odoo", t)
        e = {"motivo": "conflicto", "ns": [80, 88], "sin_documento": [80, 88],
             "oc_no_documenta": {88: ["P03488: recibió lo mismo (210) que P03519 (del 80)"]},
             "evidencia": {U.OC_RECIBIDA: {"80": [{}], "88": [{}]}}}
        self.assertIn("La OC del 88 no cuenta como documento: P03488", X.que_decidir(e))

    def test_sin_formulas(self):
        for ws in self.wb.worksheets:
            for fila in ws.iter_rows():
                for c in fila:
                    self.assertNotEqual(c.data_type, "f", f"{ws.title}!{c.coordinate}")


if __name__ == "__main__":
    unittest.main()
