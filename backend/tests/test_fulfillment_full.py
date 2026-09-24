"""Pruebas de la planeación semanal por tienda («Crear FULL»), su IA, su Excel y la ficha del SKU.

── QUÉ FIJAN ───────────────────────────────────────────────────────────────────
  1. «En camino» y «esta semana» por TIENDA (ML Kubera, ML San Corpe, FBA, WFS):
     abierta reciente → lo pedido; validada sin cerrar → lo que ML no avisa;
     cerrada → nada; olvidada (+21 días) → se enseña y no se resta.
  2. Los insumos no confunden un hueco con un cero, las PRUEBAS no se restan y
     el ganador agotado trae reemplazos YA PUBLICADOS con stock (mismo modelo,
     luego misma categoría), como pide el prompt estándar.
  3. Lo libre se reparte entre tiendas y luego por almacén sin prometer dos veces.
  4. Con el interruptor APAGADO nada se escribe; encendido: borrador, precio 0,
     socio fijo por tienda, «PRUEBA» en modo prueba, y la clave no duplica.
  5. La guía sólo se adjunta a órdenes del PANEL que siguen en borrador.
  6. Lo que sugiera la IA se valida: fuera de la planeación o sobre lo libre no pasa.

No se llama a Odoo, kubera ni Mercado Libre: se sustituyen por falsos.

    cd backend && python -m unittest discover -s tests -v
"""
from __future__ import annotations

import sys
import unittest
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from services import fulfillment_excel as fx  # noqa: E402
from services import fulfillment_full as ff  # noqa: E402
from services import fulfillment_ia as fia  # noqa: E402
from services import fulfillment_sku as fs  # noqa: E402

AHORA = datetime(2026, 9, 24, 18, 0, tzinfo=timezone.utc)      # jueves de la S39
TEXCO, TEXCO2 = 135, 150


def _envio(orden="S1", cuenta="Kubera", hecha=True, dias=3, lineas=(("A", 10, 10, 0),),
           cerrado=False, avisos=True, canal="meli"):
    """lineas: (sku, pedidas, enviadas, llegadas)."""
    creada = AHORA - timedelta(days=dias + 1)
    salida = AHORA - timedelta(days=dias) if hecha else None
    e = {"canal": canal, "cuenta": cuenta, "orden": orden, "salida": f"TEXCO/OUT/{orden}",
         "estado_odoo": "done" if hecha else "assigned", "pedidas": sum(l[1] for l in lineas),
         "piezas": sum(l[2] for l in lineas) if hecha else None,
         "etapas": [{"ts": creada.isoformat()}, {"ts": salida.isoformat()} if salida else None, None, None, None],
         "lineas": [{"sku": s, "pedidas": p, "enviadas": en if hecha else None, "llegadas": ll}
                    for s, p, en, ll in lineas]}
    if hecha and avisos:
        e["cobertura"] = {"fuente": "avisos", "cerrado": cerrado}
    return e


class EnCamino(unittest.TestCase):
    def test_salida_abierta_reciente_cuenta_lo_pedido(self):
        camino, zombis = ff.en_camino([_envio(hecha=False, dias=2)], AHORA)
        self.assertEqual(camino[("meli:Kubera", "A")]["piezas"], 10)
        self.assertEqual(zombis, [])

    def test_orden_olvidada_no_se_resta(self):
        camino, zombis = ff.en_camino([_envio(orden="S26441", hecha=False, dias=40)], AHORA)
        self.assertEqual(camino, {})
        self.assertEqual((zombis[0]["orden"], zombis[0]["tienda"]), ("S26441", "meli:Kubera"))

    def test_validada_sin_cerrar_cuenta_lo_que_falta_por_llegar(self):
        camino, _ = ff.en_camino([_envio(lineas=(("A", 10, 10, 4), ("B", 5, 5, 5)))], AHORA)
        self.assertEqual(camino[("meli:Kubera", "A")]["piezas"], 6)
        self.assertNotIn(("meli:Kubera", "B"), camino, "B ya llegó completo")

    def test_cerrada_no_esta_en_camino(self):
        camino, _ = ff.en_camino([_envio(lineas=(("A", 10, 10, 4),), cerrado=True, dias=15)], AHORA)
        self.assertEqual(camino, {}, "lo que falta en un envío cerrado es lo que ML no recibió")

    def test_fba_sin_avisos_toma_lo_enviado(self):
        camino, _ = ff.en_camino([_envio(canal="amazon", cuenta="San Corpe", avisos=False, dias=5)], AHORA)
        self.assertEqual(camino[("amazon", "A")]["piezas"], 10)

    def test_sin_cuenta_no_entra(self):
        self.assertEqual(ff.en_camino([_envio(cuenta=None, hecha=False, dias=1)], AHORA), ({}, []))


class EstaSemana(unittest.TestCase):
    def test_salieron_esta_semana_mas_por_validar(self):
        salio = _envio(orden="S1", dias=1, lineas=(("A", 10, 10, 0), ("B", 5, 5, 0)))     # miércoles S39
        vieja = _envio(orden="S0", dias=8, lineas=(("C", 7, 7, 0),))                        # S38: no cuenta
        abierta = _envio(orden="S2", hecha=False, dias=0, lineas=(("A", 4, 0, 0), ("D", 3, 0, 0)))
        s = ff.esta_semana([salio, vieja, abierta], AHORA)["meli:Kubera"]
        self.assertEqual((s["envios"], s["piezas"], s["skus"]), (2, 22, 3))
        self.assertEqual(s["salieron"], {"envios": 1, "piezas": 15})
        self.assertEqual(s["por_validar"], {"envios": 1, "piezas": 7})


class Utilidades(unittest.TestCase):
    def test_modelo_base(self):
        self.assertEqual(ff.modelo_base("TEC-0393-ROS"), "TEC-0393")
        self.assertEqual(ff.modelo_base("ACC-0696-ROJ-NEG-140CM"), "ACC-0696")

    def test_parecido_detecta_reciclados(self):
        self.assertLess(ff.parecido("Binoculares 10x50 profesionales", "Faros de niebla LED para auto"), 0.15)
        self.assertGreater(ff.parecido("Audífonos inalámbricos BT 5.3 invisibles",
                                       "AUDIFONOS INVISIBLES VERDADERAMENTE CON BT 5.3"), 0.3)
        self.assertIsNone(ff.parecido("X", "Algo largo aquí"), "con una palabra no se juzga")

    def test_medidas_sospechosas_del_prompt(self):
        self.assertTrue(ff.medidas_sospechosas({"largo": 41, "ancho": 60, "alto": 40, "peso": 0.3}))
        self.assertFalse(ff.medidas_sospechosas({"largo": 41, "ancho": 60, "alto": 40, "peso": 2}))
        self.assertFalse(ff.medidas_sospechosas({"largo": 20, "ancho": 10, "alto": 5, "peso": 0.2}))
        self.assertFalse(ff.medidas_sospechosas(None))


class Propuesta(unittest.TestCase):
    def setUp(self):
        ventas = [{"canal": "mercado_libre", "cuenta": "BEKURA", "sku": "TEC-0001-ROS", "vv": 60, "v7": 20,
                   "ultima": date(2026, 9, 23)},
                  {"canal": "mercado_libre", "cuenta": "BEKURA", "sku": "TEC-0002-NEG", "vv": 9, "v7": 1, "ultima": None},
                  {"canal": "mercado_libre", "cuenta": "SANCORFASHION", "sku": "TEC-0001-ROS", "vv": 30, "v7": 7,
                   "ultima": None},
                  {"canal": "amazon", "cuenta": "AMAZON", "sku": "VIA-0024-NEG", "vv": 12, "v7": 3, "ultima": None}]
        pubs = {"meli:Kubera": {
                    "TEC-0001-ROS": {"sku": "TEC-0001-ROS", "listing_id": "MLM1", "url": "u", "situacion": "paused",
                                     "en_almacen": True, "categoria": "MLM10", "stock": 12},
                    "TEC-0002-NEG": {"sku": "TEC-0002-NEG", "listing_id": "MLM2", "url": "u", "situacion": "paused",
                                     "en_almacen": True, "categoria": "MLM20", "stock": 0},
                    "TEC-0002-AZL": {"sku": "TEC-0002-AZL", "listing_id": "MLM3", "url": "u", "situacion": "active",
                                     "en_almacen": True, "categoria": "MLM20", "stock": 4},
                    "HOG-0100-BLN": {"sku": "HOG-0100-BLN", "listing_id": "MLM4", "url": "u", "situacion": "active",
                                     "en_almacen": False, "categoria": "MLM20", "stock": 0}},
                "meli:San Corpe": {}, "amazon": {}, "walmart": {}}
        vivos = {"meli:Kubera": {"MLM1": {"estado": "paused", "stock": 0, "logistica": "fulfillment",
                                          "titulo": "Set de brochas rosa", "categoria": "MLM10"}}}
        borradores = [
            {"orden": "S38878", "tienda": "meli:Kubera", "prueba": False, "lineas": [{"sku": "TEC-0001-ROS", "cantidad": 20}]},
            {"orden": "S39000", "tienda": "meli:Kubera", "prueba": True, "lineas": [{"sku": "TEC-0001-ROS", "cantidad": 99}]},
        ]
        productos = {"TEC-0001-ROS": {"id": 1, "name": "Brochas"}, "TEC-0002-NEG": {"id": 2, "name": "Tablet negra"},
                     "TEC-0002-AZL": {"id": 3, "name": "Tablet azul"}, "HOG-0100-BLN": {"id": 4, "name": "Lámpara"}}
        libres = {1: {TEXCO: 30, TEXCO2: -4}, 2: {TEXCO: 0, TEXCO2: 0}, 3: {TEXCO: 0, TEXCO2: 8}, 4: {TEXCO: 50}}
        nombres = {"TEC-0001-ROS": "SET DE BROCHAS ROSA"}
        self.p = ff.armar_propuesta(list(ff.TIENDAS), ventas, pubs, vivos, [_envio(hecha=False, dias=1,
                                    lineas=(("TEC-0001-ROS", 10, 0, 0),))], borradores, productos, libres,
                                    nombres, {}, AHORA, 30)

    def fila(self, tienda, sku):
        return next(f for f in self.p["tiendas"][tienda]["filas"] if f["sku"] == sku)

    def test_insumos_con_ml_en_vivo(self):
        a = self.fila("meli:Kubera", "TEC-0001-ROS")
        self.assertEqual(a["stock"], 0, "ML en vivo (0) manda sobre la copia del sync (12)")
        self.assertTrue(a["verificada"])
        self.assertEqual((a["vv"], a["en_camino"], a["borrador"]), (60, 10, 20), "la prueba (99) no se resta")
        self.assertEqual(a["libre"], {"TEXCO": 30, "TEXCO II": 0}, "libre negativo es cero para planear")
        self.assertEqual(a["nombre"], "SET DE BROCHAS ROSA", "el nombre de Omnicanal, no el de Odoo")

    def test_ganador_agotado_trae_reemplazos_publicados(self):
        g = self.fila("meli:Kubera", "TEC-0002-NEG")
        self.assertEqual([(r["sku"], r["tipo"]) for r in g["reemplazos"]],
                         [("TEC-0002-AZL", "mismo modelo"), ("HOG-0100-BLN", "misma categoría")])

    def test_cada_tienda_con_sus_skus(self):
        self.assertEqual([f["sku"] for f in self.p["tiendas"]["meli:San Corpe"]["filas"]], ["TEC-0001-ROS"])
        fba = self.fila("amazon", "VIA-0024-NEG")
        self.assertIsNone(fba["stock"], "sin publicación FBA registrada no se sabe (no es 0)")
        self.assertIsNone(fba["libre"], "sin producto en Odoo no se sabe")
        self.assertEqual(self.p["tiendas"]["walmart"]["filas"], [])

    def test_semana_y_ventana(self):
        self.assertEqual((self.p["semana"]["semana"], self.p["semana"]["lunes"]), ("S39", "2026-09-21"))
        self.assertEqual((self.p["ventana"]["dias"], self.p["ventana"]["hasta"]), (30, "2026-09-23"))


class Almacenes(unittest.TestCase):
    def _l(self, sku, pid, n):
        return {"sku": sku, "product_id": pid, "cantidad": n}

    def test_un_almacen_que_cubre_todo_gana(self):
        r = ff.repartir_almacenes([self._l("A", 1, 5), self._l("B", 2, 5)],
                                  {1: {TEXCO: 9, TEXCO2: 50}, 2: {TEXCO: 5, TEXCO2: 0}})
        self.assertEqual([p["almacen"] for p in r["partes"]], ["TEXCO"])

    def test_cada_renglon_entero_donde_alcance(self):
        r = ff.repartir_almacenes([self._l("A", 1, 10), self._l("B", 2, 5)],
                                  {1: {TEXCO: 4, TEXCO2: 20}, 2: {TEXCO: 5, TEXCO2: 0}})
        partes = {p["almacen"]: {l["sku"]: l["cantidad"] for l in p["lineas"]} for p in r["partes"]}
        self.assertEqual(partes, {"TEXCO": {"B": 5}, "TEXCO II": {"A": 10}})

    def test_lo_que_no_hay_se_recorta_y_se_dice(self):
        r = ff.repartir_almacenes([self._l("A", 1, 10), self._l("B", 2, 3)],
                                  {1: {TEXCO: 4, TEXCO2: 2}, 2: {TEXCO: 0, TEXCO2: 0}})
        self.assertEqual([(x["sku"], x["van"], x["porque"]) for x in r["recortes"]],
                         [("A", 6, "sólo hay 6 libres en Odoo"), ("B", 0, "Odoo no tiene libre")])

    def test_entre_tiendas_se_reparte_en_proporcion(self):
        pedidos = {"meli:Kubera": [self._l("A", 1, 30)], "meli:San Corpe": [self._l("A", 1, 10)]}
        ajustados, recortes = ff.repartir_entre_tiendas(pedidos, {1: {TEXCO: 20, TEXCO2: 0}})
        self.assertEqual((ajustados["meli:Kubera"][0]["cantidad"], ajustados["meli:San Corpe"][0]["cantidad"]), (15, 5))
        self.assertEqual(len(recortes), 2)


class Limpiar(unittest.TestCase):
    def test_un_renglon_por_sku_y_solo_tiendas_conocidas(self):
        r = ff.limpiar_tiendas([{"tienda": "meli:Kubera", "lineas": [{"sku": " A ", "cantidad": 3, "sugerido": 5},
                                                                     {"sku": "A", "cantidad": 2},
                                                                     {"sku": "B", "cantidad": 0}]},
                                {"tienda": "temu", "lineas": [{"sku": "X", "cantidad": 9}]}])
        self.assertEqual(r, {"meli:Kubera": [{"sku": "A", "cantidad": 5, "sugerido": 5}]},
                         "Temu es DROP: no es tienda de la planeación")


class Buscar(unittest.TestCase):
    """Buscar para agregar: sólo lo PUBLICADO en la tienda; lo que falta se le pregunta a ML."""

    def _buscar(self, texto):
        pubs = [{"codigo": "BEKURA", "sku": s, "listing_id": f"MLM{i}", "url": None, "status": "active",
                 "situacion": "active", "en_almacen": True, "logistic_type": "fulfillment", "category_id": "MLM1",
                 "product_type": None, "stock_full": 3, "stock_fba": 0}
                for i, s in enumerate(["JUGU-0100-ROJ", "JUGU-0100-AZL", "TEC-0001-NEG"])]
        preguntados: list[str] = []

        def fetch_all(sql, params=None):
            return [dict(p) for p in pubs] if sql is ff._SQL_PUBLICACIONES else []

        def en_ml(codigo, sku):
            preguntados.append(sku)
            return []

        with mock.patch.object(ff.sdb, "fetch_all", fetch_all),              mock.patch.object(ff, "verificar_ml", lambda c, ids: {}),              mock.patch.object(ff, "buscar_ml_por_sku", en_ml),              mock.patch.object(ff.odoo_ventas, "productos_por_sku", lambda skus: {}),              mock.patch.object(ff.odoo_ventas, "libre_por_almacen", lambda ids: {}):
            return ff.buscar("meli:Kubera", texto, [], 30, AHORA), preguntados

    def test_varios_skus_aceptan_la_base_de_sus_variantes(self):
        r, preguntados = self._buscar("JUGU-0100, TEC-0001-NEG, NO-EXISTE")
        self.assertEqual([f["sku"] for f in r["filas"]], ["JUGU-0100-AZL", "JUGU-0100-ROJ", "TEC-0001-NEG"],
                         "pegar la base trae sus variantes publicadas; antes la daba por «no publicada»")
        self.assertEqual(r["no_publicados"], ["NO-EXISTE"])
        self.assertEqual(preguntados, ["NO-EXISTE"], "sólo lo que no está en la copia se le pregunta a ML")

    def test_una_tienda_que_no_existe_no_busca(self):
        r = ff.buscar("temu", "JUGU", [], 30, AHORA)
        self.assertFalse(r["ok"], "Temu es DROP: no es tienda de la planeación")


class OdooFalso:
    """Anota cada llamada. Sólo contesta lo que `crear` y `adjuntar_guia` necesitan."""

    def __init__(self, existentes=None, socio=None, orden=None):
        self.llamadas: list[tuple] = []
        self.existentes = existentes or []
        self.socio = socio
        self.orden = orden or {}
        self.siguiente = 900

    def __call__(self, modelo, metodo, args, kwargs=None):
        self.llamadas.append((modelo, metodo, args, kwargs or {}))
        if (modelo, metodo) == ("product.product", "search_read"):
            return [{"id": 1, "default_code": "A", "name": "Producto A"},
                    {"id": 2, "default_code": "B", "name": "Producto B"}]
        if (modelo, metodo) == ("product.product", "read"):
            wid = (kwargs or {}).get("context", {}).get("warehouse")
            return [{"id": 1, "free_qty": 50 if wid == TEXCO else 0},
                    {"id": 2, "free_qty": 0 if wid == TEXCO else 8}]
        if (modelo, metodo) == ("sale.order", "search_read"):
            return self.existentes
        if (modelo, metodo) == ("res.partner", "search_read"):
            return [{"id": self.socio}] if self.socio else []
        if (modelo, metodo) == ("res.partner", "create"):
            return 777
        if (modelo, metodo) == ("sale.order", "create"):
            self.siguiente += 1
            return self.siguiente
        if (modelo, metodo) == ("sale.order", "read"):
            oid = args[0][0]
            if self.orden:
                return [{**self.orden, "id": oid, "meli_etiqueta_file": "12 KB",
                         "meli_etiqueta_filename": self.orden.get("_nombre_pdf")}]
            return [{"id": oid, "name": f"S{oid}", "state": "draft", "picking_ids": [], "order_line": [1]}]
        if (modelo, metodo) == ("sale.order", "write"):
            self.orden.update({k: v for k, v in args[1].items() if k == "client_order_ref"})
            if "meli_etiqueta_filename" in args[1]:
                self.orden["_nombre_pdf"] = args[1]["meli_etiqueta_filename"]
            return True
        raise AssertionError(f"llamada no esperada a Odoo: {modelo}.{metodo}")

    def escrituras(self):
        return [(m, x) for m, x, *_ in self.llamadas if x not in ("search_read", "read", "search", "search_count")]


PEDIDO = [{"tienda": "meli:Kubera", "lineas": [{"sku": "A", "cantidad": 10, "sugerido": 12},
                                               {"sku": "B", "cantidad": 5, "sugerido": 5}]}]


class Crear(unittest.TestCase):
    def _crear(self, odoo, encendido, clave="a1b2c3d4e5", prueba=True):
        publicadas = [{"sku": "A", "listing_id": "MLM1"}, {"sku": "B", "listing_id": "MLM2"}]
        with mock.patch.object(ff.odoo_ventas, "_kw", odoo), \
             mock.patch.object(ff, "habilitado", lambda refrescar=False: encendido), \
             mock.patch.object(ff, "_guardar_solicitud", lambda *a, **k: True), \
             mock.patch.object(ff.sdb, "fetch_all", lambda *a, **k: publicadas), \
             mock.patch.object(ff, "verificar_ml", lambda c, ids: {i: {"estado": "active"} for i in ids}), \
             mock.patch.object(ff.odoo_ventas, "url_orden_publica", lambda: "https://odoo/{id}"):
            return ff.crear(PEDIDO, "brandon@kubera.mx", clave, {"cobertura_dias": 30}, prueba)

    def test_apagado_no_escribe_nada(self):
        odoo = OdooFalso()
        r = self._crear(odoo, encendido=False)
        self.assertEqual(r["accion"], "apagado")
        self.assertEqual(odoo.escrituras(), [], "con el interruptor apagado sólo se lee")
        self.assertEqual({p["almacen"]: p["piezas"] for p in r["tiendas"][0]["partes"]},
                         {"TEXCO": 10, "TEXCO II": 5})

    def test_encendido_crea_borradores_de_prueba_uno_por_almacen(self):
        odoo = OdooFalso()
        r = self._crear(odoo, encendido=True)
        self.assertEqual((r["accion"], r["ok"]), ("creada", True))
        self.assertEqual(odoo.escrituras(), [("res.partner", "create"), ("sale.order", "create"),
                                             ("sale.order", "create")], "nunca action_confirm")
        v = next(x for x in odoo.llamadas if x[:2] == ("sale.order", "create"))[2][0]
        self.assertEqual(v["partner_id"], 777, "el socio fijo de la tienda")
        self.assertIn("· a1b2c3d4e5 ·", v["origin"])
        self.assertTrue(v["origin"].endswith("· PRUEBA"))
        self.assertEqual(v["client_order_ref"], ff.PRUEBA_REF)
        renglon = v["order_line"][0][2]
        self.assertEqual((renglon["price_unit"], renglon["tax_id"]), (0.0, [(6, 0, [])]))
        self.assertEqual([o["orden"] for o in r["tiendas"][0]["ordenes"]], ["S901", "S902"])

    def test_real_no_lleva_prueba(self):
        odoo = OdooFalso()
        self._crear(odoo, encendido=True, prueba=False)
        v = next(x for x in odoo.llamadas if x[:2] == ("sale.order", "create"))[2][0]
        self.assertNotIn("client_order_ref", v)
        self.assertNotIn("PRUEBA", v["origin"])

    def test_la_misma_clave_no_crea_dos_veces(self):
        existentes = [{"id": 5, "name": "S5", "state": "draft", "warehouse_id": [TEXCO, "TEXCO"], "partner_id": [321, "x"]},
                      {"id": 6, "name": "S6", "state": "draft", "warehouse_id": [TEXCO2, "TEXCO II"], "partner_id": [321, "x"]}]
        odoo = OdooFalso(existentes=existentes, socio=321)
        r = self._crear(odoo, encendido=True)
        self.assertEqual(r["accion"], "ya_existia")
        self.assertEqual(odoo.escrituras(), [])

    def test_sin_clave_no_hace_nada(self):
        odoo = OdooFalso()
        self.assertEqual(self._crear(odoo, encendido=True, clave="x")["accion"], "sin_clave")
        self.assertEqual(odoo.llamadas, [])


class Guia(unittest.TestCase):
    PDF = b"%PDF-1.4 guia"

    def _adjuntar(self, orden, encendido=True, numero="76309173", pdf=PDF):
        odoo = OdooFalso(orden=orden)
        with mock.patch.object(ff.odoo_ventas, "_kw", odoo), \
             mock.patch.object(ff, "habilitado", lambda refrescar=False: encendido):
            return ff.adjuntar_guia(10, numero, pdf, "etiquetas.pdf", "brandon@kubera.mx"), odoo

    def test_orden_del_panel_en_borrador(self):
        orden = {"name": "S10", "state": "draft", "origin": "Panel FULLFILMENT · ML Kubera · b · c · TEXCO",
                 "client_order_ref": False}
        r, odoo = self._adjuntar(orden)
        self.assertTrue(r["ok"], r)
        w = next(x for x in odoo.llamadas if x[:2] == ("sale.order", "write"))[2][1]
        self.assertEqual((w["client_order_ref"], w["meli_etiqueta_filename"]), ("76309173", "etiquetas.pdf"))

    def test_en_prueba_el_numero_conserva_la_marca(self):
        orden = {"name": "S10", "state": "draft", "origin": "Panel FULLFILMENT · ML Kubera · b · c · TEXCO · PRUEBA",
                 "client_order_ref": ff.PRUEBA_REF}
        r, odoo = self._adjuntar(orden)
        w = next(x for x in odoo.llamadas if x[:2] == ("sale.order", "write"))[2][1]
        self.assertEqual(w["client_order_ref"], "PRUEBA · 76309173")

    def test_ajena_confirmada_apagado_o_sin_pdf_no_se_toca(self):
        ajena = {"name": "S11", "state": "draft", "origin": "S11", "client_order_ref": False}
        confirmada = {"name": "S12", "state": "sale", "origin": "Panel FULLFILMENT · x", "client_order_ref": False}
        for orden, accion in ((ajena, "ajena"), (confirmada, "confirmada")):
            r, odoo = self._adjuntar(orden)
            self.assertEqual(r["accion"], accion)
            self.assertEqual(odoo.escrituras(), [])
        r, odoo = self._adjuntar(ajena, encendido=False)
        self.assertEqual((r["accion"], odoo.llamadas), ("apagado", []))
        r, odoo = self._adjuntar(ajena, pdf=b"no soy pdf")
        self.assertEqual((r["accion"], odoo.llamadas), ("sin_pdf", []))


class ValidarIA(unittest.TestCase):
    DATOS = {"tiendas": {"meli:Kubera": {
        "renglones": [{"sku": "A", "libre": 10}, {"sku": "B", "libre": 0}],
        "ganadores_agotados": [{"sku": "B", "candidatos": [{"sku": "C", "libre": 7}]}]}}}

    def test_lo_que_no_esta_en_la_planeacion_no_pasa(self):
        r = fia.validar({"ajustes": [{"tienda": "meli:Kubera", "sku": "A", "cantidad": 25, "motivo": "sube"},
                                     {"tienda": "meli:Kubera", "sku": "Z", "cantidad": 3, "motivo": "?"},
                                     {"tienda": "temu", "sku": "A", "cantidad": 3, "motivo": "?"},
                                     {"tienda": "meli:Kubera", "sku": "C", "cantidad": 4, "motivo": "reemplazo"}],
                         "reemplazos": [{"tienda": "meli:Kubera", "agotado": "B", "reemplazo": "C",
                                         "tipo_match": "mismo modelo", "motivo": "ok"},
                                        {"tienda": "meli:Kubera", "agotado": "B", "reemplazo": "Q",
                                         "tipo_match": "misma por nombre", "motivo": "inventado"}],
                         "alertas": [], "confirmacion": "c", "resumen": "r"}, self.DATOS)
        self.assertEqual([(a["sku"], a["cantidad"]) for a in r["ajustes"]], [("A", 10), ("C", 4)],
                         "A se topa a lo libre; C es un candidato válido")
        self.assertIn("se topó", r["ajustes"][0]["nota"])
        self.assertEqual([x["reemplazo"] for x in r["reemplazos"]], ["C"])
        self.assertEqual(len(r["descartados"]), 3)

    def test_la_tienda_por_su_nombre_tambien_vale(self):
        datos = {"tiendas": {"meli:Kubera": {"nombre": "ML Kubera", "destino": "meli_bekura",
                                             **self.DATOS["tiendas"]["meli:Kubera"]}}}
        r = fia.validar({"ajustes": [{"tienda": "ML Kubera", "sku": "A", "cantidad": 4, "motivo": "sube"},
                                     {"tienda": "BEKURA", "sku": "A", "cantidad": 5, "motivo": "sube"}],
                         "reemplazos": [{"tienda": "meli_bekura", "agotado": "B", "reemplazo": "C",
                                         "tipo_match": "mismo modelo", "motivo": "ok"}],
                         "alertas": [], "confirmacion": "c", "resumen": "r"}, datos)
        self.assertEqual([(a["tienda"], a["cantidad"]) for a in r["ajustes"]], [("meli:Kubera", 4), ("meli:Kubera", 5)],
                         "la 1ª corrida real descartó todo por escribir «ML Kubera»")
        self.assertEqual([x["tienda"] for x in r["reemplazos"]], ["meli:Kubera"])
        self.assertEqual(r["descartados"], [])

    def test_el_esquema_exige_la_llave_de_la_tienda(self):
        ajuste = fia.ESQUEMA["properties"]["ajustes"]["items"]["properties"]["tienda"]
        self.assertEqual(ajuste["enum"], ["meli:Kubera", "meli:San Corpe", "amazon", "walmart"])

    def test_sin_clave_no_arranca(self):
        with mock.patch.object(fia.settings, "anthropic_api_key", ""):
            self.assertFalse(fia.iniciar(self.DATOS)["ok"])


class Excel(unittest.TestCase):
    def test_arma_un_xlsx_con_hoja_por_tienda(self):
        from io import BytesIO
        from openpyxl import load_workbook
        contenido = fx.armar({"semana": "S39", "tiendas": [
            {"tienda": "meli:Kubera", "nombre": "ML Kubera", "totales": {"renglones": 1},
             "renglones": [{"sku": "A", "nombre": "Á con acento", "a_mandar": 5}]}],
            "ganadores": [{"tienda": "ML Kubera", "sku": "B", "candidatos": [{"sku": "C", "tipo": "mismo modelo"}]}],
            "alertas": [{"tienda": "ML Kubera", "sku": "A", "tipo": "reciclado", "detalle": "x"}]})
        wb = load_workbook(BytesIO(contenido))
        self.assertEqual(wb.sheetnames, ["Resumen", "ML Kubera", "Ganadores agotados", "SKUs", "Alertas"])
        self.assertEqual(wb["ML Kubera"]["B2"].value, "Á con acento")
        self.assertEqual(wb["SKUs"]["B2"].value, "A")


class FichaSku(unittest.TestCase):
    def test_semanas_de_venta_seguidas_con_ceros(self):
        hoy = date(2026, 9, 24)
        filas = [{"cuenta": "BEKURA", "date": date(2026, 9, 22), "unidades": 3},
                 {"cuenta": "BEKURA", "date": date(2026, 9, 23), "unidades": 2},
                 {"cuenta": "BEKURA", "date": date(2026, 9, 9), "unidades": 4}]
        s = fs.semanas_de_venta(filas, hoy, n=3)
        self.assertEqual([(x["semana"], x["unidades"]) for x in s["Kubera"]], [("S37", 4), ("S38", 0), ("S39", 5)])

    def test_envios_del_sku(self):
        e1 = _envio(orden="S1", lineas=(("A", 10, 10, 10), ("B", 5, 5, 0)))
        e2 = _envio(orden="S2", dias=1, lineas=(("A", 3, 3, 0),))
        self.assertEqual([x["orden"] for x in fs.envios_del_sku("A", [e1, e2])], ["S2", "S1"])


if __name__ == "__main__":
    unittest.main()
