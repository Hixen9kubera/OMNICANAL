"""
excel_pl.py — El inventario de la página, en un Excel con la FOTO de cada producto.

Una fila por cada fila de la página (los SKUs y los productos de packing list sin SKU), con lo que
se pidió y nada más: foto, SKU, título, categoría, subcategoría, piezas (comprado, salió, queda y
lo que Odoo tiene libre) y la media de Mercado Libre, con su leyenda.

No pregunta nada a ningún sistema: lee `datos_pl.json` y las hojas de fotos grandes que ya armaron
`fotos_grandes` y `pagina_pl` (`img/gNNN.jpg` y `img/qNNN.jpg`), de donde recorta la foto de cada fila.

Tres cosas que no son obvias:
  · La foto va ANCLADA a su celda («mover y cambiar tamaño con las celdas»). Así, al filtrar por
    categoría, la foto de una fila oculta se esconde con ella en vez de quedar encimada.
  · openpyxl guarda una copia de la imagen por cada fila, aunque veinte tallas compartan la misma
    foto. Al final se reescribe el .xlsx para que las fotos repetidas apunten a un solo archivo:
    pesa casi la mitad.
  · Tiene que cuadrar al 100% con la página. La media es EXACTAMENTE la que la página pone en su
    columna «media» (incluido lo estimado, que lleva su leyenda), y al terminar el Excel se vuelve a
    leer y se compara fila por fila contra `datos_pl.json`: si algo no coincide, la etapa falla.
"""
from __future__ import annotations

import hashlib
import io
import re
import zipfile
from pathlib import Path
from typing import Any

from comun import aviso, leer_json

LADO = 160                       # píxeles de la foto que se guarda (se muestra a 100)
CAJA = 100                       # píxeles a los que se ve en la hoja
ANCHO_COL = 15.0                 # ancho de la columna de la foto (≈ 110 px)
ALTO_FILA = 80.0                 # puntos (≈ 107 px)


def _recortes(salida: Path, filas: list[dict[str, Any]], doc: dict[str, Any]) -> dict[tuple[str, int], Path]:
    """La foto de cada fila, recortada de su hoja grande. Una por foto distinta (g = Odoo, q = packing list)."""
    from PIL import Image

    destino = salida / "cache" / "excel_img"
    destino.mkdir(parents=True, exist_ok=True)
    quiero: dict[str, set[int]] = {"g": set(), "q": set()}
    for r in filas:
        if r.get("i") is not None:
            quiero["g"].add(r["i"])
        elif r.get("j") is not None:
            quiero["q"].add(r["j"])
    rutas: dict[tuple[str, int], Path] = {}
    for pre, sprite in (("g", doc.get("spriteG") or {}), ("q", doc.get("spriteQ") or {})):
        if not sprite or not sprite.get("hojas"):
            continue
        lado, cols, por = sprite["lado"], sprite["cols"], sprite["por_hoja"]
        por_hoja: dict[int, list[int]] = {}
        for i in quiero[pre]:
            por_hoja.setdefault(i // por, []).append(i)
        for h, indices in sorted(por_hoja.items()):
            ruta_hoja = salida / "img" / f"{pre}{h:03d}.jpg"
            if not ruta_hoja.exists():
                continue
            hoja = None
            for i in indices:
                ruta = destino / f"{pre}_{i}.jpg"
                rutas[(pre, i)] = ruta
                if ruta.exists() and ruta.stat().st_mtime >= ruta_hoja.stat().st_mtime:
                    continue
                hoja = hoja or Image.open(ruta_hoja).convert("RGB")
                n = i % por
                x, y = (n % cols) * lado, (n // cols) * lado
                hoja.crop((x, y, x + lado, y + lado)).resize((LADO, LADO), Image.LANCZOS).save(
                    ruta, "JPEG", quality=82, optimize=True)
    return rutas


CABECERA = ["Foto", "SKU", "Título", "Categoría", "Subcategoría", "Comprado", "Salió", "Queda", "Odoo libre",
            "Media de Mercado Libre", "De dónde sale la media", "Valor a precio de ML", "Estado"]


def media_de_la_pagina(r: dict[str, Any]) -> tuple[float | None, str]:
    """Lo que la página muestra en la columna «media» de «Mercado Libre por pieza», y su leyenda.

    Es el mismo orden de `celdaPrecio` en la plantilla: (1) precio del archivo de Eduardo: la media de sus
    publicaciones, o el precio usado si no trae publicaciones; (2) medido por Competencia del panel;
    (3) estimado; (4) nada."""
    p = r.get("ml")
    if p and p.get("o") == "e":
        valor = p.get("me") if p.get("lo") is not None else p.get("m")
        cl = p.get("cl")
        de = ("revisado a mano" if cl == "r" else
              "mismo producto" + (" (del modelo)" if p.get("nv") == "x" else "") if cl == "f" else
              "banda de su categoría" + (f" «{p['cg']}»" if p.get("cg") else "") if cl == "b" else
              "sin publicaciones: nuestro precio")
        # Cuando bodega contó distinto que el proveedor, la página topa el precio que multiplica: se dice cuál.
        return valor, de + (" · sin revisar" if p.get("sr") else "") + (
            f" · topado a 10× FOB: se multiplica por ${p['m']:,.2f}" if p.get("tp") else "")
    firme = bool(p and not r.get("sin") and p.get("n", 0) >= 2 and not p.get("x") and p.get("o") != "c")
    if firme:
        return p.get("m"), "mismo producto (Competencia del panel)"
    if r.get("em"):
        return r["em"][0], "estimado"
    return None, ("insumo: no se valúa" if r.get("nv") else "")


def fila_de_la_pagina(r: dict[str, Any], rutas_cat: list[list[str]]) -> tuple:
    """Los datos de una fila tal como los muestra la página (sin la foto). Lo usa el Excel y su comprobación."""
    k = rutas_cat[r["k"]] if r.get("k") is not None else ["Sin categoría", "", ""]
    con_pl = r.get("pc") is not None
    valor, de = media_de_la_pagina(r)
    # Igual que la tabla de la página: sin packing list, «comprado» es una raya, o «≈ N según costos» si la
    # base de costos lo manda a un contenedor (esas piezas no suman: ya están entre los productos sin SKU).
    comprado = r["pc"] if con_pl else (f"≈ {round(r['pe']):,} según costos" if r.get("pe") is not None else None)
    estado = ", ".join(x for x in (
        "sin SKU" if r.get("sin") else "",
        "insumo, no se valúa" if r.get("nv") else "",
        "SOLD OUT" if r.get("so") else "",
        "renglón no ubicado" if (not con_pl and r.get("pe") is not None) else "",
        "sin packing list" if (not con_pl and r.get("pe") is None) else "") if x)
    # El valor de la fila, como en la página: piezas que quedan × el precio que usa (la media, o el topado, o
    # el estimado). Los insumos y lo que no tiene packing list no se valúan.
    p = r.get("ml")
    con_precio = bool(p and (p.get("o") == "e" or (not r.get("sin") and p.get("n", 0) >= 2 and not p.get("x") and p.get("o") != "c")))
    usado = p["m"] if con_precio else (r["em"][0] if r.get("em") else None)
    total = round(r["pq"] * usado, 2) if (con_pl and usado is not None) else None
    return ("Sin SKU" if r.get("sin") else r["s"], r.get("t") or "", k[0], k[1] or "",
            comprado,
            (r.get("pz") or r.get("ps")) if r.get("ps") else 0,
            r.get("pq") if con_pl else None,
            r.get("lb") if not r.get("sin") else None,
            valor, de, total, estado)


def comprobar(ruta: Path, doc: dict[str, Any]) -> dict[str, Any]:
    """Vuelve a leer el Excel ya escrito y lo compara, fila por fila, contra los datos de la página.
    Deben coincidir al 100%: mismos SKUs, mismos títulos, misma categoría, mismas piezas, misma media."""
    from collections import Counter

    from openpyxl import load_workbook

    def norma(t: tuple) -> tuple:
        return tuple(round(float(x), 2) if isinstance(x, (int, float)) and not isinstance(x, bool) else (x if x not in ("", None) else None)
                     for x in t)

    esperado = Counter(norma(fila_de_la_pagina(r, doc["K"])) for r in doc["R"])
    hoja = load_workbook(ruta, read_only=True)["Productos"]
    leido: Counter = Counter()
    n = 0
    for fila in hoja.iter_rows(min_row=2, min_col=2, max_col=len(CABECERA), values_only=True):
        if fila[0] is None:
            continue
        leido[norma(fila)] += 1
        n += 1
    sobran, faltan = leido - esperado, esperado - leido
    skus_pag = Counter(("Sin SKU" if r.get("sin") else r["s"]) for r in doc["R"])
    skus_xls = Counter(t[0] for t in leido.elements())
    return {"filas_excel": n, "filas_pagina": len(doc["R"]), "iguales": sum((leido & esperado).values()),
            "sobran": sum(sobran.values()), "faltan": sum(faltan.values()),
            "skus_iguales": skus_pag == skus_xls,
            "ejemplo_distinto": [list(map(str, x)) for x in list(faltan)[:2]] + [list(map(str, x)) for x in list(sobran)[:2]]}


def _sin_repetir(ruta: Path) -> tuple[int, int]:
    """Reescribe el .xlsx para que las fotos idénticas compartan un solo archivo. Devuelve (antes, después)."""
    with zipfile.ZipFile(ruta) as z:
        nombres = z.namelist()
        contenido = {n: z.read(n) for n in nombres}
    medios = [n for n in nombres if n.startswith("xl/media/")]
    primero: dict[str, str] = {}
    cambia: dict[str, str] = {}
    for n in medios:
        h = hashlib.md5(contenido[n]).hexdigest()
        if h in primero:
            cambia[n.split("/")[-1]] = primero[h].split("/")[-1]
        else:
            primero[h] = n
    if not cambia:
        return len(medios), len(medios)
    patron = re.compile(rb'Target="(?:\.\./|/xl/)media/([^"]+)"')
    for n in nombres:
        if n.startswith("xl/drawings/_rels/") and n.endswith(".rels"):
            contenido[n] = patron.sub(
                lambda m: m.group(0).replace(m.group(1), cambia.get(m.group(1).decode(), m.group(1).decode()).encode()),
                contenido[n])
    quitar = {"xl/media/" + k for k in cambia}
    tmp = ruta.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for n in nombres:
            if n in quitar:
                continue
            # las fotos ya vienen comprimidas: guardarlas tal cual abre más rápido
            z.writestr(n, contenido[n], compress_type=zipfile.ZIP_STORED if n.startswith("xl/media/") else zipfile.ZIP_DEFLATED)
    tmp.replace(ruta)
    return len(medios), len(medios) - len(cambia)


def construir(salida: Path) -> dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as Foto
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.units import pixels_to_EMU

    doc = leer_json(salida / "datos_pl.json")
    if not doc:
        raise RuntimeError("falta datos_pl.json: corre antes `pagina_pl`")
    rutas_cat = doc["K"]

    def cat(r: dict[str, Any]) -> tuple[str, str]:
        if r.get("k") is None:
            return ("Sin categoría", "")
        k = rutas_cat[r["k"]]
        return (k[0], k[1] or "")

    def precio(r: dict[str, Any]) -> tuple[float | None, str]:
        """La MEDIA de Mercado Libre de la fila, EXACTAMENTE la que la página pone en su columna «media»
        (la función `celdaPrecio` de la plantilla), y de dónde sale. Si se cambia allá, se cambia aquí."""
        return media_de_la_pagina(r)

    filas = sorted(doc["R"], key=lambda r: (cat(r)[0] == "Sin categoría", cat(r)[0], cat(r)[1],
                                           -(r.get("pq") or 0), r["s"]))
    fotos = _recortes(salida, filas, doc)
    aviso(f"excel: {len(filas)} filas · {len(fotos)} fotos distintas recortadas")

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Productos"
    cab = CABECERA
    hoja.append(cab)
    oscuro = PatternFill("solid", fgColor="1F2937")
    raya = Side(style="thin", color="E5E7EB")
    for c in hoja[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = oscuro
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hoja.row_dimensions[1].height = 34
    for col, ancho in zip("ABCDEFGHIJKLM", (ANCHO_COL, 24, 52, 28, 30, 13, 11, 12, 12, 16, 36, 17, 22)):
        hoja.column_dimensions[col].width = ancho

    margen = pixels_to_EMU(4)
    con_foto = 0
    centro = Alignment(vertical="center", wrap_text=True)
    derecha = Alignment(vertical="center", horizontal="right")
    for n, r in enumerate(filas, start=2):
        hoja.append([None, *fila_de_la_pagina(r, rutas_cat)])
        hoja.row_dimensions[n].height = ALTO_FILA
        for j in range(2, 6):
            hoja.cell(row=n, column=j).alignment = centro
        hoja.cell(row=n, column=2).font = Font(bold=True, name="Consolas")
        for j in range(6, 10):
            celda = hoja.cell(row=n, column=j)
            celda.alignment, celda.number_format = derecha, "#,##0"
        celda = hoja.cell(row=n, column=10)
        celda.alignment, celda.number_format = derecha, '"$"#,##0.00'
        hoja.cell(row=n, column=11).alignment = centro
        celda = hoja.cell(row=n, column=12)
        celda.alignment, celda.number_format = derecha, '"$"#,##0'
        hoja.cell(row=n, column=13).alignment = centro
        for j in range(1, 14):
            hoja.cell(row=n, column=j).border = Border(bottom=raya)
        ruta = fotos.get(("g", r["i"])) if r.get("i") is not None else (
            fotos.get(("q", r["j"])) if r.get("j") is not None else None)
        if ruta is not None and ruta.exists():
            foto = Foto(str(ruta))
            ancla = TwoCellAnchor(editAs="twoCell")                # se mueve y cambia de tamaño con su celda
            ancla._from = AnchorMarker(col=0, colOff=margen, row=n - 1, rowOff=margen)
            ancla.to = AnchorMarker(col=0, colOff=margen + pixels_to_EMU(CAJA), row=n - 1, rowOff=margen + pixels_to_EMU(CAJA))
            foto.anchor = ancla
            hoja.add_image(foto)
            con_foto += 1
    hoja.freeze_panes = "C2"
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{len(filas) + 1}"

    # ── Por categoría: el mismo resumen de la pestaña de la página, plegable ──────────────
    res = wb.create_sheet("Por categoría")
    cab2 = ["Categoría", "Subcategoría", "SKUs", "Productos sin SKU", "Con mercancía hoy", "SOLD OUT",
            "Comprado", "Salió", "Queda"]
    res.append(cab2)
    for c in res[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), oscuro
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    res.row_dimensions[1].height = 32
    grupos: dict[str, dict[str, list[dict[str, Any]]]] = {}
    for r in filas:
        if r.get("pc") is None:
            continue
        c1, c2 = cat(r)
        grupos.setdefault(c1, {}).setdefault(c2 or "(sin subcategoría)", []).append(r)

    def suma(l: list[dict[str, Any]]) -> list[Any]:
        return [sum(1 for r in l if not r.get("sin")), sum(1 for r in l if r.get("sin")),
                sum(1 for r in l if not r.get("so") and (r.get("pq") or 0) > 0), sum(1 for r in l if r.get("so")),
                round(sum(r["pc"] for r in l)), round(sum(min(r.get("pz") or r.get("ps") or 0, r["pc"]) for r in l)),
                round(sum(r["pq"] for r in l))]

    for c1 in sorted(grupos, key=lambda k: -sum(r["pq"] for l in grupos[k].values() for r in l)):
        todas = [r for l in grupos[c1].values() for r in l]
        res.append([c1, "", *suma(todas)])
        for c in res[res.max_row]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="F3F4F6")
        for c2 in sorted(grupos[c1], key=lambda k: -sum(r["pq"] for r in grupos[c1][k])):
            res.append(["", c2, *suma(grupos[c1][c2])])
            res.row_dimensions[res.max_row].outlineLevel = 1
    total = [r for r in filas if r.get("pc") is not None]
    res.append(["TOTAL", "", *suma(total)])
    for c in res[res.max_row]:
        c.font = Font(bold=True)
    for fila in res.iter_rows(min_row=2, min_col=3):
        for c in fila:
            c.number_format, c.alignment = "#,##0", Alignment(horizontal="right")
    for col, ancho in zip("ABCDEFGHI", (34, 38, 10, 14, 14, 11, 13, 12, 13)):
        res.column_dimensions[col].width = ancho
    res.freeze_panes = "C2"
    res.sheet_properties.outlinePr.summaryBelow = False

    destino = salida / "inventario_kubera_skus.xlsx"
    wb.save(destino)
    antes, despues = _sin_repetir(destino)
    peso = destino.stat().st_size / 1e6
    aviso(f"excel: {destino.name} · {len(filas)} filas · {con_foto} con foto · {antes} imágenes → {despues} sin repetir · {peso:.1f} MB")
    chequeo = comprobar(destino, doc)
    aviso(f"excel: comprobación contra la página: {chequeo['iguales']} de {chequeo['filas_pagina']} filas idénticas · "
          f"sobran {chequeo['sobran']} · faltan {chequeo['faltan']} · mismos SKUs: {chequeo['skus_iguales']}")
    if chequeo["sobran"] or chequeo["faltan"] or not chequeo["skus_iguales"]:
        raise RuntimeError(f"el Excel NO cuadra con la página: {chequeo}")
    return {"filas": len(filas), "con_foto": con_foto, "imagenes": despues, "mb": round(peso, 1),
            "con_media_ml": sum(1 for r in filas if precio(r)[0] is not None), "comprobacion": chequeo}
