"""
pl_leer.py — Lee los packing lists ya bajados y deja, por renglón, qué se compró.

De dónde sale la lógica
-----------------------
Del proyecto `kubera-exit` de José (github.com/joseKubera/kubera-exit, commit
42ae62b, 6-oct-2026), que ya resolvió lo difícil: cada proveedor y cada persona de
bodega arma el Excel a su manera. Aquí se copia su lectura —no se importa— con dos
cambios: no depende del backend (usa `pl_parser.py`, la copia del lector de
producción) y de cada foto guarda además su HUELLA, que es con lo que después se
empata un renglón con su SKU.

Dos clases de archivo
---------------------
  · ORIGINAL    el packing list del proveedor: descripción, cajas, piezas, precio.
                No trae SKU. Es «lo que se compró».
  · VALIDADO    el mismo archivo con lo que agregó bodega (carpeta Ferraforme):
                la columna «SKU ODOO» y las piezas contadas en físico. Es el puente
                entre un renglón del proveedor y un SKU.

Los ajustes por archivo (`OVERRIDES`, `GEN_OVR`, `FORZAR_GENERICO`) son los de José:
encabezados que mienten, columnas corridas, el nombre del producto solo en la
cabecera del grupo. Van por `id` de `costing.packing_archivos`.

Cada archivo leído se guarda en `<cache>/leido/`, así que volver a correr solo lee
lo que cambió.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import posixpath
import re
import warnings
import zipfile
from collections import Counter
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from comun import ahora_iso, aviso, escribir_json, leer_json

LADO = 72
NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
RE_NUM = re.compile(r"cont(?:enedor)?\.?\s*#?\s*(\d{1,3})\b", re.I)
RE_SKU = re.compile(r"^[A-Z]{2,5}-\d{3,5}(-[A-Z0-9.#/]+)*$", re.I)
HOJAS_RESUMEN = re.compile(r"oc vs|tabla|^hoja|analisis|ingreso|limpieza|pendiente|costo|variacion|checklist", re.I)

# Encabezados (normalizados) por prioridad. Primero lo que capturó bodega.
PIEZAS = ["total de piezas", "total de pz", "piezas totales", "piezas totales en físico", "pz total",
          "total pz", "total pzs", "pz * total", "piezas físicas", "cantidad validada",
          "validación unidades", "pzs", "piezas", "total"]
FISICO_AMBIGUO = ["cantidad fisica", "cantidades fisicas", "cantidad física", "conteo fisico", "fisico"]
CAJAS = ["total de cajas", "cantidad cajas", "cantidad de cajas", "cajas fisicas", "no. de cajas",
         "cantidad fisica (cajas)", "cantidad fisica cajas", "ctns fisico", "cajas completas2",
         "qty cajas", "total cntnr", "cajas", "caja"]
PZ_X_CAJA = ["piezas por caja", "piezas x caja", "pz x caja", "pack * caja", "*pz por caja",
             "cantidad de piezas x caja", "pz*caja", "piezas en cajas", "ct x caja", "cantidad x caja"]
QTY_PL = ["total qty", "总个数", "total pcs", "总产品数量", "总数量", "piezas pl", "qty(pcs)", "t.qty", "qty"]
CTN_PL = ["totai ctns", "total ctns", "total ctn", "箱数", "cajas pl", "packgage(ctns)", "件数", "ctn"]
PRECIO = ["单价", "price", "precio/usd", "costo usd", "costo dlls"]
TITULO = ["descripción", "descripcion", "nombre producto", "description of goods", "英文品名",
          "english name", "description", "des.", "name", "spanish", "产品英文品名", "chino", "中文品名",
          "chinese name"]
FOTO = ["picture", "图片", "product photo", "foto", "产品图片", "imagen"]
# id de costing.packing_archivos → columna forzada cuando el encabezado miente.
OVERRIDES = {154: {"pz": 24}}

# Respaldo para los originales que el lector de producción no entiende.
G_QTY = ["t.qty", "total pcs", "总件数", "产品总件数", "piezas_totales", "total qty", "总个数", "总数量", "总数", "qty(pcs)"]
G_CTN = ["totai ctns", "total ctns", "num_cajas", "箱数", "ctns", "件数"]
G_PXC = ["quantity per box", "pcs/ctn", "每箱件数", "产品申报数量", "装箱数", "单箱数量"]
G_PRE = ["precio_usd", "price", "单价", "货值 valor usd"]
G_TIT = ["descripcion en es", "spanish", "description of goods", "producto", "产品英文品名", "英文品名", "品名", "产品中文品名"]
GEN_OVR = {3: {"shift": 1}, 35: {"pxc": 20, "qty": None}, 70: {"carry": True}}
FORZAR_GENERICO = {35, 50}


# ── utilidades ────────────────────────────────────────────────────────────────

def num_contenedor(nombre: str) -> int | None:
    m = RE_NUM.search(nombre or "")
    return int(m.group(1)) if m else None


def norm(x: Any) -> str:
    return re.sub(r"\s+", " ", str(x or "").replace("\xa0", " ")).strip().lower()


def num(v: Any) -> float:
    if v is None or isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))
    try:
        return float(s) if s not in ("", ".", "-") else 0.0
    except ValueError:
        return 0.0


def buscar(h: list[str], claves: list[str], exacto: bool = False, excluir: tuple = (),
           cortas: bool = False) -> int | None:
    for k in claves:
        if cortas and len(k) <= 6:      # «total», «pzs», «cajas»: solo exactas, o pescan «total ctns»
            i = next((i for i, c in enumerate(h) if c == k), None)
            if i is not None:
                return i
            continue
        for i, c in enumerate(h):
            if not c or any(e in c for e in excluir):
                continue
            if (c == k) if exacto else (c == k or c.startswith(k + " ") or k in c):
                return i
    return None


def _resolver(base_dir: str, target: str, nombres: set[str]) -> str | None:
    t = target.replace("\\", "/")
    cands = [t.lstrip("/")] if t.startswith("/") else [
        posixpath.normpath(posixpath.join(base_dir, t)), f"xl/{t}", t]
    return next((c for c in cands if c in nombres), None)


def _rels(z: zipfile.ZipFile, path: str, nombres: set[str]) -> dict[str, str | None]:
    if path not in nombres:
        return {}
    base = posixpath.dirname(posixpath.dirname(path))
    return {r.attrib["Id"]: _resolver(base, r.attrib["Target"], nombres)
            for r in ET.fromstring(z.read(path)).findall("rel:Relationship", NS)}


def fotos_hoja(xlsx: bytes, hoja: str, columna: int | None = None) -> dict[int, bytes]:
    """{fila_0based: bytes} de las fotos ancladas en la hoja `hoja` (los validados
    traen varias hojas con fotos; solo se mira el dibujo de esa)."""
    with zipfile.ZipFile(io.BytesIO(xlsx)) as z:
        nombres = set(z.namelist())
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        wbrels = _rels(z, "xl/_rels/workbook.xml.rels", nombres)
        hoja_path = None
        for s in wb.find("m:sheets", NS):
            if s.attrib.get("name") == hoja:
                hoja_path = wbrels.get(s.attrib[f"{{{NS['r']}}}id"])
        if not hoja_path:
            return {}
        srels = _rels(z, posixpath.join(posixpath.dirname(hoja_path), "_rels",
                                        posixpath.basename(hoja_path) + ".rels"), nombres)
        anclas = []
        for draw in {v for v in srels.values() if v and "/drawings/drawing" in v}:
            drels = _rels(z, posixpath.join(posixpath.dirname(draw), "_rels",
                                            posixpath.basename(draw) + ".rels"), nombres)
            root = ET.fromstring(z.read(draw))
            for a in list(root.findall("xdr:twoCellAnchor", NS)) + list(root.findall("xdr:oneCellAnchor", NS)):
                frm, blip = a.find("xdr:from", NS), a.find(".//a:blip", NS)
                if frm is None or blip is None:
                    continue
                media = drels.get(blip.attrib.get(f"{{{NS['r']}}}embed"))
                if media:
                    anclas.append((int(frm.find("xdr:col", NS).text),
                                   int(frm.find("xdr:row", NS).text), media))
        if not anclas:
            return {}
        if columna is None or not any(c == columna for c, _, _ in anclas):
            columna = Counter(c for c, _ in {(c, r) for c, r, _ in anclas}).most_common(1)[0][0]
        out: dict[int, bytes] = {}
        for c, r, media in anclas:
            if c == columna and r not in out:
                try:
                    out[r] = z.read(media)
                except Exception:  # noqa: BLE001
                    pass
        return out


LADO_GRANDE = 320
_GRANDES: Any = False          # (carpeta, nombres que se quieren) · None si no se pidió · False = sin mirar


def _grandes() -> tuple[Path, set[str]] | None:
    """¿Se pidió guardar también la foto a tamaño completo? Lo dice la variable de entorno
    `PL_FOTOS_GRANDES` (una carpeta con `_quiero.json`: los nombres de miniatura que hacen falta).
    Va por el entorno para que llegue a los procesos hijos."""
    global _GRANDES
    if _GRANDES is False:
        carpeta = os.environ.get("PL_FOTOS_GRANDES") or ""
        lista = Path(carpeta) / "_quiero.json" if carpeta else None
        _GRANDES = (Path(carpeta), set(json.loads(lista.read_text(encoding="utf-8")))) if lista and lista.exists() else None
    return _GRANDES


def _guardar_grande(datos: bytes, destino: Path) -> None:
    """La misma foto del renglón, sin achicarla a miniatura: hasta LADO_GRANDE px por lado."""
    from PIL import Image

    try:
        im = Image.open(io.BytesIO(datos)).convert("RGBA")
    except Exception:  # noqa: BLE001
        return
    fondo = Image.new("RGB", im.size, (255, 255, 255))
    fondo.paste(im, mask=im.split()[3])
    fondo.thumbnail((LADO_GRANDE, LADO_GRANDE), Image.LANCZOS)
    destino.parent.mkdir(parents=True, exist_ok=True)
    fondo.save(destino, "JPEG", quality=88)


def huella(datos: bytes, destino: Path) -> dict[str, str] | None:
    """Miniatura a disco + las dos huellas de la foto: `s` (sha1 del archivo tal cual
    viene incrustado: si coincide con la de Odoo es EL MISMO archivo) y `d` (dHash
    de 64 bits: la misma imagen aunque la hayan vuelto a guardar)."""
    from PIL import Image

    try:
        im = Image.open(io.BytesIO(datos))
        im.draft("RGB", (LADO * 2, LADO * 2))
        im = im.convert("RGBA")
    except Exception:  # noqa: BLE001 — una foto rota no tumba el archivo
        return None
    fondo = Image.new("RGB", im.size, (255, 255, 255))
    fondo.paste(im, mask=im.split()[3])
    gris = fondo.convert("L").resize((9, 8), Image.LANCZOS)
    px = list(gris.getdata())
    bits = 0
    for y in range(8):
        for x in range(8):
            bits = (bits << 1) | (1 if px[y * 9 + x] > px[y * 9 + x + 1] else 0)
    fondo.thumbnail((LADO, LADO))
    lienzo = Image.new("RGB", (LADO, LADO), (255, 255, 255))
    lienzo.paste(fondo, ((LADO - fondo.width) // 2, (LADO - fondo.height) // 2))
    destino.parent.mkdir(parents=True, exist_ok=True)
    lienzo.save(destino, "JPEG", quality=62)
    pedidas = _grandes()
    if pedidas and destino.name in pedidas[1]:
        _guardar_grande(datos, pedidas[0] / destino.name)
    return {"s": hashlib.sha1(datos).hexdigest(), "d": f"{bits:016x}"}


# ── VALIDADO (lo que contó bodega) ────────────────────────────────────────────

def _col_sku(h: list[str], filas: list[tuple]) -> int | None:
    cand = [i for i, c in enumerate(h) if "sku" in c and not re.search(
        r"padre|pdre|\bpl\b|origen|duplicado|similar|\boc\b|#", c)]
    if not cand:
        return None

    def puntos(i: int) -> tuple[int, int]:
        n = sum(1 for f in filas if i < len(f) and f[i] and RE_SKU.match(re.sub(r"[\[\]\s]", "", str(f[i]))))
        return (n, 1 if re.search(r"od|dd|0dd", h[i]) else 0)
    return max(cand, key=puntos)


def _elegir_hoja(wb: Any) -> tuple | None:
    mejor = None
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        for hi, r in enumerate(filas[:20]):
            h = [norm(c) for c in r]
            if not any("sku" in c for c in h):
                continue
            datos = filas[hi + 1:]
            i = _col_sku(h, datos)
            if i is None:
                continue
            n = sum(1 for f in datos if i < len(f) and f[i] and RE_SKU.match(re.sub(r"[\[\]\s]", "", str(f[i]))))
            # Bodega llenó piezas pero no SKU: la hoja cuenta igual, para no perder sus piezas.
            tiene_pz = buscar(h, PIEZAS[:6], cortas=True) is not None
            clave = (0 if HOJAS_RESUMEN.search(ws.title) else 1, tiene_pz, n)
            if (n or tiene_pz) and (mejor is None or clave > mejor[0]):
                mejor = (clave, ws.title, hi, h, datos, i)
            break
    return mejor


def _validado(meta: dict[str, Any], ruta: Path, thumbs: Path) -> dict[str, Any]:
    import openpyxl

    nombre = meta.get("miembro") or meta["nombre"] or ""
    res: dict[str, Any] = {"id": meta["id"], "archivo": nombre, "contenedor": meta.get("contenedor_base"),
                           "num": num_contenedor(nombre), "filas": [], "aviso": ""}
    crudo = ruta.read_bytes()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(crudo), data_only=True)
    except Exception as e:  # noqa: BLE001
        res["aviso"] = f"ilegible: {e}"
        return res
    m = _elegir_hoja(wb)
    if not m:
        res["aviso"] = "sin columna SKU ODOO"
        return res
    _, hoja, hi, h, datos, i_sku = m
    i_pz = buscar(h, PIEZAS, cortas=True, excluir=("x caja", "por caja", "x pack", "* pack", "en cajas", "pl", "qty", "总",
                                                   "solicitadas", "*pack", "por tarima", "por estiba", "por paquete"))
    i_pz = OVERRIDES.get(meta["id"], {}).get("pz", i_pz)
    i_fis = buscar(h, FISICO_AMBIGUO, exacto=True) if i_pz is None else None
    i_cj = buscar(h, CAJAS, cortas=True, excluir=("x caja", "por caja", "pl", "por tarima", "en tarima", "en cajas", "*", "箱数"))
    i_pxc = buscar(h, PZ_X_CAJA)
    i_qty = buscar(h, QTY_PL, cortas=True, excluir=("value", "货值", "por", "x caja", "validado"))
    i_ctn = buscar(h, CTN_PL, cortas=True, excluir=("por", "x caja", "tarima"))
    i_pr = buscar(h, PRECIO, excluir=("total", "总价", "货值"))
    i_ti = buscar(h, TITULO, excluir=("sku",))
    i_foto = buscar(h, FOTO, excluir=("imagenar",))

    # «cantidad física» a veces son cajas y a veces piezas: se decide comparando
    # contra las columnas del proveedor, renglón por renglón.
    fis_es_cajas = False
    if i_fis is not None and i_qty is not None and i_ctn is not None:
        vs_c = vs_q = 0
        for f in datos:
            g = lambda j: num(f[j]) if j < len(f) else 0  # noqa: E731
            fv, q, c = g(i_fis), g(i_qty), g(i_ctn)
            if fv and q and c and q != c:
                if abs(fv - c) < abs(fv - q):
                    vs_c += 1
                else:
                    vs_q += 1
        fis_es_cajas = vs_c > vs_q
        if fis_es_cajas and i_cj is None:
            i_cj = i_fis

    fotos = fotos_hoja(crudo, hoja, i_foto)
    for k, f in enumerate(datos):
        fila0 = hi + 1 + k
        g = lambda j: (f[j] if j is not None and j < len(f) else None)  # noqa: E731
        sku = re.sub(r"[\[\]\s]", "", str(g(i_sku) or "")).upper()
        titulo = str(g(i_ti) or "").strip()
        sku_ok = bool(RE_SKU.match(sku))
        if not sku_ok and not titulo:
            continue
        qty_pl, ctn_pl = num(g(i_qty)), num(g(i_ctn))
        cajas = num(g(i_cj))
        if i_pz is not None:
            piezas = num(g(i_pz))
        elif i_fis is not None:
            fv = num(g(i_fis))
            if fis_es_cajas:
                pxc = num(g(i_pxc)) or (qty_pl / ctn_pl if ctn_pl else 0)
                piezas = fv * pxc
                cajas = cajas or fv
            else:
                piezas = fv
        elif i_cj is not None and (i_pxc is not None or ctn_pl):
            pxc = num(g(i_pxc)) or (qty_pl / ctn_pl if ctn_pl else 0)
            piezas = cajas * pxc
        else:
            piezas = qty_pl
        if not sku_ok and not piezas and not qty_pl:
            continue
        if titulo.lower().startswith(("total", "合计", "observ", "variante")):
            continue
        fila: dict[str, Any] = {"fila": fila0, "sku": sku if sku_ok else "", "titulo": titulo[:140],
                                "cajas": round(cajas, 3), "piezas": round(piezas, 3),
                                "qty_pl": round(qty_pl, 3), "ctn_pl": round(ctn_pl, 3),
                                "precio_usd": round(num(g(i_pr)), 4)}
        if not sku_ok and sku:
            fila["sku_raw"] = sku[:40]
        if fila0 in fotos:
            hu = huella(fotos[fila0], thumbs / f"v_{meta['id']}_{fila0}.jpg")
            if hu:
                fila["foto"] = hu
        res["filas"].append(fila)
    nom = lambda i: f"{i}:{h[i][:22]}" if i is not None else "-"  # noqa: E731
    res["cols"] = (f"hoja={hoja} enc={hi} sku={nom(i_sku)} pz={nom(i_pz)} fis={nom(i_fis)}"
                   f"{'(=cajas)' if fis_es_cajas else ''} cj={nom(i_cj)} qty={nom(i_qty)} ctn={nom(i_ctn)} "
                   f"precio={nom(i_pr)} tit={nom(i_ti)} foto={nom(i_foto)}")
    return res


# ── ORIGINAL (lo que mandó el proveedor) ──────────────────────────────────────

def _generico(meta: dict[str, Any], crudo: bytes) -> list[dict[str, Any]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(crudo), data_only=True)
    mejor = None
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        for hi, r in enumerate(filas[:20]):
            h = [norm(c) for c in r]
            puntos = sum(buscar(h, L) is not None for L in (G_QTY, G_CTN, G_PXC, G_TIT))
            if puntos >= 2 and (mejor is None or puntos > mejor[0]):
                mejor = (puntos, ws.title, hi, h, filas[hi + 1:])
    if not mejor:
        return []
    _, hoja, hi, h, cuerpo = mejor
    ovr = GEN_OVR.get(meta["id"], {})
    sh = ovr.get("shift", 0)
    col = lambda L, **k: (lambda i: None if i is None else i + sh)(buscar(h, L, **k))  # noqa: E731
    i_q, i_c, i_p = col(G_QTY, excluir=("value", "货值")), col(G_CTN, excluir=("per", "每")), col(G_PXC)
    i_pr, i_t = col(G_PRE, excluir=("total", "总")), buscar(h, G_TIT)
    i_t = None if i_t is None else i_t + sh
    i_q = ovr.get("qty", i_q)
    i_p = ovr.get("pxc", i_p)
    fotos = fotos_hoja(crudo, hoja)
    out = []
    previo = ""
    for k, f in enumerate(cuerpo):
        g = lambda j: f[j] if j is not None and j < len(f) else None  # noqa: E731
        # Renglones de un mismo grupo traen el nombre solo en la cabecera.
        titulo = str(g(i_t) or "").strip() or (previo if ovr.get("carry") else "")
        previo = titulo
        cajas, q = num(g(i_c)), num(g(i_q))
        piezas = q or cajas * num(g(i_p))      # en GEN_OVR, `pxc` y `qty` son ÍNDICES de columna
        if not titulo or not piezas or titulo.lower().startswith(("total", "合计")):
            continue
        out.append({"fila": hi + 1 + k, "titulo": titulo, "cajas": cajas, "piezas": piezas,
                    "precio_usd": num(g(i_pr)), "img": fotos.get(hi + 1 + k)})
    return out


def _original(meta: dict[str, Any], ruta: Path, thumbs: Path) -> dict[str, Any]:
    import pl_parser

    res: dict[str, Any] = {"id": meta["id"], "archivo": meta["nombre"] or "",
                           "contenedor": meta.get("contenedor_base"),
                           "num": num_contenedor(meta["nombre"] or ""), "filas": [], "aviso": ""}
    crudo = ruta.read_bytes()
    try:
        out = pl_parser.leer(crudo)
    except Exception as e:  # noqa: BLE001
        out = {"filas": [], "imagenes": {}, "avisos": [f"leer: {str(e)[:160]}"]}
    for f in out["filas"]:
        titulo = (f["producto"] or "").strip()
        if titulo.lower().startswith(("total", "合计", "小计")):
            continue
        if not (f["piezas"] or f["cajas"]):
            continue
        fila: dict[str, Any] = {
            "fila": f["fila_idx"], "titulo": titulo[:140], "titulo_chn": (f["producto_chn"] or "")[:80],
            # caja compartida: las cajas se reparten entre los renglones que la comparten
            "cajas": round(f["cajas"] / f["tam_grupo"] if f["comparte_caja"] else f["cajas"], 3),
            "piezas": round(f["piezas"], 3), "precio_usd": round(f["precio_usd"] or 0, 4),
            "cbm": round(f["cbm_por_pieza"] or 0, 8)}
        img = out["imagenes"].get(f["fila_idx"])
        if img:
            hu = huella(img, thumbs / f"o_{meta['id']}_{f['fila_idx']}.jpg")
            if hu:
                fila["foto"] = hu
        res["filas"].append(fila)
    res["avisos"] = out["avisos"][:6]
    if meta["id"] in FORZAR_GENERICO or not sum(x["piezas"] for x in res["filas"]):
        res["filas"] = []
        for f in _generico(meta, crudo):
            fila = {"fila": f["fila"], "titulo": f["titulo"][:140], "titulo_chn": "",
                    "cajas": round(f["cajas"], 3), "piezas": round(f["piezas"], 3),
                    "precio_usd": round(f["precio_usd"], 4), "cbm": 0}
            if f["img"]:
                hu = huella(f["img"], thumbs / f"o_{meta['id']}_{f['fila']}.jpg")
                if hu:
                    fila["foto"] = hu
            res["filas"].append(fila)
        res["aviso"] = "lector genérico"
    return res


# ── orquesta ──────────────────────────────────────────────────────────────────

def _uno(tarea: tuple[dict[str, Any], str, str]) -> dict[str, Any]:
    meta, cache_txt, _ = tarea
    warnings.simplefilter("ignore")
    cache = Path(cache_txt)
    ruta = cache / meta["tipo"] / f"{meta['id']}.xlsx"
    guardado = cache / "leido" / f"{meta['tipo']}_{meta['id']}.json"
    firma = f"{ruta.stat().st_size}:{int(ruta.stat().st_mtime)}:v2"
    previo = leer_json(guardado)
    if previo and previo.get("_firma") == firma:
        return previo
    thumbs = cache / "thumbs"
    try:
        res = (_original if meta["tipo"] == "original" else _validado)(meta, ruta, thumbs)
    except Exception as e:  # noqa: BLE001 — un archivo roto no detiene a los demás
        res = {"id": meta["id"], "archivo": meta.get("nombre") or "", "contenedor": meta.get("contenedor_base"),
               "num": None, "filas": [], "aviso": f"FALLÓ: {type(e).__name__}: {str(e)[:160]}"}
    res["tipo"], res["_firma"] = meta["tipo"], firma
    guardado.parent.mkdir(parents=True, exist_ok=True)
    guardado.write_text(json.dumps(res, ensure_ascii=False), encoding="utf-8")
    return res


def leer_todos(salida: Path, cache: Path, procesos: int = 4) -> dict[str, Any]:
    indice = leer_json(salida / "datos" / "pl_indice.json")
    if not indice:
        raise RuntimeError("falta datos/pl_indice.json: corre primero la etapa `pl_bajar`")
    metas = [a for a in indice["archivos"] if a["local"].get("ok")]
    # Los grandes primero: así los procesos no se quedan todos con un archivo de 100 MB al final.
    metas.sort(key=lambda a: -(a["local"].get("bytes") or 0))
    tareas = [(a, str(cache), "") for a in metas]
    resultados: list[dict[str, Any]] = []
    with ProcessPoolExecutor(max_workers=max(1, procesos)) as pool:
        for i, r in enumerate(pool.map(_uno, tareas), 1):
            resultados.append(r)
            if i % 10 == 0 or i == len(tareas):
                aviso(f"packing lists leídos: {i}/{len(tareas)}")
    orig = sorted((r for r in resultados if r["tipo"] == "original"), key=lambda r: r["id"])
    val = sorted((r for r in resultados if r["tipo"] == "ferraforme"), key=lambda r: r["id"])
    for r in resultados:
        r.pop("_firma", None)
    doc = {"leido": ahora_iso(), "cache": str(cache), "originales": orig, "validados": val}
    escribir_json(salida / "datos" / "pl.json", doc)

    def _res(lista: list[dict[str, Any]], campo: str) -> dict[str, Any]:
        return {"archivos": len(lista), "con_renglones": sum(1 for r in lista if r["filas"]),
                "renglones": sum(len(r["filas"]) for r in lista),
                "piezas": round(sum(f[campo] for r in lista for f in r["filas"])),
                "con_foto": sum(1 for r in lista for f in r["filas"] if f.get("foto"))}

    resumen = {"originales": _res(orig, "piezas"), "validados": _res(val, "piezas"),
               "validados_qty_pl": round(sum(f["qty_pl"] for r in val for f in r["filas"])),
               "validados_con_sku": sum(1 for r in val for f in r["filas"] if f["sku"]),
               "skus_distintos": len({f["sku"] for r in val for f in r["filas"] if f["sku"]}),
               "sin_leer": [f"{r['tipo']} #{r['id']} {r['archivo'][:40]}: {r['aviso']}"
                            for r in resultados if not r["filas"]][:40]}
    aviso(f"packing lists: originales {resumen['originales']} · validados {resumen['validados']}")
    return resumen
